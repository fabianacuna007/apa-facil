"""
Paso 9.7: convierte un texto pegado en un documento Word, para procesarlo con el
mismo núcleo que un .docx subido.

Cada línea del texto es un párrafo. Las líneas vacías (o solo con espacios)
separan párrafos y no se copian: el núcleo quitaría igualmente los párrafos
vacíos. El texto de cada línea se copia tal cual, con sus espacios y sus
tabuladores. Los saltos de línea de Word (Mayús+Intro) y de página que traiga
el texto también separan párrafos.

El texto pegado no trae formato (ni negritas ni tamaños), así que los títulos
se reconocen solo por su forma: líneas cortas sin punto final, con numeración o
en mayúsculas (paso 3.2).

La primera línea es el título del trabajo (la web lo pide así: «pega tu trabajo
empezando por el título, sin la portada»). Si parece un título (corta y sin punto
final), lleva detrás un salto de página: así el núcleo la trata como la portada
que solo tiene el título, la repite al principio del texto (paso 4.4) y, si hay
datos de portada, la sustituye por la portada APA (paso 5.3). Sin el salto, el
núcleo tomaría por portada todas las líneas cortas del principio, también
«Introducción», porque sin formato no puede distinguirlas. Si la primera línea
es un párrafo, no hay salto: el núcleo decide como con cualquier documento.

El documento es A4 (el tamaño de papel de España y de casi toda Hispanoamérica;
python-docx crea uno en tamaño carta) y sin estilos propios: las reglas ponen
después los márgenes, la fuente y el resto del formato APA.
"""

import re

import docx
from docx.enum.text import WD_BREAK
from docx.shared import Cm

from apa_facil.documento import ErrorDocumento
from apa_facil.estructura import FINAL_DE_FRASE, PALABRAS_DE_TITULO

NOMBRE = "texto_pegado.docx"
MAXIMO = 500_000   # caracteres: un TFG largo tiene unos 150 000

# Saltos de línea de cualquier sistema, y los de Word: \x0b (Mayús+Intro) y \x0c (salto de página).
SALTOS = re.compile(r"\r\n|[\r\n\x0b\x0c\u2028\u2029\x85]")
# Caracteres que Word no admite en un documento (XML 1.0): los de control, salvo el tabulador.
NO_ADMITIDOS = re.compile(r"[\x00-\x08\x0e-\x1f\x7f\ud800-\udfff\ufffe\uffff]")


def parrafos(texto):
    """Los párrafos del texto: una línea, un párrafo, sin las líneas vacías."""
    return [linea for linea in SALTOS.split(texto) if linea.strip()]


def empieza_por_el_titulo(lineas):
    """Si la primera línea parece el título del trabajo: corta, sin punto final y con texto detrás."""
    primera = lineas[0].strip() if lineas else ""
    return len(lineas) > 1 and len(primera.split()) < PALABRAS_DE_TITULO and not FINAL_DE_FRASE.search(primera)


def comprobar(texto):
    """Si el texto no se puede convertir, un ErrorDocumento que explica qué hacer."""
    if not texto or not texto.strip():
        raise ErrorDocumento("No has pegado ningún texto. Copia tu trabajo y pégalo en el cuadro.")
    if len(texto) > MAXIMO:
        miles = lambda n: f"{n:,}".replace(",", ".")
        raise ErrorDocumento(f"El texto tiene {miles(len(texto))} caracteres, y el máximo es {miles(MAXIMO)}. "
                             "Si es tan largo, sube el documento de Word.")
    raro = NO_ADMITIDOS.search(texto)
    if raro:
        antes = texto[max(0, raro.start() - 30):raro.start()].split("\n")[-1].strip()
        donde = f", después de «{antes}»" if antes else ""
        raise ErrorDocumento(f"El texto tiene un carácter invisible que Word no admite (código "
                             f"U+{ord(raro.group()):04X}){donde}. Bórralo y vuelve a pegarlo, o sube el "
                             "documento de Word.")


def crear(texto, ruta):
    """Guarda en `ruta` un .docx con el texto, un párrafo por línea."""
    comprobar(texto)
    documento = docx.Document()
    seccion = documento.sections[0]
    seccion.page_width, seccion.page_height = Cm(21), Cm(29.7)
    lineas = parrafos(texto)
    for parrafo in lineas:
        documento.add_paragraph(parrafo)
    if empieza_por_el_titulo(lineas):
        documento.paragraphs[0].add_run().add_break(WD_BREAK.PAGE)
    documento.save(ruta)
    return ruta
