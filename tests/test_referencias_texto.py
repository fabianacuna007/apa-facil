"""
Paso 7.2: lector de referencias (apa_facil/revision/referencias.py), según la web
oficial de APA (Elements of reference list entries, Missing reference
information, DOIs and URLs, y los ejemplos de Journal article, Book/ebook y
Webpage on a website references). Los ejemplos son los de esas páginas (no los
de normas-apa.org: ver la fase 7 del ROADMAP), y sus versiones en español.

El lector solo lee y nunca adivina: guarda cada parte tal como está escrita, deja
vacío lo que no encuentra y marca como «no reconocida» lo que no puede leer.
En los ejemplos, «*…*» marca lo que va en cursiva.
"""

import hashlib
from pathlib import Path

import docx
import pytest
from lxml import etree

from apa_facil import documento
from apa_facil.__main__ import main
from apa_facil.revision import referencias
from apa_facil.revision.referencias import ARTICULO, CAPITULO, OBRA, Autor

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {ruta.name[:2]: ruta for ruta in sorted((RAIZ / "pruebas").glob("*.docx"))}


def leer(marcado):
    """Lee una referencia escrita con «*…*» alrededor de lo que va en cursiva."""
    texto, cursiva = "", []
    for i, trozo in enumerate(marcado.split("*")):
        texto += trozo
        cursiva += [i % 2 == 1] * len(trozo)
    return referencias.leer_texto(texto, cursiva)


def partes(ref, *nombres):
    return {nombre: getattr(ref, nombre) for nombre in nombres}


# --- Ejemplos de la web oficial de APA ------------------------------------------------

def test_articulo_con_cinco_autores_y_doi():
    ref = leer("Grady, J. S., Her, M., Moreno, G., Perez, C., & Yelinek, J. (2019). Emotions in storybooks: A "
               "comparison of storybooks that represent ethnic and racial groups in the United States. *Psychology "
               "of Popular Media Culture, 8*(3), 207–217. https://doi.org/10.1037/ppm0000185")
    assert ref.autores == [Autor("Grady", "J. S."), Autor("Her", "M."), Autor("Moreno", "G."),
                           Autor("Perez", "C."), Autor("Yelinek", "J.")]
    assert partes(ref, "tipo", "conector", "año", "titulo", "titulo_en_cursiva", "fuente", "fuente_en_cursiva",
                  "volumen", "numero_de_revista", "paginas", "doi", "url", "confianza") == {
        "tipo": ARTICULO, "conector": "&", "año": "2019",
        "titulo": "Emotions in storybooks: A comparison of storybooks that represent ethnic and racial groups in "
                  "the United States", "titulo_en_cursiva": False,
        "fuente": "Psychology of Popular Media Culture", "fuente_en_cursiva": True,
        "volumen": "8", "numero_de_revista": "3", "paginas": "207–217",
        "doi": "https://doi.org/10.1037/ppm0000185", "url": None, "confianza": "alta"}


def test_articulo_con_numero_de_articulo():
    ref = leer("Jerrentrup, A., Mueller, T., Glowalla, U., Herder, M., Henrichs, N., Neubauer, A., & Schaefer, "
               "J. R. (2018). Teaching medicine with the help of “Dr. House.” *PLoS ONE, 13*(3), Article e0193972. "
               "https://doi.org/10.1371/journal.pone.0193972")
    assert len(ref.autores) == 7 and ref.autores[-1] == Autor("Schaefer", "J. R.")
    assert partes(ref, "titulo", "fuente", "volumen", "numero_de_revista", "paginas", "articulo") == {
        "titulo": "Teaching medicine with the help of “Dr. House.”", "fuente": "PLoS ONE", "volumen": "13",
        "numero_de_revista": "3", "paginas": None, "articulo": "Article e0193972"}


def test_articulo_sin_volumen_y_con_estacion():
    ref = leer("Lipscomb, A. Y. (2021, Winter). Addressing trauma in the college essay writing process. *The "
               "Journal of College Admission*, (249), 30–33. https://www.catholiccollegesonline.org/pdf/x.pdf")
    assert partes(ref, "fecha_escrita", "año", "fuente", "volumen", "numero_de_revista", "paginas", "url") == {
        "fecha_escrita": "2021, Winter", "año": "2021", "fuente": "The Journal of College Admission",
        "volumen": None, "numero_de_revista": "249", "paginas": "30–33",
        "url": "https://www.catholiccollegesonline.org/pdf/x.pdf"}


def test_libro_con_edicion_y_doi():
    ref = leer("Jackson, L. M. (2019). *The psychology of prejudice: From attitudes to social action* (2nd ed.). "
               "American Psychological Association. https://doi.org/10.1037/0000168-000")
    assert partes(ref, "tipo", "autores", "titulo", "titulo_en_cursiva", "edicion", "fuente", "doi") == {
        "tipo": OBRA, "autores": [Autor("Jackson", "L. M.")],
        "titulo": "The psychology of prejudice: From attitudes to social action", "titulo_en_cursiva": True,
        "edicion": "2nd ed.", "fuente": "American Psychological Association",
        "doi": "https://doi.org/10.1037/0000168-000"}


def test_libro():
    ref = leer("Sapolsky, R. M. (2017). *Behave: The biology of humans at our best and worst*. Penguin Books.")
    assert partes(ref, "titulo", "fuente", "lugar", "doi", "url", "confianza") == {
        "titulo": "Behave: The biology of humans at our best and worst", "fuente": "Penguin Books",
        "lugar": None, "doi": None, "url": None, "confianza": "alta"}


def test_libro_editado():
    ref = leer("Hygum, E., & Pedersen, P. M. (Eds.). (2010). *Early childhood education: Values and practices in "
               "Denmark*. Hans Reitzels Forlag. https://earlychildhoodeducation.digi.hansreitzel.dk/")
    assert partes(ref, "autores", "autores_escritos", "rol", "año", "titulo", "fuente") == {
        "autores": [Autor("Hygum", "E."), Autor("Pedersen", "P. M.")],
        "autores_escritos": "Hygum, E., & Pedersen, P. M.", "rol": "Eds.", "año": "2010",
        "titulo": "Early childhood education: Values and practices in Denmark", "fuente": "Hans Reitzels Forlag"}


def test_pagina_web_de_una_organizacion():
    ref = leer("National Institute of Mental Health. (2018, July). *Anxiety disorders*. U.S. Department of Health "
               "and Human Services, National Institutes of Health. "
               "https://www.nimh.nih.gov/health/topics/anxiety-disorders/index.shtml")
    assert partes(ref, "autores", "grupo", "fecha_escrita", "año", "titulo", "fuente", "lugar", "url") == {
        "autores": [Autor("National Institute of Mental Health", None)], "grupo": True,
        "fecha_escrita": "2018, July", "año": "2018", "titulo": "Anxiety disorders",
        "fuente": "U.S. Department of Health and Human Services, National Institutes of Health", "lugar": None,
        "url": "https://www.nimh.nih.gov/health/topics/anxiety-disorders/index.shtml"}


def test_pagina_web_de_un_sitio_de_noticias():
    ref = leer("Bologna, C. (2019, October 31). *Why some people with anxiety love watching horror movies*. "
               "HuffPost. https://www.huffpost.com/entry/anxiety")
    assert partes(ref, "fecha_escrita", "año", "titulo", "fuente", "url") == {
        "fecha_escrita": "2019, October 31", "año": "2019",
        "titulo": "Why some people with anxiety love watching horror movies", "fuente": "HuffPost",
        "url": "https://www.huffpost.com/entry/anxiety"}


def test_capitulo_de_un_libro_editado():
    ref = leer("Pérez, A. (2019). Título del capítulo. In G. R. Samanez-Larkin (Ed.), *The aging brain: Functional "
               "adaptation across adulthood* (pp. 9–43). American Psychological Association. "
               "https://doi.org/10.1037/0000143-002")
    assert partes(ref, "tipo", "titulo", "paginas") == {
        "tipo": CAPITULO, "titulo": "Título del capítulo", "paginas": "pp. 9–43"}
    assert ref.fuente.startswith("In G. R. Samanez-Larkin (Ed.), The aging brain")


@pytest.mark.parametrize("marcado", [
    "American Psychological Association. (2012). *APA handbook of industrial and organizational psychology: Vol. 1. "
    "Building and developing the organization*. American Psychological Association.",
    "American Psychological Association. (2012). APA handbook of industrial and organizational psychology: Vol. 1. "
    "Building and developing the organization. American Psychological Association.",
])
def test_volumen_con_titulo_propio(marcado):
    """APA (Elements, Title): el volumen con título propio va dentro del título; sin cursiva, «Vol. 1.» no lo corta."""
    ref = leer(marcado)
    assert ref.titulo == ("APA handbook of industrial and organizational psychology: Vol. 1. Building and "
                          "developing the organization")
    assert ref.fuente == "American Psychological Association"


def test_organizacion_editora():
    ref = leer("Organización Mundial de la Salud (Ed.). (2020). *Informe mundial*. Organización Mundial de la Salud.")
    assert partes(ref, "autores_escritos", "grupo", "rol", "titulo") == {
        "autores_escritos": "Organización Mundial de la Salud", "grupo": True, "rol": "Ed.",
        "titulo": "Informe mundial"}


def test_titulo_que_acaba_en_interrogacion():
    ref = leer("Gómez, A. (2020). Late-onset unexplained epilepsy: What are we missing? Revista de Neurología, "
               "3(1), 1–9.")
    assert partes(ref, "titulo", "fuente", "volumen") == {
        "titulo": "Late-onset unexplained epilepsy: What are we missing?", "fuente": "Revista de Neurología",
        "volumen": "3"}


# --- En español -----------------------------------------------------------------------

def test_en_español():
    ref = leer("Santos-García, S. y Velasco Rodríguez, M. L. (s. f.). *Título de la obra* (2.ª ed.). Editorial "
               "Ejemplo.")
    assert partes(ref, "autores", "conector", "año", "titulo", "edicion", "fuente") == {
        "autores": [Autor("Santos-García", "S."), Autor("Velasco Rodríguez", "M. L.")], "conector": "y",
        "año": "s. f.", "titulo": "Título de la obra", "edicion": "2.ª ed.", "fuente": "Editorial Ejemplo"}
    coordinado = leer("Ramírez, P. (Coord.). (2020). *Libro coordinado*. Editorial Ejemplo. Recuperado de "
                      "https://ejemplo.org/libro")
    assert partes(coordinado, "rol", "url", "recuperado_de") == {
        "rol": "Coord.", "url": "https://ejemplo.org/libro", "recuperado_de": "Recuperado de"}
    capitulo = leer("Pérez, A. (2019). Título del capítulo. En B. Gómez (Ed.), *Libro* (pp. 9–43). Editorial.")
    assert (capitulo.tipo, capitulo.titulo, capitulo.paginas) == (CAPITULO, "Título del capítulo", "pp. 9–43")


def test_el_titulo_en_español_con_interrogacion_en_medio():
    """«¿…?» en medio del título: sin cursiva, el título sigue hasta el punto."""
    ref = leer("Collantes, F. y Pinilla, V. (2011). ¿Lugares que no importan? La despoblación de la España rural. "
               "Prensas Universitarias de Zaragoza.")
    assert ref.titulo == "¿Lugares que no importan? La despoblación de la España rural"
    assert ref.fuente == "Prensas Universitarias de Zaragoza"


@pytest.mark.parametrize("fuente, editorial, lugar", [
    ("Paris: Éditions Exemple", "Éditions Exemple", "Paris"),
    ("Editorial Síntesis, Madrid", "Editorial Síntesis", "Madrid"),
    ("Madrid", None, "Madrid"),
    ("Cambridge University Press", "Cambridge University Press", None),
    ("Penguin Books", "Penguin Books", None),
])
def test_lugar_de_edicion(fuente, editorial, lugar):
    ref = leer(f"Pérez, A. (2019). *Título*. {fuente}.")
    assert (ref.fuente, ref.lugar) == (editorial, lugar)


# --- Lo que falta o está fuera de su sitio --------------------------------------------

def test_sin_fecha_entre_parentesis():
    ref = leer("Lefèvre, A. 2009. Industrialisation et travail en Europe. Paris: Éditions Exemple.")
    assert partes(ref, "año", "año_entre_parentesis", "año_tras_el_autor", "titulo", "confianza") == {
        "año": "2009", "año_entre_parentesis": False, "año_tras_el_autor": True,
        "titulo": "Industrialisation et travail en Europe", "confianza": "media"}
    detras = leer("Martínez, J. Redes sociales y adolescencia. 2019. https://www.example.org/martinez-redes")
    assert partes(detras, "año", "año_tras_el_autor", "titulo", "url") == {
        "año": "2019", "año_tras_el_autor": False, "titulo": "Redes sociales y adolescencia",
        "url": "https://www.example.org/martinez-redes"}
    sin_año = leer("Herranz, M. Turismo rural. Revista de Estudios, 41, 77-96. doi:10.0000/rerf.2015.041")
    assert partes(sin_año, "año", "fecha_escrita", "titulo", "fuente", "volumen", "doi", "dudas") == {
        "año": None, "fecha_escrita": None, "titulo": "Turismo rural", "fuente": "Revista de Estudios",
        "volumen": "41", "doi": "doi:10.0000/rerf.2015.041", "dudas": ["no tiene fecha"]}


def test_el_titulo_y_la_revista_separados_por_una_coma():
    ref = leer("Pérez, L.; Romero, A.; Saiz, T.; Moya, P. (2020) Autoestima digital, Revista Ejemplo Joven, nº 7.")
    assert partes(ref, "conector", "titulo", "fuente", "numero_de_revista", "confianza") == {
        "conector": ";", "titulo": "Autoestima digital", "fuente": "Revista Ejemplo Joven",
        "numero_de_revista": "7", "confianza": "media"}
    assert len(ref.autores) == 4


@pytest.mark.parametrize("final", [", 12, nº 7, 1-9.", ", vol. 12, nº 7, 1-9.", ", vol. 12(7), 1-9."])
def test_volumen_y_numero_escritos_con_abreviaturas(final):
    """Formas que no son de APA, pero que el lector debe entender (el 7.4 las señala)."""
    ref = leer(f"Pérez, L. (2020). Autoestima digital. Revista Joven{final}")
    assert partes(ref, "fuente", "volumen", "numero_de_revista", "paginas") == {
        "fuente": "Revista Joven", "volumen": "12", "numero_de_revista": "7", "paginas": "1-9"}


def test_sin_autor_el_titulo_va_delante():
    """APA (Missing reference information): «Título. (Fecha). Fuente.» Por la forma no se distingue de una organización."""
    ref = leer("*Merriam-Webster’s collegiate dictionary*. (2019). Merriam-Webster.")
    assert (ref.grupo, ref.autores_escritos, ref.año) == (True, "Merriam-Webster’s collegiate dictionary", "2019")


@pytest.mark.parametrize("texto, motivo", [
    ("https://www.historiaejemplo.es/revolucion-industrial-trabajo", "Solo tiene una dirección web"),
    ("www.example.org/datos-uso-redes-2024", "Solo tiene una dirección web"),
    ("Redes sociales y adolescencia, varios autores, sin fecha.", "No se distingue el autor del título"),
])
def test_no_reconocidas(texto, motivo):
    ref = leer(texto)
    assert not ref.reconocida and ref.motivo.startswith(motivo)
    assert ref.titulo is None and ref.autores == [] and ref.confianza is None


# --- Los cinco documentos de prueba ---------------------------------------------------

CAMPOS = ("numero", "reconocida", "tipo", "autores_escritos", "año", "titulo", "fuente", "volumen",
          "numero_de_revista", "paginas", "lugar", "doi", "url", "confianza")
ESPERADAS = {
    "01": [(20, True, OBRA, "Ramírez Soto, M.", "2015", "Fábricas y ciudades: el trabajo en la Europa industrial",
            "Editorial Síntesis Ficticia", None, None, None, "Madrid", None, None, "alta"),
           (21, True, ARTICULO, "Ortega, L. y Blanco, R.", "2018", "El tiempo de la fábrica",
            "Revista de Historia Social Inventada", "34", "2", "45-67", None, None, None, "alta"),
           (22, True, OBRA, "Lefèvre, A.", "2009", "Industrialisation et travail en Europe", "Éditions Exemple",
            None, None, None, "Paris", None, None, "media"),
           (23, False, None, None, None, None, None, None, None, None, None, None,
            "https://www.historiaejemplo.es/revolucion-industrial-trabajo", None)],
    "02": [(67, True, ARTICULO, "Castro, P., Llorente, M. y Benítez, J.", "2019",
            "Calidad de sueño en población universitaria: una revisión sistemática", "Revista Ejemplo de Psicología",
            "12", "3", "201-219", None, None, None, "alta"),
           (68, True, OBRA, "Navarro, A. y Gil, S.", "2020", "Sueño, memoria y aprendizaje",
            "Editorial Pirámide Inventada", None, None, None, None, None, None, "alta")],
    "03": [(28, True, OBRA, "Herrero, C.", "2021", "Energía y vivienda en España", "Ediciones Ejemplo",
            None, None, None, None, None, None, "alta"),
           (29, True, ARTICULO, "Campos, F., & Vidal, N.", "2022", "Autoconsumo solar en entornos rurales",
            "Revista Ficticia de Energía", "8", None, "15-29", None, None, None, "alta")],
    "04": [(66, True, OBRA, "Martínez, J.", "2019", "Redes sociales y adolescencia", None, None, None, None,
            None, None, "https://www.example.org/martinez-redes", "media"),
           (67, True, ARTICULO, "Pérez, L.; Romero, A.; Saiz, T.; Moya, P.", "2020", "Autoestima digital",
            "Revista Ejemplo Joven", None, "7", None, None, None, None, "media"),
           (68, False, None, None, None, None, None, None, None, None, None, None,
            "www.example.org/datos-uso-redes-2024", None)],
    "05": [(68, True, OBRA, "Molina, R.", "2017", "Turismo y territorio: límites del desarrollo rural",
            "Editorial Ejemplo", None, None, None, None, None, None, "alta"),
           (69, True, OBRA, "Collantes, F. y Pinilla, V.", "2011",
            "¿Lugares que no importan? La despoblación de la España rural", "Prensas Universitarias de Ejemplo",
            None, None, None, None, None, None, "alta"),
           (70, True, OBRA, "Consejo Económico de Ejemplo", "2022",
            "Informe sobre el medio rural y su vertebración social y territorial", None, None, None, None,
            "Madrid", None, None, "alta"),
           (71, True, ARTICULO, "Herranz, M.", None, "Turismo rural y servicios públicos",
            "Revista de Estudios Regionales Ficticios", "41", None, "77-96", None, "doi:10.0000/rerf.2015.041",
            None, "media"),
           (72, True, ARTICULO, "Sáenz Gil, P.", "2020", "Nuevos pobladores en el medio rural",
            "Cuadernos de Geografía Inventada", "18", "1", "33-51", None, "https://doi.org/10.0000/cgi.2020.18.1",
            None, "alta"),
           (73, True, OBRA, "Torres, E.", "2019", "Agroturismo y empleo femenino", "Editorial Ejemplo",
            None, None, None, None, None, None, "alta")],
}


@pytest.mark.parametrize("clave", sorted(ESPERADAS))
def test_cinco_documentos(clave):
    leidas = referencias.leer(documento.abrir(PRUEBAS[clave]))
    assert [tuple(getattr(r, c) for c in CAMPOS) for r in leidas] == ESPERADAS[clave]


def test_la_cursiva_se_lee_del_documento():
    """En el 03 y el 05 los títulos de los libros y los nombres de las revistas van en cursiva; en el 01, no."""
    tercero = referencias.leer(documento.abrir(PRUEBAS["03"]))
    assert [(r.titulo_en_cursiva, r.fuente_en_cursiva) for r in tercero] == [(True, False), (False, True)]
    primero = referencias.leer(documento.abrir(PRUEBAS["01"]))
    assert [r.titulo_en_cursiva for r in primero if r.reconocida] == [False, False, False]


def test_el_rotulo_de_la_lista():
    assert {r.rotulo for r in referencias.leer(documento.abrir(PRUEBAS["04"]))} == {"Webgrafía:"}
    assert {r.rotulo for r in referencias.leer(documento.abrir(PRUEBAS["05"]))} == {"BIBLIOGRAFÍA"}


def test_varias_listas():
    doc = docx.Document()
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph("El turismo rural es un fenómeno estudiado desde hace décadas por la geografía y la economía, "
                      "y su relación con la despoblación sigue abierta (Molina, 2017).")
    doc.add_heading("Bibliografía", level=1)
    doc.add_paragraph("Molina, R. (2017). Turismo y territorio. Editorial Ejemplo.")
    doc.add_heading("Webgrafía", level=1)
    doc.add_paragraph("Instituto Nacional de Estadística. (2024). Padrón municipal. https://www.ine.es/padron")
    leidas = referencias.leer(doc)
    assert [(r.lista, r.rotulo, r.autores_escritos) for r in leidas] == [
        (1, "Bibliografía", "Molina, R."), (2, "Webgrafía", "Instituto Nacional de Estadística")]


def test_no_cambia_el_documento():
    doc = documento.abrir(PRUEBAS["05"])
    antes = etree.tostring(doc.element)
    huella = hashlib.sha256(PRUEBAS["05"].read_bytes()).hexdigest()
    referencias.leer(doc)
    assert etree.tostring(doc.element) == antes
    assert hashlib.sha256(PRUEBAS["05"].read_bytes()).hexdigest() == huella


def test_un_documento_sin_referencias():
    doc = docx.Document()
    doc.add_paragraph("Un párrafo sin lista de referencias.")
    assert referencias.leer(doc) == []
    assert referencias.como_texto([]) == "No se ha encontrado ninguna lista de referencias."


# --- Opción --referencias -------------------------------------------------------------

def test_opcion_referencias_no_genera_copias(tmp_path, capsys):
    assert main([str(PRUEBAS["05"]), "--referencias"], carpeta_salida=tmp_path) == 0
    salida = capsys.readouterr().out
    assert list(tmp_path.iterdir()) == []
    assert "Referencias de 05_tfg_completo.docx" in salida and "── Lista 1 («BIBLIOGRAFÍA»)" in salida
    assert "DOI: doi:10.0000/rerf.2015.041" in salida and "Fuera de su sitio: no tiene fecha" in salida
    assert "Resumen: 6 referencia(s): 6 reconocida(s) (confianza alta 5, media 1) y 0 no reconocida(s)." in salida
    assert "1 de 1 documento(s) revisado(s). No se ha generado ninguna copia." in salida


def test_opcion_referencias_con_comodin_y_otra_opcion(capsys):
    assert main([str(RAIZ / "pruebas" / "*.docx"), "--referencias", "--fuente", "Arial"]) == 0
    salida = capsys.readouterr().out
    assert "la opción --fuente no se usa en la revisión de referencias" in salida
    assert "No reconocida. Solo tiene una dirección web" in salida
    assert "5 de 5 documento(s) revisado(s)" in salida


def test_opcion_referencias_con_un_archivo_que_no_existe(capsys):
    assert main(["no_existe.docx", "--referencias"]) == 1


def test_la_ayuda_explica_la_opcion_referencias(capsys):
    with pytest.raises(SystemExit):
        main(["--help"])
    ayuda = " ".join(capsys.readouterr().out.split())
    assert "--referencias" in ayuda and "muestra las partes de cada referencia" in ayuda
