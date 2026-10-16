"""
Paso 0.4: abrir un documento y guardar una copia sin perder nada.

Las copias se guardan en una carpeta temporal (tmp_path) para no ensuciar
salida/, y los archivos falsos para los errores también se crean ahí:
la carpeta pruebas/ no se toca nunca.
"""

import hashlib
import zipfile
from pathlib import Path

import docx
import pytest

from apa_facil import documento, pipeline
from apa_facil.__main__ import main
from apa_facil.documento import ErrorDocumento

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))


def huella(ruta):
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


def test_estan_los_cinco_documentos():
    assert len(PRUEBAS) == 5


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_ida_y_vuelta(original, tmp_path):
    antes = huella(original)
    salida, _ = pipeline.procesar(original, tmp_path)

    assert salida == tmp_path / f"{original.stem}_APA.docx"
    copia, orig = docx.Document(str(salida)), docx.Document(str(original))
    assert len(copia.paragraphs) == len(orig.paragraphs)
    assert len(copia.tables) == len(orig.tables)
    assert len(copia.sections) == len(orig.sections)
    assert huella(original) == antes, "el original no debe cambiar"


def test_la_copia_nunca_es_el_original(tmp_path):
    original = tmp_path / "trabajo.docx"
    assert documento.ruta_salida(original, tmp_path) != original.resolve()


def _zip_sin_word(ruta):
    with zipfile.ZipFile(ruta, "w") as z:
        z.writestr("hola.txt", "no soy un Word")


@pytest.mark.parametrize("nombre, preparar, mensaje", [
    ("no_existe.docx", None, "No se encuentra"),
    ("notas.txt", lambda r: r.write_text("hola"), "Solo se admiten archivos .docx"),
    ("antiguo.doc", lambda r: r.write_bytes(b"\xd0\xcf\x11\xe0"), "formato antiguo"),
    ("falso.docx", lambda r: r.write_text("esto es texto, no un Word"), "dañado"),
    ("zip_raro.docx", _zip_sin_word, "dañado"),
    ("~$abierto.docx", lambda r: r.write_bytes(b"x"), "archivo temporal"),
])
def test_errores_explicados(tmp_path, nombre, preparar, mensaje):
    ruta = tmp_path / nombre
    if preparar:
        preparar(ruta)
    with pytest.raises(ErrorDocumento, match=mensaje):
        documento.abrir(ruta)


def test_linea_de_comandos_sigue_tras_un_error(tmp_path, capsys):
    falso = tmp_path / "falso.docx"
    falso.write_text("no soy un Word")
    codigo = main([str(PRUEBAS[0]), str(falso)], carpeta_salida=tmp_path / "salida")

    salida = capsys.readouterr().out
    assert codigo == 1
    assert "1 de 2 documento(s) procesado(s)" in salida
    assert (tmp_path / "salida" / f"{PRUEBAS[0].stem}_APA.docx").exists()


def test_linea_de_comandos_sin_archivos(capsys):
    with pytest.raises(SystemExit) as fin:
        main([])
    assert fin.value.code == 2
    assert "Falta indicar qué documento procesar" in capsys.readouterr().err
