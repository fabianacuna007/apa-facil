"""
Paso 4.1: estilos de título de Word (Título 1 a Título 5) con el formato APA,
cada título con el estilo de su nivel y avisos de estructura.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

import copy
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK

from apa_facil import estructura, pipeline
from apa_facil.documento import Estilos, comparar, radiografia, w
from apa_facil.informe import Informe
from apa_facil.reglas import titulos
from herramientas.verificar import FORMATO_TITULOS, TITULOS, verificar

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
TEXTO = ("Este es un párrafo normal de texto con bastantes palabras para que no parezca un título "
         "ni una línea de portada, y que termina en punto como cualquier frase.")


def documento_con_portada():
    """Un documento cuya primera página es una portada (para que no estorbe en las pruebas)."""
    doc = docx.Document()
    doc.add_paragraph("Portada").runs[0].add_break(WD_BREAK.PAGE)
    return doc


def aplicar(doc):
    informe = Informe()
    titulos.aplicar(doc, {}, informe)
    return informe


def titulo_falso(doc, texto):
    """Un título hecho a mano: corto, en negrita, con el estilo Normal y seguido de texto."""
    p = doc.add_paragraph()
    p.add_run(texto).bold = True
    doc.add_paragraph(TEXTO)
    return p


def estilo_de(doc, p):
    return Estilos(doc).estilo_de(p._p)


def unico(doc, nivel):
    estilos = titulos.estilos_de_titulo(doc)[nivel]
    assert len(estilos) == 1
    return estilos[0]


def propio(estilo, capa, etiqueta, atributo="val"):
    el = estilo.find(f"{w(capa)}/{w(etiqueta)}")
    if el is None:
        return None
    return el.get(w(atributo), "activado")


def problemas_del_verificador(doc, regla=TITULOS):
    return [h.mensaje for h in verificar(doc).incumplimientos if h.regla == regla]


# --- Formato de los estilos ---------------------------------------------------------------

@pytest.mark.parametrize("nivel, alineacion, cursiva, sangria", [
    (1, "center", "0", "0"),
    (2, "left", "0", "0"),
    (3, "left", "activado", "0"),
    (4, "left", "0", "720"),
    (5, "left", "activado", "720"),
])
def test_formato_apa_de_cada_nivel(nivel, alineacion, cursiva, sangria):
    doc = documento_con_portada()
    doc.add_heading("Un título", level=nivel)
    aplicar(doc)
    estilo = unico(doc, nivel)
    assert propio(estilo, "pPr", "jc") == alineacion
    assert propio(estilo, "rPr", "b") == "activado"
    assert propio(estilo, "rPr", "i") == cursiva
    assert propio(estilo, "pPr", "ind", "firstLine") == sangria
    assert propio(estilo, "pPr", "keepNext") == "activado", "no se queda solo al final de la página"
    assert propio(estilo, "pPr", "outlineLvl") == str(nivel - 1), "nivel de esquema: panel de navegación"


def test_se_quitan_color_subrayado_mayusculas_y_bordes():
    doc = documento_con_portada()
    doc.add_heading("Un título", level=1)
    estilo = unico(doc, 1)
    rpr, ppr = estilo.get_or_add_rPr(), estilo.get_or_add_pPr()
    for etiqueta in ("u", "caps", "smallCaps"):
        rpr.append(copy.deepcopy(rpr.makeelement(w(etiqueta), {w("val"): "single"} if etiqueta == "u" else {})))
    ppr.append(ppr.makeelement(w("pBdr"), {}))
    assert estilo.find(f"{w('rPr')}/{w('color')}") is not None   # el azul de Word
    aplicar(doc)
    for etiqueta in ("color", "u", "caps", "smallCaps"):
        assert estilo.find(f"{w('rPr')}/{w(etiqueta)}") is None, etiqueta
    assert estilo.find(f"{w('pPr')}/{w('pBdr')}") is None


def test_el_estilo_de_caracter_vinculado_pierde_el_color():
    doc = documento_con_portada()
    doc.add_heading("Un título", level=1)
    vinculado = doc.styles.element.makeelement(w("style"), {w("type"): "character", w("styleId"): "Heading1Char"})
    rpr = vinculado.makeelement(w("rPr"), {})
    rpr.append(rpr.makeelement(w("color"), {w("val"): "2E74B5"}))
    vinculado.append(rpr)
    doc.styles.element.append(vinculado)
    unico(doc, 1).get_or_add_pPr().addprevious(unico(doc, 1).makeelement(w("link"), {w("val"): "Heading1Char"}))
    aplicar(doc)
    assert vinculado.find(f"{w('rPr')}/{w('color')}") is None


def test_estilos_duplicados_quedan_en_uno():
    """En 02, 03 y 05 hay dos definiciones de «Heading 1»: el resultado dependería de cuál lea Word."""
    doc = documento_con_portada()
    p = doc.add_heading("Un título", level=1)
    doc.styles.element.append(copy.deepcopy(unico(doc, 1)))
    assert len(titulos.estilos_de_titulo(doc)[1]) == 2
    assert any("definido varias veces" in m for m in problemas_del_verificador(doc))
    aplicar(doc)
    unico(doc, 1)
    assert estilo_de(doc, p) == "Heading1"
    assert problemas_del_verificador(doc) == []


def test_estilo_que_falta_se_crea_basado_en_normal():
    doc = documento_con_portada()
    for estilo in titulos.estilos_de_titulo(doc)[1]:
        estilo.getparent().remove(estilo)
    p = titulo_falso(doc, "Introducción")
    informe = aplicar(doc)
    estilo = unico(doc, 1)
    assert estilo.get(w("styleId")) == "Heading1"
    assert propio(estilo, "basedOn", "basedOn") is None   # basedOn es hijo directo del estilo
    assert estilo.find(w("basedOn")).get(w("val")) == "Normal"
    assert estilo_de(doc, p) == "Heading1"
    assert propio(estilo, "pPr", "jc") == "center"
    assert any("1 título («Introducción») usa ahora el estilo de título" in c for c in informe.cambios)


def test_estilo_con_identificador_en_espanol():
    """Word en español guarda «Título 1» con el identificador «Ttulo1» y el nombre interno «heading 1»."""
    doc = documento_con_portada()
    estilo = unico(doc, 1)
    estilo.set(w("styleId"), "Ttulo1")
    estilo.find(w("name")).set(w("val"), "heading 1")
    p = titulo_falso(doc, "Introducción")
    aplicar(doc)
    assert unico(doc, 1).get(w("styleId")) == "Ttulo1", "no se crea otro «Título 1»"
    assert estilo_de(doc, p) == "Ttulo1"


def test_numeracion_automatica_del_estilo_solo_se_avisa():
    doc = documento_con_portada()
    doc.add_heading("Un título", level=1)
    ppr = unico(doc, 1).get_or_add_pPr()
    ppr.get_or_add_numPr().get_or_add_numId().set(w("val"), "1")
    informe = aplicar(doc)
    assert ppr.find(w("numPr")) is not None
    assert any("numera automáticamente" in a for a in informe.avisos)


# --- Asignación ---------------------------------------------------------------------------

def test_titulos_falsos_reciben_el_estilo_de_su_nivel():
    doc = documento_con_portada()
    p1 = titulo_falso(doc, "1. Marco teórico")
    p2 = titulo_falso(doc, "1.1. Antecedentes")
    p3 = titulo_falso(doc, "1.2. Estado de la cuestión")
    aplicar(doc)
    assert [estilo_de(doc, p) for p in (p1, p2, p3)] == ["Heading1", "Heading2", "Heading2"]
    assert problemas_del_verificador(doc) == []


def test_los_que_ya_tenian_su_estilo_no_cuentan_como_cambio():
    doc = documento_con_portada()
    doc.add_heading("Método", level=1)
    doc.add_paragraph(TEXTO)
    informe = aplicar(doc)
    assert len(informe.cambios) == 1 and informe.cambios[0].startswith("Títulos: los estilos")


def test_titulo_dudoso_no_recibe_estilo_y_se_avisa():
    doc = documento_con_portada()
    p = titulo_falso(doc, "Esto es importante.")
    c = next(c for c in estructura.clasificar(doc) if c.elemento is p._p)
    assert (c.tipo, c.confianza) == ("titulo", estructura.BAJA)
    informe = aplicar(doc)
    assert estilo_de(doc, p) == "Normal"
    assert informe.avisos == ["«Esto es importante.» parece un título, pero no es seguro, así que se ha dejado "
                              "como estaba. Si es un título, ponle en Word el estilo «Título 1»."]
    resultado = verificar(doc)
    assert [h.regla for h in resultado.excepciones] == [TITULOS]
    assert not [h for h in resultado.incumplimientos if h.regla == TITULOS]


def test_sin_titulos_no_se_informa_de_nada():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    informe = aplicar(doc)
    assert informe.cambios == informe.avisos == []


def test_el_texto_no_cambia():
    doc = documento_con_portada()
    titulo_falso(doc, "INTRODUCCIÓN")
    titulo_falso(doc, "1.1. Antecedentes")
    doc.add_heading("Método", level=1)
    antes = radiografia(doc)
    aplicar(doc)
    assert comparar(antes, radiografia(doc)) == []


def test_el_verificador_ve_un_titulo_sin_estilo():
    doc = documento_con_portada()
    titulo_falso(doc, "Introducción")
    assert "Título de nivel 1 sin el estilo «Título 1» de Word." in problemas_del_verificador(doc)


def test_el_verificador_ve_un_estilo_estropeado():
    doc = documento_con_portada()
    doc.add_heading("Método", level=1)
    doc.add_heading("Participantes", level=2)
    aplicar(doc)
    assert problemas_del_verificador(doc) == []
    unico(doc, 2).find(f"{w('rPr')}/{w('b')}").set(w("val"), "0")
    unico(doc, 1).find(f"{w('pPr')}/{w('keepNext')}").getparent().remove(unico(doc, 1).find(f"{w('pPr')}/{w('keepNext')}"))
    assert problemas_del_verificador(doc) == [
        "El estilo «Título 1» no tiene el formato APA: no tiene «Conservar con el siguiente».",
        "El estilo «Título 2» no tiene el formato APA: no va en negrita.",
    ]


# --- 4.2 Formato puesto a mano ----------------------------------------------------------------

def titulo(doc, texto, nivel):
    p = doc.add_heading(texto, level=nivel)
    doc.add_paragraph(TEXTO)
    return p


def poner(el, etiqueta, **atributos):
    """Añade a mano un ajuste de formato (w:jc, w:i…) a unas propiedades."""
    hijo = el.makeelement(w(etiqueta), {w(k): v for k, v in atributos.items()})
    el.append(hijo)
    return hijo


def formato_corregido(informe):
    return [c for c in informe.cambios if "formato puesto a mano" in c]


def test_alineacion_a_mano_que_contradice_el_nivel():
    doc = documento_con_portada()
    p = titulo(doc, "Participantes", 2)
    poner(p._p.get_or_add_pPr(), "jc", val="center")
    informe = aplicar(doc)
    assert p._p.pPr.find(w("jc")) is None
    assert formato_corregido(informe) == ["Títulos: se ha quitado el formato puesto a mano que no corresponde a "
                                          "su nivel en 1 título («Participantes»: la alineación)."]
    assert problemas_del_verificador(doc, FORMATO_TITULOS) == []


def test_lo_que_coincide_con_el_nivel_no_se_cuenta():
    doc = documento_con_portada()
    p1 = titulo(doc, "Método", 1)
    poner(p1._p.get_or_add_pPr(), "jc", val="center")
    p4 = titulo(doc, "Edad.", 4)
    p4._p.get_or_add_pPr().get_or_add_ind().set(w("firstLine"), "720")
    for r in p1.runs:
        r.bold = True
    informe = aplicar(doc)
    assert p1._p.pPr.find(w("jc")) is not None and p4._p.pPr.find(w("ind")) is not None
    assert formato_corregido(informe) == []


def test_sangria_a_mano_en_un_titulo():
    doc = documento_con_portada()
    p1 = titulo(doc, "Método", 1)
    titulo(doc, "Muestra", 2)
    titulo(doc, "Edad", 3)
    p4 = titulo(doc, "Niños.", 4)
    aplicar(doc)   # los estilos, ya con el formato APA
    p1._p.get_or_add_pPr().get_or_add_ind().set(w("left"), "720")
    p4._p.get_or_add_pPr().get_or_add_ind().set(w("firstLine"), "0")
    assert problemas_del_verificador(doc, FORMATO_TITULOS) == [
        "Título de nivel 1: la sangría no es la de su nivel.", "Título de nivel 4: la sangría no es la de su nivel."]
    informe = aplicar(doc)
    assert p1._p.pPr.find(w("ind")) is None and p4._p.pPr.find(w("ind")) is None
    assert "«Método»: la sangría; «Niños.»: la sangría" in formato_corregido(informe)[0]
    assert problemas_del_verificador(doc, FORMATO_TITULOS) == []


def test_cursiva_en_todo_el_titulo_se_quita():
    doc = documento_con_portada()
    p = titulo(doc, "Método", 1)
    p.runs[0].italic = True
    informe = aplicar(doc)
    assert p.runs[0].italic is None
    assert "«Método»: la cursiva)" in formato_corregido(informe)[0]


def test_cursiva_desactivada_en_un_nivel_3():
    doc = documento_con_portada()
    titulo(doc, "Método", 1)
    titulo(doc, "Muestra", 2)
    p = titulo(doc, "Edad", 3)
    p.runs[0].italic = False
    informe = aplicar(doc)
    assert p.runs[0].italic is None
    assert "«Edad»: la cursiva desactivada" in formato_corregido(informe)[0]


def test_cursiva_en_una_parte_del_titulo_se_respeta():
    """APA usa la cursiva dentro de un título para el título de un libro o el nombre de una especie."""
    doc = documento_con_portada()
    p = doc.add_heading("Lecturas de ", level=1)
    p.add_run("Cien años de soledad").italic = True
    doc.add_paragraph(TEXTO)
    informe = aplicar(doc)
    assert p.runs[1].italic is True
    assert formato_corregido(informe) == []
    assert problemas_del_verificador(doc, FORMATO_TITULOS) == []


def test_negrita_desactivada_mayusculas_bordes_y_conservar():
    doc = documento_con_portada()
    p = titulo(doc, "Método", 1)
    rpr = p.runs[0]._r.get_or_add_rPr()
    poner(rpr, "b", val="0")
    poner(rpr, "smallCaps")
    ppr = p._p.get_or_add_pPr()
    poner(ppr, "keepNext", val="0")
    poner(ppr, "pBdr")
    antes = radiografia(doc)
    assert problemas_del_verificador(doc, FORMATO_TITULOS)
    informe = aplicar(doc)
    for etiqueta in ("b", "smallCaps"):
        assert rpr.find(w(etiqueta)) is None, etiqueta
    for etiqueta in ("keepNext", "pBdr"):
        assert ppr.find(w(etiqueta)) is None, etiqueta
    assert formato_corregido(informe) == [
        "Títulos: se ha quitado el formato puesto a mano que no corresponde a su nivel en 1 título («Método»: los "
        "bordes o el sombreado, «Conservar con el siguiente» desactivado, la negrita desactivada y las "
        "mayúsculas forzadas)."]
    assert comparar(antes, radiografia(doc)) == [], "las mayúsculas son formato: el texto no cambia"
    assert problemas_del_verificador(doc, FORMATO_TITULOS) == []


def test_el_verificador_ve_el_formato_a_mano():
    doc = documento_con_portada()
    p = titulo(doc, "Método", 1)
    aplicar(doc)
    p.runs[0].italic = True
    p.runs[0].underline = True
    assert problemas_del_verificador(doc, FORMATO_TITULOS) == ["Título de nivel 1: va en cursiva; va subrayado."]


def test_numeracion_de_word_en_un_titulo_solo_se_avisa():
    doc = documento_con_portada()
    p = titulo(doc, "Método", 1)
    p._p.get_or_add_pPr().get_or_add_numPr().get_or_add_numId().set(w("val"), "1")
    informe = aplicar(doc)
    assert p._p.pPr.find(w("numPr")) is not None
    assert [a for a in informe.avisos if "numeración automática" in a] == [
        "1 título («Método») lleva la numeración automática de Word. APA no numera los títulos, pero se ha "
        "dejado así: quítala en Word si no la exige tu universidad."]


def test_aviso_de_nivel_4_y_5_sin_punto_final():
    """Decisión del propietario: el punto se añade en el paso 4.5; de momento, solo se avisa."""
    doc = documento_con_portada()
    for texto, nivel in (("Método", 1), ("Muestra", 2), ("Edad", 3), ("Niños", 4), ("Niñas.", 4),
                         ("Menores de 6", 5), ("Mayores de 6.", 5)):
        titulo(doc, texto, nivel)
    antes = radiografia(doc)
    avisos = [a for a in aplicar(doc).avisos if "punto" in a]
    assert avisos == ["2 títulos de nivel 4 o 5 («Niños», «Menores de 6») no terminan en punto. En APA, estos "
                      "títulos terminan en punto y el texto sigue en la misma línea; de momento se han dejado "
                      "como estaban."]
    assert comparar(antes, radiografia(doc)) == []


# --- Avisos de estructura ---------------------------------------------------------------------

def test_aviso_de_un_solo_subapartado():
    doc = documento_con_portada()
    for texto, nivel in (("Método", 1), ("Participantes", 2), ("Resultados", 1), ("Datos", 2), ("Análisis", 2)):
        doc.add_heading(texto, level=nivel)
        doc.add_paragraph(TEXTO)
    informe = aplicar(doc)
    assert informe.avisos == ["El apartado «Método» tiene un solo subapartado («Participantes»). APA recomienda "
                              "dividir un apartado solo si tiene al menos dos subapartados: añade otro o integra "
                              "su contenido en el apartado."]


def test_los_nietos_no_cuentan_como_subapartados():
    doc = documento_con_portada()
    for texto, nivel in (("Método", 1), ("Muestra", 2), ("Edad", 3), ("Sexo", 3), ("Instrumentos", 2)):
        doc.add_heading(texto, level=nivel)
        doc.add_paragraph(TEXTO)
    assert aplicar(doc).avisos == []


def test_aviso_de_salto_de_nivel():
    doc = documento_con_portada()
    for texto, nivel in (("Método", 1), ("Edad", 3), ("Sexo", 3)):
        doc.add_heading(texto, level=nivel)
        doc.add_paragraph(TEXTO)
    avisos = aplicar(doc).avisos
    assert avisos == ["El título «Edad» es de nivel 3, pero el anterior («Método») es de nivel 1: se salta un "
                      "nivel. APA pide no saltarse niveles: revisa si debería ser de nivel 2."]


def test_aviso_si_el_primer_titulo_no_es_de_nivel_1():
    doc = documento_con_portada()
    for texto in ("Muestra", "Instrumentos"):
        doc.add_heading(texto, level=2)
        doc.add_paragraph(TEXTO)
    assert aplicar(doc).avisos == ["El primer título del trabajo, «Muestra», es de nivel 2. APA pide empezar por "
                                   "el nivel 1: revisa su nivel."]


# --- Documentos de prueba ---------------------------------------------------------------------

UN_SOLO_SUBAPARTADO = {   # apartado → su único subapartado
    "01": [],
    "02": [],
    "03": [("2. Datos y resultados", "2.1. Fuentes de energía")],
    "04": [("Efectos sobre la autoestima", "Factores de riesgo:")],
    "05": [("1. Introducción", "1.1. Objetivos")],
}
SIN_ESTILO_EN_EL_ORIGINAL = {"01": 5, "02": 0, "03": 0, "04": 4, "05": 1}   # títulos falsos (05: «RESUMEN»)


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    clave = original.name[:2]
    sin_estilo = [m for m in problemas_del_verificador(docx.Document(str(original))) if "sin el estilo" in m]
    assert len(sin_estilo) == SIN_ESTILO_EN_EL_ORIGINAL[clave]

    salida, informe = pipeline.procesar(original, tmp_path)
    copia = docx.Document(str(salida))
    assert problemas_del_verificador(copia) == problemas_del_verificador(copia, FORMATO_TITULOS) == []
    assert formato_corregido(informe) == [], "en los 5 documentos no queda formato a mano que contradiga"
    sin_punto = [a for a in informe.avisos if "no terminan en punto" in a or "no termina en punto" in a]
    assert len(sin_punto) == (1 if clave == "05" else 0)   # «a) Evolución demográfica» y «b) Percepción…»
    for nivel in titulos.NIVELES:
        ids = [e.get(w("styleId")) for e in titulos.estilos_de_titulo(copia)[nivel]]
        assert len(ids) == len(set(ids)), f"«Título {nivel}» sin duplicados"

    estilos = Estilos(copia)
    por_id = {e.get(w("styleId")): e for e in copia.styles.element.iter(w("style"))}
    for c in estructura.clasificar(copia):
        if c.tipo == "titulo":
            nombre = por_id[estilos.estilo_de(c.elemento)].find(w("name")).get(w("val")).lower()
            assert nombre == f"heading {c.nivel}", c.texto

    esperados = [f"El apartado «{a}» tiene un solo subapartado («{b}»)" for a, b in UN_SOLO_SUBAPARTADO[clave]]
    de_estructura = [a for a in informe.avisos if "subapartado" in a or "se salta un nivel" in a]
    assert [a.split(". APA")[0] for a in de_estructura] == esperados
