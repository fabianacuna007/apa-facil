"""
Paso 1.6: número de página arriba a la derecha en todas las páginas, y nada
más en el encabezado.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

import base64
import io
from pathlib import Path

import docx
import pytest
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm

from apa_facil import pipeline
from apa_facil.documento import Estilos, texto_legible, w
from apa_facil.informe import Informe
from apa_facil.reglas import paginacion

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def aplicar(doc):
    informe = Informe()
    paginacion.aplicar(doc, {}, informe)
    return informe


def campo(p, codigo, antes="", despues=""):
    """Añade a un párrafo un campo de Word (p. ej. PAGE) con texto antes y después."""
    if antes:
        p.add_run(antes)
    for tipo, texto in (("begin", None), ("instr", f" {codigo} "), ("separate", None), ("t", "3"), ("end", None)):
        run = OxmlElement("w:r")
        if tipo == "instr":
            hijo = OxmlElement("w:instrText")
            hijo.text = texto
        elif tipo == "t":
            hijo = OxmlElement("w:t")
            hijo.text = texto
        else:
            hijo = OxmlElement("w:fldChar")
            hijo.set(w("fldCharType"), tipo)
        run.append(hijo)
        p._p.append(run)
    if despues:
        p.add_run(despues)
    return p


def comprobar_encabezado_apa(doc):
    """Un solo encabezado (el de la primera sección), con un párrafo a la derecha que solo lleva PAGE."""
    primera = doc.sections[0]
    assert not primera.header.is_linked_to_previous
    for seccion in doc.sections:
        assert not seccion.different_first_page_header_footer
    for seccion in doc.sections[1:]:
        assert seccion.header.is_linked_to_previous
    parrafos = primera.header.paragraphs
    assert len(parrafos) == 1 and not primera.header.tables
    p = parrafos[0]
    assert p.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert paginacion.tiene_numero_de_pagina(p._p)
    assert texto_legible(p._p) == "[nº de página]"
    assert not doc.settings.odd_and_even_pages_header_footer


# --- El encabezado -------------------------------------------------------------------

def test_documento_sin_encabezado():
    doc = docx.Document()
    doc.add_paragraph("Texto")
    informe = aplicar(doc)
    comprobar_encabezado_apa(doc)
    assert informe.cambios == ["Número de página: arriba a la derecha en todas las páginas, también en la portada."]
    assert informe.avisos == []


def test_varias_secciones_y_horizontal():
    doc = docx.Document()
    doc.add_paragraph("Uno")
    horizontal = doc.add_section(WD_SECTION.NEW_PAGE)
    horizontal.orientation = WD_ORIENT.LANDSCAPE
    horizontal.page_width, horizontal.page_height = Cm(29.7), Cm(21)
    horizontal.header.is_linked_to_previous = False
    horizontal.header.paragraphs[0].text = "Otro encabezado"
    doc.add_section(WD_SECTION.NEW_PAGE)
    aplicar(doc)
    comprobar_encabezado_apa(doc)
    assert doc.sections[1].orientation == WD_ORIENT.LANDSCAPE


def test_texto_del_encabezado_se_quita_y_se_avisa():
    doc = docx.Document()
    doc.sections[0].header.paragraphs[0].text = "Facultad de Ciencias"
    campo(doc.sections[0].header.add_paragraph(), "PAGE")
    informe = aplicar(doc)
    comprobar_encabezado_apa(doc)
    assert informe.avisos == ["El encabezado contenía «Facultad de Ciencias». Se ha eliminado porque APA de "
                              "estudiante no lleva encabezado de página (solo el número)."]


def test_imagen_y_tabla_del_encabezado():
    doc = docx.Document()
    encabezado = doc.sections[0].header
    encabezado.paragraphs[0].add_run().add_picture(io.BytesIO(PNG), width=Cm(2))
    encabezado.add_table(rows=1, cols=2, width=Cm(10))
    informe = aplicar(doc)
    comprobar_encabezado_apa(doc)
    assert any("tenía una imagen" in a for a in informe.avisos)
    assert any("tenía una tabla" in a for a in informe.avisos)


def test_primera_pagina_diferente_y_pares():
    doc = docx.Document()
    seccion = doc.sections[0]
    seccion.different_first_page_header_footer = True
    seccion.first_page_header.paragraphs[0].text = "Portada sin número"
    doc.settings.odd_and_even_pages_header_footer = True
    seccion.even_page_header.paragraphs[0].text = "Páginas pares"
    informe = aplicar(doc)
    comprobar_encabezado_apa(doc)
    assert seccion.first_page_header.is_linked_to_previous
    assert "«Portada sin número», «Páginas pares»" in informe.avisos[0]


def test_numeracion_continua():
    doc = docx.Document()
    doc.add_paragraph("Índice")
    doc.add_section(WD_SECTION.NEW_PAGE)
    tipo = OxmlElement("w:pgNumType")
    tipo.set(w("fmt"), "lowerRoman")
    doc.sections[0]._sectPr.append(tipo)
    reinicio = OxmlElement("w:pgNumType")
    reinicio.set(w("start"), "1")
    doc.sections[1]._sectPr.append(reinicio)
    informe = aplicar(doc)
    assert tipo.get(w("fmt")) is None and reinicio.get(w("start")) is None
    assert any("ya no se reinicia" in c for c in informe.cambios)
    assert any("romanos" in c for c in informe.cambios)


# --- Los pies ----------------------------------------------------------------------------

@pytest.mark.parametrize("antes, despues", [("", ""), ("- ", " -"), ("Página ", ""), ("Pág. ", "")])
def test_pie_con_solo_el_numero_se_vacia(antes, despues):
    doc = docx.Document()
    campo(doc.sections[0].footer.paragraphs[0], "PAGE", antes, despues)
    informe = aplicar(doc)
    pie = doc.sections[0].footer
    assert [texto_legible(p._p) for p in pie.paragraphs] == [""]
    assert any(c.startswith("Se ha quitado el número de página del pie") for c in informe.cambios)


def test_pie_pagina_x_de_y():
    doc = docx.Document()
    p = campo(doc.sections[0].footer.paragraphs[0], "PAGE", "Página ")
    campo(p, "NUMPAGES", " de ")
    aplicar(doc)
    assert [texto_legible(p._p) for p in doc.sections[0].footer.paragraphs] == [""]


def test_pie_con_otro_texto_se_conserva():
    doc = docx.Document()
    pie = doc.sections[0].footer
    campo(pie.paragraphs[0], "PAGE", "Universidad de Ejemplo · ")
    pie.add_paragraph("Curso 2025-2026")
    informe = aplicar(doc)
    textos = [texto_legible(p._p) for p in pie.paragraphs]
    assert textos == ["Universidad de Ejemplo · ", "Curso 2025-2026"]
    assert not any(paginacion.tiene_numero_de_pagina(p._p) for p in pie.paragraphs)
    assert [a for a in informe.avisos if "pie de página contiene" in a] == [
        "El pie de página contiene «Universidad de Ejemplo ·». Se ha dejado porque no es un número de página, "
        "pero APA de estudiante no usa pie de página: revisa si lo necesitas.",
        "El pie de página contiene «Curso 2025-2026». Se ha dejado porque no es un número de página, "
        "pero APA de estudiante no usa pie de página: revisa si lo necesitas."]


# --- Documentos de prueba ------------------------------------------------------------------

AVISOS_DE_ENCABEZADO = {"03": "Facultad de Ciencias · Medio Ambiente y Sostenibilidad",
                        "05": "Turismo rural y despoblación — L. Fernández Prieto"}
PIES_QUITADOS = {"02": "[nº de página]", "03": "- [nº de página] -",
                 "05": "Página [nº de página] de [total de páginas]"}


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    clave = original.name[:2]
    salida, informe = pipeline.procesar(original, tmp_path)
    copia = docx.Document(str(salida))
    comprobar_encabezado_apa(copia)

    for seccion in copia.sections:
        if not seccion.footer.is_linked_to_previous:
            assert not any(paginacion.tiene_numero_de_pagina(p._p) for p in seccion.footer.paragraphs)
            assert all(texto_legible(p._p).strip() == "" for p in seccion.footer.paragraphs)

    avisos = [a for a in informe.avisos if "encabezado contenía" in a]
    if clave in AVISOS_DE_ENCABEZADO:
        assert avisos == [f"El encabezado contenía «{AVISOS_DE_ENCABEZADO[clave]}». Se ha eliminado porque APA "
                          "de estudiante no lleva encabezado de página (solo el número)."]
    else:
        assert avisos == []
    quitados = [c for c in informe.cambios if c.startswith("Se ha quitado el número de página del pie")]
    assert quitados == ([f"Se ha quitado el número de página del pie («{PIES_QUITADOS[clave]}»)."]
                        if clave in PIES_QUITADOS else [])

    # El número usa la fuente y el tamaño del documento, a interlineado sencillo y sin sangría.
    estilos = Estilos(copia)
    p = copia.sections[0].header.paragraphs[0]._p
    run = p.find(w("r"))
    cuerpo = next(q for q in copia.paragraphs if q.text.strip())
    assert estilos.fuente(run, p) == estilos.fuente(cuerpo.runs[0]._r, cuerpo._p)
    assert estilos.parrafo(p, "spacing", "line") == "240"
    assert estilos.parrafo(p, "ind", "firstLine") == "0"
