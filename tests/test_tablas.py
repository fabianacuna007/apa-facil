"""
Paso 6.1: las tablas, según la web oficial de APA (Table setup, Sample tables y
Paragraph alignment and indentation): «Tabla N» en negrita encima y el título en
cursiva en la línea siguiente, a la izquierda y sin sangría; la nota debajo, con
«Nota.» en cursiva; la fila de encabezados repetida en cada página; y, por defecto,
solo líneas horizontales y sin colores de fondo (--tablas-tal-cual los deja).

Lo que cambia el texto (subir la leyenda, separar el número del título, sus
mayúsculas, añadir «Tabla N» y «Nota.») lo comprueba la integridad: cada caso pasa
por pipeline.procesar, que no entrega la copia si el contenido cambia de otra forma.
"""

from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from apa_facil import pipeline
from apa_facil.documento import CambioPrevisto, Estilos, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import tablas as regla
from apa_facil.reglas.fuente import _runs_propios
from herramientas.verificar import TABLAS, verificar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {p.name[:2]: p for p in (RAIZ / "pruebas").glob("*.docx")}
TEXTO = ("Los estudiantes que duermen menos de seis horas obtienen, de media, peores notas que los que duermen "
         "más de ocho, según los datos recogidos durante el curso.")


def trabajo(antes=(), despues=(), filas=3, celdas=None):
    """Un apartado con una tabla de datos y, alrededor, los párrafos `antes` y `despues` (sus textos)."""
    doc = docx.Document()
    doc.add_heading("Resultados", level=1)
    doc.add_paragraph(TEXTO)
    for texto in antes:
        doc.add_paragraph(texto)
    tabla = doc.add_table(rows=filas, cols=3)
    for i, fila in enumerate(tabla.rows):
        for j, celda in enumerate(fila.cells):
            celda.text = (celdas or {}).get((i, j), f"{'Grupo' if i == 0 else 'Dato'} {i}.{j}")
    for texto in despues:
        doc.add_paragraph(texto)
    doc.add_heading("Discusión", level=1)
    doc.add_paragraph(TEXTO)
    return doc


def procesar(doc, tmp_path, opciones=None, nombre="trabajo.docx"):
    """Guarda el documento y lo pasa por el programa entero (con la comprobación de integridad)."""
    ruta = tmp_path / nombre
    doc.save(str(ruta))
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", opciones)
    return docx.Document(str(salida)), informe


def alrededor(doc, k=0):
    """Los textos de los párrafos del cuerpo justo encima y justo debajo de la tabla k (sin los vacíos)."""
    cuerpo = list(doc.element.body.iterchildren())
    tbl = doc.tables[k]._tbl
    i = cuerpo.index(tbl)
    encima = [texto_parrafo(p) for p in cuerpo[:i] if p.tag == w("p") and texto_parrafo(p).strip()]
    debajo = [texto_parrafo(p) for p in cuerpo[i + 1:] if p.tag == w("p") and texto_parrafo(p).strip()]
    return encima[-2:], debajo[:1]


def parrafo(doc, inicio):
    return next(p for p in doc.element.body.iter(w("p")) if texto_parrafo(p).startswith(inicio)
                and next(p.iterancestors(w("tc")), None) is None)


def formato(doc, p):
    """[(texto, cursiva, negrita)] de cada fragmento con texto."""
    estilos = Estilos(doc)
    return [("".join(t.text or "" for t in r.iter(w("t"))), regla._activado(estilos, r, p, "i"),
             regla._activado(estilos, r, p, "b")) for r in _runs_propios(p)
            if "".join(t.text or "" for t in r.iter(w("t")))]


def incumplimientos(doc):
    return [h.mensaje for h in verificar(doc).incumplimientos if h.regla == TABLAS]


# --- Los documentos de prueba --------------------------------------------------------------------

@pytest.fixture(scope="module")
def copias(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("tablas")
    return {clave: pipeline.procesar(PRUEBAS[clave], carpeta) for clave in ("02", "04", "05")}


def test_el_02_sube_y_separa_sus_leyendas(copias):
    salida, informe = copias["02"]
    doc = docx.Document(str(salida))
    assert alrededor(doc, 0)[0] == ["Tabla 1", "Estadísticos Descriptivos de las Variables del Estudio (N = 180)"]
    assert alrededor(doc, 1)[0] == ["Tabla 2", "Rendimiento Académico Según las Horas de Sueño"]
    assert formato(doc, parrafo(doc, "Tabla 1")) == [("Tabla 1", False, True)]
    assert formato(doc, parrafo(doc, "Rendimiento")) == [("Rendimiento Académico Según las Horas de Sueño", True,
                                                          False)]
    assert any(c.startswith("Tablas: las leyendas «Tabla 1: Estadísticos") for c in informe.cambios)
    # La tabla 2 tiene los encabezados en dos filas (celdas combinadas): se repiten las dos.
    filas = doc.tables[1]._tbl.findall(w("tr"))
    assert [f.find(f"{w('trPr')}/{w('tblHeader')}") is not None for f in filas[:3]] == [True, True, False]


def test_el_04_recibe_su_numero(copias):
    salida, informe = copias["04"]
    doc = docx.Document(str(salida))
    assert alrededor(doc)[0] == ["Tabla 1", "Datos de Uso por Red Social"]
    assert any(c.startswith("Tablas: se ha añadido el número a 1 tabla") for c in informe.cambios)


def test_el_05_separa_el_titulo_y_pone_la_nota(copias):
    salida, informe = copias["05"]
    doc = docx.Document(str(salida))
    assert alrededor(doc) == (["Tabla 1", "Evolución de la Población (2001-2024)"],
                              ["Nota. Fuente: elaboración propia a partir del padrón municipal."])
    nota = parrafo(doc, "Nota.")
    assert formato(doc, nota)[0] == ("Nota.", True, False)
    assert all(not cursiva for texto, cursiva, _ in formato(doc, nota)[1:] if texto.strip())


@pytest.mark.parametrize("clave", ["02", "04", "05"])
def test_las_copias_cumplen_y_los_originales_no(copias, clave):
    salida, _ = copias[clave]
    assert incumplimientos(docx.Document(str(salida))) == []
    assert incumplimientos(docx.Document(str(PRUEBAS[clave])))


def test_el_estilo_apa_en_el_02(copias):
    salida, _ = copias["02"]
    tbl = docx.Document(str(salida)).tables[0]._tbl
    bordes = tbl.find(f"{w('tblPr')}/{w('tblBorders')}")
    assert {b.tag.split("}")[1]: b.get(w("val")) for b in bordes} == {
        "top": "single", "left": "nil", "bottom": "single", "right": "nil", "insideH": "nil", "insideV": "nil"}
    assert not [r for r in tbl.iter(w("r")) if r.find(f"{w('rPr')}/{w('color')}") is not None], \
        "el texto blanco de los encabezados pasa a negro"
    assert {s.get(w("fill")) for s in tbl.iter(w("shd"))} == {"auto"}


# --- Casos hechos a mano --------------------------------------------------------------------------

def test_leyenda_debajo_sube_y_se_separa(tmp_path):
    doc, informe = procesar(trabajo(despues=["Tabla 1: notas medias por grupo", TEXTO]), tmp_path)
    assert alrededor(doc) == (["Tabla 1", "Notas Medias por Grupo"], [TEXTO])
    assert incumplimientos(doc) == []


def test_numero_y_titulo_ya_separados_solo_cambian_de_formato(tmp_path):
    original = trabajo(antes=["Tabla 1", "Notas medias por grupo"])
    original.paragraphs[2].alignment = 1   # centrado
    doc, informe = procesar(original, tmp_path)
    assert alrededor(doc)[0] == ["Tabla 1", "Notas Medias por Grupo"]
    assert not any("se ha separado" in c for c in informe.cambios)
    assert incumplimientos(doc) == []


def test_con_titulos_tal_cual_el_titulo_conserva_sus_mayusculas(tmp_path):
    doc, informe = procesar(trabajo(antes=["Tabla 1. Notas medias por grupo"]), tmp_path,
                            {"titulos_tal_cual": True})
    assert alrededor(doc)[0] == ["Tabla 1", "Notas medias por grupo"]


def test_sin_numero_en_ninguna_tabla_se_numeran(tmp_path):
    original = trabajo(antes=["Notas medias por grupo"])
    original.add_paragraph("Horas de sueño por curso")
    original.add_table(rows=2, cols=2)
    doc, informe = procesar(original, tmp_path)
    assert alrededor(doc, 0)[0] == ["Tabla 1", "Notas Medias por Grupo"]
    assert alrededor(doc, 1)[0] == ["Tabla 2", "Horas de Sueño por Curso"]


def test_si_unas_tienen_numero_y_otras_no_solo_se_avisa(tmp_path):
    original = trabajo(antes=["Tabla 1. Notas medias por grupo"])
    original.add_paragraph("Horas de sueño por curso")
    original.add_table(rows=2, cols=2)
    doc, informe = procesar(original, tmp_path)
    assert alrededor(doc, 1)[0][-1] == "Horas de Sueño por Curso"
    assert "Tabla 2" not in [texto_parrafo(p) for p in doc.element.body.iter(w("p"))]
    assert any("Como otras sí lo llevan, no se ha añadido" in a for a in informe.avisos)


def test_nota_que_ya_empieza_por_nota(tmp_path):
    doc, informe = procesar(trabajo(antes=["Tabla 1", "Notas"], despues=["Nota. Datos del curso 2025."]),
                            tmp_path)
    assert alrededor(doc)[1] == ["Nota. Datos del curso 2025."]
    assert formato(doc, parrafo(doc, "Nota."))[0][1:] == (True, False)
    assert not any("«Nota.»" in c and "empieza" in c for c in informe.cambios)


def test_con_tablas_tal_cual_se_dejan_los_bordes(tmp_path):
    original = trabajo(antes=["Tabla 1", "Notas"])
    original.tables[0].style = original.styles["Table Grid"]
    doc, informe = procesar(original, tmp_path, {"tablas_tal_cual": True})
    tbl = doc.tables[0]._tbl
    assert tbl.find(f"{w('tblPr')}/{w('tblBorders')}") is None
    assert tbl.find(f"{w('tr')}/{w('trPr')}/{w('tblHeader')}") is not None, "los encabezados se repiten igual"
    assert incumplimientos(doc) == ["Tabla: tiene bordes verticales."]


def test_estilo_de_tabla_con_texto_de_color(tmp_path):
    original = trabajo(antes=["Tabla 1", "Notas"])
    original.tables[0].style = original.styles["Light Grid Accent 1"]
    doc, _ = procesar(original, tmp_path)
    tbl = doc.tables[0]._tbl
    assert incumplimientos(doc) == []
    colores = {c.get(w("val")) for c in tbl.iter(w("color"))}
    assert colores <= {"auto"}


def test_campo_de_numeracion_de_word(tmp_path):
    """«Insertar título» de Word: «Tabla {SEQ Tabla}: título». El campo se queda con el número."""
    original = trabajo()
    p = original.add_paragraph()
    tbl = original.tables[0]._tbl
    tbl.addprevious(p._p)
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t xml:space="preserve">Tabla </w:t></w:r>'))
    for xml in ('<w:fldChar w:fldCharType="begin"/>', '<w:instrText xml:space="preserve"> SEQ Tabla </w:instrText>',
                '<w:fldChar w:fldCharType="separate"/>', '<w:t>1</w:t>', '<w:fldChar w:fldCharType="end"/>'):
        p._p.append(parse_xml(f'<w:r {nsdecls("w")}>{xml}</w:r>'))
    p._p.append(parse_xml(f'<w:r {nsdecls("w")}><w:t xml:space="preserve">: notas medias</w:t></w:r>'))
    doc, informe = procesar(original, tmp_path)
    assert alrededor(doc)[0] == ["Tabla 1", "Notas Medias"]
    numero, titulo = parrafo(doc, "Tabla 1"), parrafo(doc, "Notas Medias")
    assert len(list(numero.iter(w("fldChar")))) == 3 and not list(titulo.iter(w("fldChar")))


def test_leyenda_con_cambios_marcados_no_se_separa(tmp_path):
    original = trabajo(antes=["Tabla 1: notas"])
    p = parrafo(original, "Tabla 1")
    p.append(parse_xml(f'<w:del {nsdecls("w")} w:id="1" w:author="A"><w:r><w:delText>x</w:delText></w:r></w:del>'))
    informe = Informe()
    regla.aplicar(original, {}, informe)
    assert any("no se ha podido separar" in a for a in informe.avisos)


def test_recuadros_y_tablas_con_imagenes_no_se_tocan():
    doc = docx.Document()
    doc.add_paragraph(TEXTO)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = "Un recuadro"
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    assert informe.entradas == []


def test_avisos_de_linea_en_blanco_y_de_posicion(tmp_path):
    original = trabajo(antes=["Tabla 1", "Notas"], despues=[TEXTO])
    informe = Informe()
    regla.aplicar(original, {}, informe)
    assert any(a.startswith("«Tabla 1» no tiene una línea en blanco") for a in informe.avisos)
    assert any(a.startswith("Comprueba en Word que cada tabla queda arriba o abajo") for a in informe.avisos)
    con_blancos = trabajo(antes=["", "Tabla 1", "Notas"], despues=["", TEXTO])
    informe = Informe()
    regla.aplicar(con_blancos, {}, informe)
    assert not any("línea en blanco" in a for a in informe.avisos)


def test_cuadro_y_numeracion_desordenada():
    original = trabajo(antes=["Cuadro 3: notas"])
    informe = Informe()
    regla.aplicar(original, {}, informe)
    assert any(a.startswith("APA llama «Tabla» a las tablas, y no «Cuadro»") for a in informe.avisos)
    assert any(a.startswith("La numeración de las tablas no va seguida (3)") for a in informe.avisos)


def test_tabla_sin_leyenda():
    informe = Informe()
    regla.aplicar(trabajo(), {}, informe)
    assert any(a.startswith("La tabla 1 no lleva número ni título") for a in informe.avisos)


def test_aviso_de_tabla_ancha_propone_pagina_horizontal(tmp_path):
    original = trabajo(antes=["Tabla 1", "Notas"])
    for columna in original.tables[0]._tbl.iterfind(f"{w('tblGrid')}/{w('gridCol')}"):
        columna.set(w("w"), "5000")   # 3 columnas de 8,8 cm: no caben entre los márgenes
    _, informe = procesar(original, tmp_path)
    assert any("APA recomienda ponerla en una página horizontal" in a for a in informe.avisos)


# --- Integridad: solo los cambios previstos, y tal cual ---------------------------------------------

def _radiografias(antes, despues):
    """Radiografías de dos listas de párrafos: [(texto, en_tabla)]."""
    def radio(lista):
        doc = docx.Document()
        for texto, en_tabla in lista:
            if en_tabla:
                doc.add_table(rows=1, cols=1).cell(0, 0).text = texto
            else:
                doc.add_paragraph(texto)
        return radiografia(doc)
    return radio(antes), radio(despues)


def test_integridad_admite_subir_y_separar():
    antes, despues = _radiografias([("Texto.", False), ("Dato", True), ("Tabla 1: Notas", False)],
                                   [("Texto.", False), ("Tabla 1", False), ("Notas", False), ("Dato", True)])
    previstos = [CambioPrevisto(3, "mover_leyenda", desde=(2,)), CambioPrevisto(3, "dividir_leyenda", quitado=": ",
                                                                                nuevo="Notas")]
    assert comparar(antes, despues, previstos) == []


@pytest.mark.parametrize("cambio", [
    CambioPrevisto(3, "dividir_leyenda", quitado=" ", nuevo="Notas"),             # no era eso lo que separaba
    CambioPrevisto(3, "dividir_leyenda", quitado=": ", nuevo="Otras"),            # otro título
    CambioPrevisto(1, "mover_leyenda", desde=(2,)),                               # no está debajo de la tabla
])
def test_integridad_rechaza_lo_demas(cambio):
    antes, despues = _radiografias([("Texto.", False), ("Dato", True), ("Tabla 1: Notas", False)],
                                   [("Texto.", False), ("Tabla 1", False), ("Notas", False), ("Dato", True)])
    otros = [CambioPrevisto(3, "mover_leyenda", desde=(2,)), CambioPrevisto(3, "dividir_leyenda", quitado=": ",
                                                                            nuevo="Notas")]
    otros = [c for c in otros if c.accion != cambio.accion] + [cambio]
    assert comparar(antes, despues, otros) != []


def test_integridad_numero_y_nota():
    antes, despues = _radiografias([("Notas", False), ("Dato", True), ("Fuente: el centro.", False)],
                                   [("Tabla 1", False), ("Notas", False), ("Dato", True),
                                    ("Nota. Fuente: el centro.", False)])
    bien = [CambioPrevisto(1, "añadir_numero_de_tabla", nuevo="Tabla 1"),
            CambioPrevisto(3, "añadir_nota", nuevo="Nota. ")]
    assert comparar(antes, despues, bien) == []
    assert comparar(antes, despues, [CambioPrevisto(1, "añadir_numero_de_tabla", nuevo="Tabla uno"), bien[1]])
    antes_texto, despues_texto = _radiografias([("Texto.", False)], [("Nota. Texto.", False)])
    assert comparar(antes_texto, despues_texto, [CambioPrevisto(1, "añadir_nota", nuevo="Nota. ")])


@pytest.mark.parametrize("antes, despues", [
    # Lo que se sube no es una leyenda con número.
    ([("Texto.", False), ("Dato", True), ("Otro texto.", False)],
     [("Texto.", False), ("Otro texto.", False), ("Dato", True)]),
    # La leyenda no está justo debajo de la tabla: hay texto en medio.
    ([("Texto.", False), ("Dato", True), ("Otro texto.", False), ("Tabla 1", False)],
     [("Texto.", False), ("Tabla 1", False), ("Dato", True), ("Otro texto.", False)]),
])
def test_integridad_solo_sube_la_leyenda_de_justo_debajo(antes, despues):
    antes, despues = _radiografias(antes, despues)
    assert comparar(antes, despues, [CambioPrevisto(len(antes.textos), "mover_leyenda", desde=(2,))])


def test_el_verificador_ve_un_titulo_sin_cursiva(tmp_path):
    doc, _ = procesar(trabajo(antes=["Tabla 1", "Notas"]), tmp_path)
    for r in _runs_propios(parrafo(doc, "Notas")):
        regla._poner(r, "i", False)
    assert incumplimientos(doc) == ["El título de la tabla: no va en cursiva."]
