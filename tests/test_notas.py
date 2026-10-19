"""
Paso 6.6: las notas al pie, según la web oficial de APA (Font y Line spacing): se
usa el formato que Word da a las notas; su letra puede ser más pequeña y su
interlineado distinto del del texto, y no hace falta cambiarlos. El ROADMAP pide
la misma fuente que el documento y no tocar el resto del formato (decisión del
propietario: el interlineado se deja como está).

La fuente la pone el paso 1.2 (reglas/fuente.py) también en las notas; estas
pruebas lo comprueban como paso propio, en el 05 (tres notas al pie) y en una
versión del 05 con las notas en otras fuentes.
"""

from pathlib import Path

import docx
import pytest
from docx.oxml import OxmlElement

from apa_facil import documento, pipeline
from apa_facil.documento import Estilos, comparar, radiografia, texto_parrafo, w
from herramientas.verificar import FUENTE, verificar

RAIZ = Path(__file__).resolve().parent.parent
TFG = RAIZ / "pruebas" / "05_tfg_completo.docx"
NOTAS = "/word/footnotes.xml"


def notas(doc):
    return documento._notas(doc, NOTAS, "footnote")


def fragmentos(nota):
    return [(r, p) for p in nota.iter(w("p")) for r in p.iter(w("r"))
            if "".join(t.text or "" for t in r.iter(w("t"))).strip()]


def formato_de_las_notas(doc):
    """[(fuentes, tamaños, interlineados)] de cada nota."""
    estilos = Estilos(doc)
    return [({estilos.fuente(r, p) for r, p in fragmentos(n)}, {estilos.tamano(r, p) for r, p in fragmentos(n)},
             {(estilos.parrafo(p, "spacing", "line"), estilos.parrafo(p, "spacing", "lineRule")) for p in n.iter(w("p"))})
            for n in notas(doc)]


def llamadas(doc):
    """Los fragmentos del texto con la llamada a una nota (el número volado)."""
    return [(r, next(r.iterancestors(w("p")))) for r in doc.element.body.iter(w("r"))
            if r.find(w("footnoteReference")) is not None]


def con_otra_fuente(ruta, destino, fuente):
    """Una copia del documento con todo el texto de sus notas en `fuente`, puesta a mano."""
    doc = docx.Document(str(ruta))

    def cambiar(nombre, raiz):
        if nombre != NOTAS:
            return
        for r in raiz.iter(w("r")):
            rpr = r.find(w("rPr"))
            if rpr is None:
                rpr = OxmlElement("w:rPr")
                r.insert(0, rpr)
            fuentes = rpr.find(w("rFonts"))
            if fuentes is None:
                fuentes = OxmlElement("w:rFonts")
                rpr.insert(0, fuentes)
            for atributo in ("ascii", "hAnsi", "cs", "eastAsia"):
                fuentes.set(w(atributo), fuente)
    documento.editar_partes_de_texto(doc, cambiar)
    doc.save(str(destino))
    return destino


@pytest.fixture(scope="module")
def tfg(tmp_path_factory):
    return pipeline.procesar(TFG, tmp_path_factory.mktemp("notas"))


def test_las_notas_del_05_llevan_la_fuente_del_texto(tfg):
    salida, _ = tfg
    copia = docx.Document(str(salida))
    assert len(notas(copia)) == 3
    for fuentes, _, _ in formato_de_las_notas(copia):
        assert fuentes == {"Times New Roman"}


def test_el_resto_del_formato_no_cambia(tfg):
    """El tamaño (10 pt) y el interlineado (1,5, heredado del documento) se quedan como estaban."""
    salida, _ = tfg
    antes = formato_de_las_notas(docx.Document(str(TFG)))
    despues = formato_de_las_notas(docx.Document(str(salida)))
    assert [t for _, t, _ in despues] == [t for _, t, _ in antes] == [{10.0}] * 3
    assert [i for _, _, i in despues] == [i for _, _, i in antes]


def test_las_llamadas_a_las_notas(tfg):
    """El número volado del texto: la fuente del texto y sigue volado."""
    salida, _ = tfg
    copia = docx.Document(str(salida))
    estilos = Estilos(copia)
    assert len(llamadas(copia)) == 3
    for r, p in llamadas(copia):
        assert estilos.fuente(r, p) == "Times New Roman"
        volado = next((capa.find(w("vertAlign")) for capa in estilos._capas_texto(r, p)
                       if capa is not None and capa.find(w("vertAlign")) is not None), None)
        assert volado is not None and volado.get(w("val")) == "superscript"


def test_el_texto_de_las_notas_no_cambia(tfg):
    salida, informe = tfg
    antes, despues = radiografia(docx.Document(str(TFG))), radiografia(docx.Document(str(salida)))
    assert [t for d, t in antes.textos if d.startswith("Nota")] == [t for d, t in despues.textos if d.startswith("Nota")]
    assert antes.notas_al_pie == despues.notas_al_pie == 3


def test_notas_en_otra_fuente_pasan_a_la_del_texto(tmp_path):
    original = con_otra_fuente(TFG, tmp_path / "notas_en_verdana.docx", "Verdana")
    fallos = [h.mensaje for h in verificar(str(original)).incumplimientos
              if h.regla == FUENTE and h.donde.startswith("Nota al pie")]
    assert len(fallos) == 3, "el verificador ve las tres notas en otra fuente"
    salida, informe = pipeline.procesar(original, tmp_path / "salida")
    copia = docx.Document(str(salida))
    assert [f for f, _, _ in formato_de_las_notas(copia)] == [{"Times New Roman"}] * 3
    assert [h for h in verificar(copia).incumplimientos if h.regla == FUENTE] == []
    assert comparar(radiografia(docx.Document(str(original))), radiografia(copia), informe.cambios_previstos) == []
