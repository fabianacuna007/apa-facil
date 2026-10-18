"""
Paso 5.2: la portada APA de estudiante, según la web oficial de APA (Title page
setup y Title case capitalization).

Los 5 documentos de pruebas/ ya tienen portada, y sustituirla o mantenerla es el
paso 5.3 (tests/test_portada_existente.py). Aquí se prueba la portada APA en un
documento sin portada. Para probarla se usa una copia temporal del 01 sin su portada (los 5
primeros párrafos); pruebas/ no se modifica.
"""

import copy
import json
from datetime import date
from pathlib import Path

import docx
import pytest

from apa_facil import documento, pipeline
from apa_facil.__main__ import main
from apa_facil.datos_portada import (DatosPortada, afiliacion, asignatura, combinar, fecha_en_espanol,
                                     parrafos_de_portada, solo_cambian_mayusculas, unir_nombres)
from apa_facil.documento import CambioPrevisto, comparar, radiografia, texto_parrafo, w
from apa_facil.estructura import clasificar
from apa_facil.reglas import portada as regla
from apa_facil.reglas.portada import mayusculas_apa
from herramientas.verificar import PORTADA, verificar

RAIZ = Path(__file__).resolve().parent.parent
DOC_01 = RAIZ / "pruebas" / "01_ensayo_basico.docx"
DATOS_01 = RAIZ / "tests" / "datos" / "01_datos.json"
TITULO_01 = "La Revolución Industrial y sus efectos en el mundo del trabajo"
TITULO_APA_01 = "La Revolución Industrial y Sus Efectos en el Mundo del Trabajo"
DEBAJO_01 = ["Javier Moreno Díaz", "Departamento de Ejemplo, Universidad de Ejemplo", "EJM 101: Historia Contemporánea",
             "Elena Sáez", "16 de octubre de 2026"]


def datos(**cambios):
    base = dict(titulo="Título", autores=("Ana Pérez",), departamento="Departamento de Historia",
                universidad="Universidad de Ejemplo", codigo_asignatura="HIS 201", asignatura="Historia",
                profesores=("Dra. López",), fecha=date(2026, 10, 16))
    base.update(cambios)
    return DatosPortada(**base)


@pytest.fixture(scope="module")
def sin_portada(tmp_path_factory):
    """El 01 sin los 5 párrafos de su portada, en una carpeta temporal."""
    ruta = tmp_path_factory.mktemp("sin_portada") / "01_sin_portada.docx"
    doc = docx.Document(str(DOC_01))
    for p in doc.paragraphs[:5]:
        p._element.getparent().remove(p._element)
    doc.save(str(ruta))
    return ruta


@pytest.fixture(scope="module")
def procesado(sin_portada, tmp_path_factory):
    salida, informe = pipeline.procesar(sin_portada, tmp_path_factory.mktemp("salida"),
                                        {"portada": combinar(DATOS_01)})
    return salida, informe, docx.Document(str(salida))


def textos(doc, cuantos):
    return [texto_parrafo(p) for p in list(doc.element.body.iter(w("p")))[:cuantos]]


# --- Las líneas de la portada ---------------------------------------------------------------

@pytest.mark.parametrize("fecha, texto", [(date(2026, 10, 16), "16 de octubre de 2026"),
                                          (date(2027, 1, 1), "1 de enero de 2027"),
                                          (date(2026, 9, 30), "30 de septiembre de 2026")])
def test_fecha_en_el_formato_de_espana(fecha, texto):
    # APA: «Use the date format commonly used in your country».
    assert fecha_en_espanol(fecha) == texto


@pytest.mark.parametrize("nombres, unidos", [
    (["Ana Pérez"], "Ana Pérez"),
    (["Ana Pérez", "Luis Gómez"], "Ana Pérez y Luis Gómez"),
    (["Ana Pérez", "Luis Gómez", "Marta Pino"], "Ana Pérez, Luis Gómez y Marta Pino"),
    (["Pedro Ruiz", "Inés Gil"], "Pedro Ruiz e Inés Gil"),
    (["Pedro Ruiz", "Hilario Gil"], "Pedro Ruiz e Hilario Gil"),
    (["Pedro Ruiz", "Íñigo Gil"], "Pedro Ruiz e Íñigo Gil"),
    (["Pedro Ruiz", "Hierro Gil"], "Pedro Ruiz y Hierro Gil"),
    (["Pedro Ruiz", "Yolanda Gil"], "Pedro Ruiz y Yolanda Gil"),
    ([], ""),
])
def test_autores_unidos_como_pide_apa_en_espanol(nombres, unidos):
    assert unir_nombres(nombres) == unidos


def test_afiliacion_y_asignatura():
    assert afiliacion(datos()) == "Departamento de Historia, Universidad de Ejemplo"
    assert afiliacion(datos(departamento=None)) == "Universidad de Ejemplo"
    assert afiliacion(datos(universidad=None)) == "Departamento de Historia"
    assert asignatura(datos()) == "HIS 201: Historia"
    assert asignatura(datos(codigo_asignatura=None)) == "Historia"
    assert asignatura(datos(asignatura=None)) == "HIS 201"
    assert asignatura(datos(codigo_asignatura="HIS 201:")) == "HIS 201: Historia"   # sin dos puntos dobles


def test_parrafos_de_la_portada():
    assert parrafos_de_portada(datos(), "Título") == [
        "", "", "", "Título", "", "Ana Pérez", "Departamento de Historia, Universidad de Ejemplo",
        "HIS 201: Historia", "Dra. López", "16 de octubre de 2026"]


def test_lo_que_falta_no_se_escribe():
    # Ni huecos ni nada en su lugar.
    assert parrafos_de_portada(datos(profesores=(), fecha=None, departamento=None, universidad=None), "T") == [
        "", "", "", "T", "", "Ana Pérez", "HIS 201: Historia"]


# --- Mayúsculas del título (title case) ------------------------------------------------------

@pytest.mark.parametrize("titulo, esperado", [
    (TITULO_01, TITULO_APA_01),
    # Artículos y conjunciones y preposiciones de 3 letras o menos, en minúscula; «Para», «Desde» y «Entre», no.
    ("el papel de la mujer: una mirada desde la historia", "El Papel de la Mujer: Una Mirada Desde la Historia"),
    ("educación para la paz entre los pueblos", "Educación Para la Paz Entre los Pueblos"),
    ("turismo con o sin guía, por y para todos", "Turismo con o sin Guía, por y Para Todos"),
    # Las palabras menores escritas con mayúscula pasan a minúscula.
    ("Turismo En La Montaña Y El Mar", "Turismo en la Montaña y el Mar"),
    # Pronombres, aunque sean cortos: palabras principales.
    ("cómo se aprende a leer", "Cómo Se Aprende a Leer"),
    # La primera palabra del subtítulo y la que sigue a una interrogación o a una raya.
    ("¿por qué leer? una propuesta", "¿Por Qué Leer? Una Propuesta"),
    ("la ciudad — un estudio", "La Ciudad — Un Estudio"),
    ("título principal\nun subtítulo en otra línea", "Título Principal\nUn Subtítulo en Otra Línea"),
    # La segunda parte de una palabra con guion («Self-Report»).
    ("análisis socio-económico", "Análisis Socio-Económico"),
    # Siglas y palabras con mayúsculas dentro, como estén.
    ("el impacto de la COVID-19 en la UE", "El Impacto de la COVID-19 en la UE"),
    ("el uso del iPhone en clase", "El Uso del iPhone en Clase"),
    ("la obra de O'Donnell", "La Obra de O'Donnell"),
    ("un siglo XXI distinto", "Un Siglo XXI Distinto"),
])
def test_mayusculas_apa(titulo, esperado):
    resultado, _ = mayusculas_apa(titulo)
    assert resultado == esperado
    assert solo_cambian_mayusculas(titulo, resultado), "solo cambian mayúsculas y minúsculas"


def test_palabras_dudosas_se_dejan_y_se_avisan():
    # «que» y «lo» pueden ser pronombres (mayúscula) o conjunción y artículo (minúscula).
    assert mayusculas_apa("lo que dicen los datos") == ("Lo que Dicen los Datos", ["que"])
    assert mayusculas_apa("todo lo que Queda") == ("Todo lo que Queda", ["lo", "que"])


def test_titulo_en_mayusculas_con_los_nombres_propios_del_trabajo():
    doc = docx.Document()
    doc.add_paragraph("En la España interior, el turismo rural crece desde hace años en varios municipios.")
    resultado, _ = mayusculas_apa("TURISMO RURAL Y DESPOBLACIÓN EN LA ESPAÑA INTERIOR", clasificar(doc))
    assert resultado == "Turismo Rural y Despoblación en la España Interior"


def test_el_articulo_de_un_nombre_propio():
    # Si el trabajo escribe «La Rioja» a mitad de frase, el artículo es parte del nombre.
    doc = docx.Document()
    doc.add_paragraph("Las bodegas de La Rioja exportan cada año más vino que las de otras regiones.")
    assert mayusculas_apa("Vinos de La Rioja", clasificar(doc))[0] == "Vinos de La Rioja"
    assert mayusculas_apa("Vinos de La Rioja")[0] == "Vinos de la Rioja", "sin el trabajo, no se sabe"


def test_solo_cambian_mayusculas():
    assert solo_cambian_mayusculas("la casa", "La Casa")
    assert not solo_cambian_mayusculas("la casa", "La Casa.")
    assert not solo_cambian_mayusculas("la casa", "La Cosa")


# --- La portada en el documento --------------------------------------------------------------

def test_la_portada_del_01(procesado):
    _, _, doc = procesado
    assert textos(doc, 12)[:10] == ["", "", "", TITULO_APA_01, ""] + DEBAJO_01
    parrafos = list(doc.element.body.iter(w("p")))
    estilos = documento.Estilos(doc)
    titulo = parrafos[3]
    assert estilos.estilo_de(titulo) == regla._id_del_estilo_del_titulo(doc)
    assert estilos.parrafo(titulo, "jc", "val") == "center"
    assert titulo.find(f"{w('pPr')}/{w('outlineLvl')}") is None, "no sale en el panel de navegación"
    for p in parrafos[5:10]:
        assert estilos.parrafo(p, "jc", "val") == "center"


def test_el_texto_empieza_con_el_mismo_titulo_en_pagina_nueva(procesado):
    _, informe, doc = procesado
    parrafos = list(doc.element.body.iter(w("p")))
    assert texto_parrafo(parrafos[10]) == TITULO_APA_01
    assert parrafos[10].find(f"{w('pPr')}/{w('outlineLvl')}").get(w("val")) == "0"
    assert parrafos[10].find(f"{w('pPr')}/{w('pageBreakBefore')}") is not None
    assert texto_parrafo(parrafos[11]) == "Introducción"
    assert not any("No se ha encontrado el título del trabajo" in a for a in informe.avisos)


def test_la_copia_cumple_el_verificador(procesado):
    salida, _, _ = procesado
    resultado = verificar(str(salida))
    assert resultado.incumplimientos == []
    assert not [h for h in resultado.excepciones if h.regla == PORTADA]


def test_el_informe_de_la_portada(procesado):
    salida, informe, _ = procesado
    assert ("Portada: se ha añadido la portada APA de estudiante con el título, el autor, la afiliación, la "
            "asignatura, el profesor y la fecha de entrega, centrados y a doble espacio, con el título en negrita "
            "tres líneas por debajo del margen superior. El texto empieza en una página nueva.") in informe.cambios
    assert any(c.startswith("Portada: el título lleva ahora mayúscula inicial en las palabras principales")
               and "--titulo-tal-cual" in c for c in informe.cambios)
    assert not any(a.startswith("Faltan datos para la portada") for a in informe.avisos)
    datos_json = json.loads(salida.with_suffix(".json").read_text(encoding="utf-8"))
    assert datos_json["opciones"]["titulo_tal_cual"] is False


def test_titulo_tal_cual(sin_portada, tmp_path):
    _, informe = pipeline.procesar(sin_portada, tmp_path, {"portada": combinar(DATOS_01), "titulo_tal_cual": True})
    doc = docx.Document(str(tmp_path / "01_sin_portada_APA.docx"))
    assert textos(doc, 11)[3] == textos(doc, 11)[10] == TITULO_01
    assert not any(c.startswith("Portada: el título lleva ahora mayúscula inicial") for c in informe.cambios)


def test_datos_que_faltan_y_palabras_dudosas(sin_portada, tmp_path):
    portada = combinar(None, {"titulo": "lo que cambió la industria", "autores": ["Ana Pérez", "Inés Gil"]})
    salida, informe = pipeline.procesar(sin_portada, tmp_path, {"portada": portada})
    doc = docx.Document(str(salida))
    assert textos(doc, 7) == ["", "", "", "Lo que Cambió la Industria", "", "Ana Pérez e Inés Gil",
                              "Lo que Cambió la Industria"]
    assert any(a.startswith("En el título de la portada se ha dejado como estaba «que»") for a in informe.avisos)
    assert any(a.startswith("Faltan datos para la portada: el departamento") for a in informe.avisos)
    assert verificar(str(salida)).incumplimientos == []


def test_sin_titulo_no_hay_portada(sin_portada, tmp_path):
    salida, informe = pipeline.procesar(sin_portada, tmp_path, {"portada": combinar(None, {"autores": ["Ana"]})})
    assert any(a.startswith("No se ha añadido la portada APA porque falta el título del trabajo, que APA exige.")
               for a in informe.avisos)
    assert textos(docx.Document(str(salida)), 1) == ["Introducción"]


def test_un_documento_con_portada_y_mantener_portada_no_cambia(tmp_path):
    # Con --mantener-portada (paso 5.3), la portada de la universidad se deja y se avisa.
    con, sin = tmp_path / "con", tmp_path / "sin"
    _, informe = pipeline.procesar(DOC_01, con, {"portada": combinar(DATOS_01), "mantener_portada": True})
    pipeline.procesar(DOC_01, sin)
    assert any(a.startswith("Se ha mantenido tu portada («La Revolución Industrial") and "como has pedido" in a
               for a in informe.avisos)
    assert radiografia(documento.abrir(con / "01_ensayo_basico_APA.docx")).textos == \
           radiografia(documento.abrir(sin / "01_ensayo_basico_APA.docx")).textos


def test_un_documento_que_empieza_con_una_tabla(tmp_path):
    doc = docx.Document()
    doc.add_table(rows=1, cols=2).cell(0, 0).text = "Dato"
    doc.add_paragraph("Un párrafo de texto con suficientes palabras para no parecer una portada de ningún tipo, "
                      "porque tiene más de veinticinco palabras seguidas en la misma frase larga.")
    ruta = tmp_path / "tabla.docx"
    doc.save(str(ruta))
    salida, _ = pipeline.procesar(ruta, tmp_path / "salida", {"portada": datos()})
    copia = docx.Document(str(salida))
    ultimo = list(copia.element.body.iter(w("p")))[9]
    assert texto_parrafo(ultimo).startswith("16 de octubre de 2026")
    assert any(b.get(w("type")) == "page" for b in ultimo.iter(w("br"))), "la tabla empieza página"
    assert [h for h in verificar(str(salida)).incumplimientos if h.regla == PORTADA] == []


# --- Integridad: solo la portada que sale de los datos ----------------------------------------------

def _portada_prevista(datos_portada, parrafos, parrafo=1):
    return CambioPrevisto(parrafo, "añadir_portada", parrafos=tuple(parrafos), datos=datos_portada)


def _con_parrafos_al_principio(textos_nuevos):
    antes = radiografia(docx.Document(str(DOC_01)))
    despues = copy.deepcopy(antes)
    despues.textos = [(f"Párrafo nuevo {n}", t) for n, t in enumerate(textos_nuevos)] + despues.textos
    return antes, despues


def test_integridad_acepta_la_portada_de_los_datos():
    lineas = parrafos_de_portada(datos(), "Título")
    antes, despues = _con_parrafos_al_principio(lineas)
    assert comparar(antes, despues, [_portada_prevista(datos(), lineas)]) == []


@pytest.mark.parametrize("cambiar", [
    lambda l: l[:5] + ["Ana Pérez (dato inventado)"] + l[6:],     # un dato que no está en los datos
    lambda l: l[:3] + ["Otro título"] + l[4:],                     # un título distinto
    lambda l: l[:3] + ["Título."] + l[4:],                         # el título con algo añadido
    lambda l: [""] + l,                                            # una línea en blanco de más
    lambda l: l + ["Curso 2025-2026"],                             # una línea de más
])
def test_integridad_rechaza_otra_portada(cambiar):
    lineas = cambiar(parrafos_de_portada(datos(), "Título"))
    antes, despues = _con_parrafos_al_principio(lineas)
    assert comparar(antes, despues, [_portada_prevista(datos(), lineas)])


def test_integridad_acepta_el_titulo_con_otras_mayusculas_y_nada_mas():
    lineas = parrafos_de_portada(datos(titulo="la casa"), "La Casa")
    antes, despues = _con_parrafos_al_principio(lineas)
    assert comparar(antes, despues, [_portada_prevista(datos(titulo="la casa"), lineas)]) == []


def test_integridad_la_portada_solo_al_principio():
    # Una portada exacta, pero detrás del primer párrafo (donde se anota): no se acepta.
    lineas = parrafos_de_portada(datos(), "Título")
    antes = radiografia(docx.Document(str(DOC_01)))
    despues = copy.deepcopy(antes)
    despues.textos[1:1] = [(f"Párrafo nuevo {n}", t) for n, t in enumerate(lineas)]
    assert comparar(antes, despues, [_portada_prevista(datos(), lineas, parrafo=2)])


def test_integridad_titulo_del_texto_solo_si_es_el_de_la_portada():
    lineas = parrafos_de_portada(datos(), "Título")
    antes, despues = _con_parrafos_al_principio(lineas + ["Título"])
    assert comparar(antes, despues, [_portada_prevista(datos(), lineas),
                                     CambioPrevisto(1, "añadir_titulo", nuevo="Título")]) == []
    antes, despues = _con_parrafos_al_principio(lineas + ["Otro"])
    assert comparar(antes, despues, [_portada_prevista(datos(), lineas),
                                     CambioPrevisto(1, "añadir_titulo", nuevo="Otro")])
    antes, despues = _con_parrafos_al_principio(["Título"])
    assert comparar(antes, despues, [CambioPrevisto(1, "añadir_titulo", nuevo="Título")]), "sin portada, no"


# --- Verificador: comprobación 5.2 Portada ------------------------------------------------------------

def _estropear(procesado, tmp_path, como):
    salida, _, _ = procesado
    doc = docx.Document(str(salida))
    como(list(doc.element.body.iter(w("p"))))
    ruta = tmp_path / "estropeado.docx"
    doc.save(str(ruta))
    return [h.mensaje for h in verificar(str(ruta)).incumplimientos if h.regla == PORTADA]


def test_verificador_dos_lineas_en_blanco(procesado, tmp_path):
    fallos = _estropear(procesado, tmp_path, lambda ps: ps[0].getparent().remove(ps[0]))
    assert "El título va 2 línea(s) por debajo del margen superior (APA: 3 o 4)." in fallos


def test_verificador_titulo_sin_centrar_ni_negrita(procesado, tmp_path):
    def estropear(ps):
        ppr = ps[3].get_or_add_pPr()
        ppr.get_or_add_jc().set(w("val"), "left")
        for r in ps[3].iter(w("r")):
            r.get_or_add_rPr().append(r.makeelement(w("b"), {w("val"): "0"}))
    assert "Título de la portada: no está centrado y no está en negrita." in _estropear(procesado, tmp_path, estropear)


def test_verificador_lineas_en_blanco_de_mas(procesado, tmp_path):
    def estropear(ps):
        ps[6].addprevious(copy.deepcopy(ps[4]))
    assert "Línea en blanco entre los datos de la portada (no debe haberla)." in _estropear(procesado, tmp_path, estropear)


def test_verificador_dato_sin_centrar(procesado, tmp_path):
    fallos = _estropear(procesado, tmp_path, lambda ps: ps[7].get_or_add_pPr().get_or_add_jc().set(w("val"), "left"))
    assert "Dato de la portada sin centrar." in fallos


def test_verificador_texto_sin_pagina_nueva(procesado, tmp_path):
    def estropear(ps):
        salto = ps[10].find(f"{w('pPr')}/{w('pageBreakBefore')}")
        salto.getparent().remove(salto)
    fallos = _estropear(procesado, tmp_path, estropear)
    assert "Después de la portada, el texto no empieza en una página nueva." in fallos


def test_verificador_no_mira_otras_portadas():
    # La portada de la universidad del 05 no sigue el esquema de APA: la comprueba 5.3 (Portada propia).
    resultado = verificar(str(RAIZ / "pruebas" / "05_tfg_completo.docx"))
    assert not [h for h in resultado.incumplimientos if h.regla == PORTADA]


# --- Línea de comandos --------------------------------------------------------------------------------

def test_linea_de_comandos_con_titulo_tal_cual(sin_portada, tmp_path, capsys):
    assert main([str(sin_portada), "--datos", str(DATOS_01), "--titulo-tal-cual"], carpeta_salida=tmp_path) == 0
    salida = capsys.readouterr().out
    assert "Portada: se ha añadido la portada APA de estudiante" in salida
    informe = json.loads((tmp_path / "01_sin_portada_APA.json").read_text(encoding="utf-8"))
    assert informe["opciones"]["titulo_tal_cual"] is True


def test_la_ayuda_explica_titulo_tal_cual(capsys):
    with pytest.raises(SystemExit):
        main(["-h"])
    assert "--titulo-tal-cual" in capsys.readouterr().out
