"""
Paso 6.4: la lista de referencias, según la web oficial de APA (Reference list
setup): en una página nueva, con el rótulo «Referencias» («Referencia» si solo hay
una) centrado y en negrita, y cada referencia con sangría francesa de 1,27 cm. El
rótulo cambia por defecto (decisión del propietario); --rotulo-tal-cual lo deja.
"""

from pathlib import Path

import docx
import pytest
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from apa_facil import estructura, pipeline
from apa_facil.documento import CambioPrevisto, Estilos, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import referencias as regla
from apa_facil.reglas.alineacion import _salto_antes
from herramientas.verificar import REFERENCIAS, verificar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {p.name[:2]: p for p in (RAIZ / "pruebas").glob("*.docx")}
TEXTO = ("El sueño cumple funciones esenciales en la consolidación de la memoria y en la regulación emocional, según "
         "los estudios revisados (Navarro y Gil, 2020).")
OBRAS = ["Castro, P. (2019). Calidad de sueño en población universitaria. Revista Ejemplo, 12(3), 201-219.",
         "Navarro, A. y Gil, S. (2020). Sueño, memoria y aprendizaje. Editorial Ejemplo."]


def trabajo(rotulo="Bibliografía", obras=OBRAS, antes_del_rotulo=None):
    doc = docx.Document()
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(TEXTO)
    if antes_del_rotulo is not None:
        antes_del_rotulo(doc)
    doc.add_paragraph(rotulo).runs[0].bold = True
    for obra in obras:
        doc.add_paragraph(obra)
    return doc


def procesar(doc, tmp_path, opciones=None):
    ruta = tmp_path / "trabajo.docx"
    doc.save(str(ruta))
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", opciones)
    return docx.Document(str(salida)), informe


def parrafo(doc, inicio):
    return next(p for p in doc.element.body.iter(w("p")) if texto_parrafo(p).startswith(inicio))


def incumplimientos(doc):
    return [h.mensaje for h in verificar(doc).incumplimientos if h.regla == REFERENCIAS]


def francesa(doc, p):
    estilos = Estilos(doc)
    return (estilos.parrafo(p, "ind", "left"), estilos.parrafo(p, "ind", "hanging"), estilos.parrafo(p, "ind", "firstLine"))


# --- Los documentos de prueba ---------------------------------------------------------------------

@pytest.fixture(scope="module")
def copias(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("referencias")
    return {clave: pipeline.procesar(ruta, carpeta) for clave, ruta in PRUEBAS.items()}


@pytest.mark.parametrize("clave, antes, obras", [("01", "BIBLIOGRAFÍA", 4), ("02", "Referencias", 2),
                                                 ("03", "Bibliografía", 2), ("04", "Webgrafía:", 3),
                                                 ("05", "BIBLIOGRAFÍA", 6)])
def test_las_listas_de_los_cinco(copias, clave, antes, obras):
    salida, informe = copias[clave]
    doc = docx.Document(str(salida))
    estilos = Estilos(doc)
    rotulo, referencias = regla.listas(estructura.clasificar(doc))[0]
    p = rotulo.elemento
    assert texto_parrafo(p) == "Referencias"
    assert estilos.nombre(estilos.estilo_de(p)).lower() == "heading 1"
    assert _salto_antes(estilos, p)
    assert len(referencias) == obras
    assert all(francesa(doc, c.elemento) == ("720", "720", None) for c in referencias)
    assert incumplimientos(doc) == []
    assert incumplimientos(docx.Document(str(PRUEBAS[clave])))
    if antes != "Referencias":
        assert any(c.startswith(f"Referencias: el rótulo «{antes}» pasa a ser «Referencias»") for c in informe.cambios)


def test_el_05_pierde_el_vacio_del_salto(copias):
    salida, informe = copias["05"]
    doc = docx.Document(str(salida))
    anterior = parrafo(doc, "Referencias").getprevious()
    assert texto_parrafo(anterior).strip(), "ya no hay un párrafo vacío con el salto de página antes del rótulo"
    assert any("se ha quitado el párrafo vacío que hacía el salto" in c for c in informe.cambios)


# --- Casos hechos a mano ---------------------------------------------------------------------------

def test_con_rotulo_tal_cual(tmp_path):
    doc, informe = procesar(trabajo(), tmp_path, {"rotulo_tal_cual": True})
    assert texto_parrafo(parrafo(doc, "Bibliograf")) == "Bibliografía"
    assert any(a.startswith("El rótulo «Bibliografía» se ha dejado como estaba porque has usado la opción "
                            "--rotulo-tal-cual") for a in informe.avisos)
    assert incumplimientos(doc) == []
    assert [h.regla for h in verificar(doc).excepciones if h.regla == REFERENCIAS] == [REFERENCIAS]


def test_una_sola_referencia(tmp_path):
    doc, _ = procesar(trabajo(obras=OBRAS[:1]), tmp_path)
    assert texto_parrafo(parrafo(doc, "Referencia")) == "Referencia"


def test_varias_listas_no_se_renombran(tmp_path):
    original = trabajo()
    original.add_paragraph("Webgrafía").runs[0].bold = True
    original.add_paragraph("Instituto Nacional de Estadística (2024). Padrón municipal. https://www.ine.es")
    doc, informe = procesar(original, tmp_path)
    assert texto_parrafo(parrafo(doc, "Bibliograf")) == "Bibliografía"
    assert texto_parrafo(parrafo(doc, "Webgraf")) == "Webgrafía"
    assert any(a.startswith("Hay 2 listas de referencias") for a in informe.avisos)


def test_texto_que_ya_acaba_en_salto_de_pagina(tmp_path):
    def salto(doc):
        doc.paragraphs[-1].add_run().add_break(docx.enum.text.WD_BREAK.PAGE)
    doc, informe = procesar(trabajo(antes_del_rotulo=salto), tmp_path)
    assert not _salto_antes(Estilos(doc), parrafo(doc, "Referencias"))
    assert not any("página nueva" in c for c in informe.cambios)
    assert incumplimientos(doc) == []


def test_los_enlaces_no_se_tocan(tmp_path):
    original = trabajo()
    p = original.add_paragraph("INE (2024). Padrón municipal. ")
    enlace = p.add_run("https://www.ine.es")
    enlace.font.color.rgb = docx.shared.RGBColor(0x05, 0x63, 0xC1)
    enlace.font.underline = True
    doc, _ = procesar(original, tmp_path)
    r = [r for r in parrafo(doc, "INE").iter(w("r")) if "ine.es" in texto_parrafo(parrafo(doc, "INE"))][-1]
    assert r.find(f"{w('rPr')}/{w('color')}") is not None and r.find(f"{w('rPr')}/{w('u')}") is not None


def test_rotulo_con_un_comentario_no_se_cambia():
    """Con un comentario de Word (o un campo, o cambios marcados), cambiar el texto podría estropearlo."""
    doc = trabajo()
    p = parrafo(doc, "Bibliograf")
    p.append(parse_xml(f'<w:r {nsdecls("w")}><w:commentReference w:id="1"/></w:r>'))
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    assert any("lleva un campo de Word o cambios marcados y no se ha podido cambiar" in a for a in informe.avisos)


# --- Integridad --------------------------------------------------------------------------------------

def _radiografias(antes, despues):
    def radio(textos):
        doc = docx.Document()
        for texto in textos:
            doc.add_paragraph(texto)
        return radiografia(doc)
    return radio(antes), radio(despues)


@pytest.mark.parametrize("antes, nuevo, admitido", [
    ("Bibliografía", "Referencias", True),
    ("WEBGRAFÍA:", "Referencia", True),
    ("Conclusiones", "Referencias", False),          # no es un rótulo de referencias
    ("Bibliografía", "Bibliografía APA", False),     # solo «Referencias» o «Referencia»
])
def test_integridad_del_rotulo(antes, nuevo, admitido):
    a, d = _radiografias([TEXTO, antes], [TEXTO, nuevo])
    diferencias = comparar(a, d, [CambioPrevisto(2, "rotulo_de_referencias", nuevo=nuevo)])
    assert (diferencias == []) == admitido


def test_el_verificador_ve_que_no_empieza_pagina(tmp_path):
    doc, _ = procesar(trabajo(), tmp_path)
    salto = parrafo(doc, "Referencias").find(f"{w('pPr')}/{w('pageBreakBefore')}")
    salto.getparent().remove(salto)
    assert incumplimientos(doc) == ["Rótulo de las referencias: la lista no empieza en una página nueva."]
