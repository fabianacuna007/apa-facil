"""
Paso 3.6: problemas que solo se avisan (saltos de línea manuales, viñetas a
mano, dobles espacios y texto en otra fuente). La regla no cambia nada.

Los casos se prueban con documentos creados aquí; después, los 5 de pruebas/.
"""

from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from docx.enum.text import WD_BREAK
from lxml import etree

from apa_facil import estructura, pipeline
from apa_facil.informe import Informe
from apa_facil.reglas import avisos, fuente

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
TEXTO = ("Este es un párrafo normal de texto con bastantes palabras para que no parezca un título "
         "ni una línea de portada, y que termina en punto como cualquier frase.")


def documento_con_portada():
    doc = docx.Document()
    doc.add_paragraph("Portada").runs[0].add_break(WD_BREAK.PAGE)
    return doc


def revisar(doc):
    informe = Informe()
    antes = etree.tostring(doc.element)
    avisos.aplicar(doc, {}, informe)
    assert etree.tostring(doc.element) == antes, "la regla de avisos no cambia nada"
    assert informe.cambios == []
    return informe.avisos


def en_fuente(doc, texto, nombre, estilo=None):
    p = doc.add_paragraph(style=estilo)
    p.add_run(texto).font.name = nombre
    return p


# --- Saltos de línea manuales -----------------------------------------------------------

def test_salto_de_linea_manual():
    doc = documento_con_portada()
    p = doc.add_paragraph("Por otro lado, las redes pueden ofrecer sentimiento de pertenencia,")
    p.runs[0].add_break()
    p.add_run("especialmente a jóvenes que se sienten aislados.")
    assert revisar(doc) == [
        "1 párrafo tiene saltos de línea manuales (Mayús+Intro) dentro del texto («…de pertenencia, ⏎ "
        "especialmente a…»). En APA el texto de un párrafo va seguido: quita el salto, o pulsa Intro si querías "
        "empezar un párrafo nuevo."]


def test_los_saltos_de_pagina_y_la_portada_no_cuentan():
    doc = docx.Document()
    portada = doc.add_paragraph("Trabajo de Psicología")
    portada.runs[0].add_break()
    portada.add_run("Nombre: Irene")
    portada.runs[-1].add_break(WD_BREAK.PAGE)
    doc.add_paragraph(TEXTO).runs[0].add_break(WD_BREAK.PAGE)
    doc.add_paragraph(TEXTO)
    assert revisar(doc) == []


# --- Viñetas escritas a mano ------------------------------------------------------------

def test_vinetas_a_mano_mezcladas():
    doc = documento_con_portada()
    for linea in ("• Comparación social", "• Búsqueda de «me gusta»", "- Ciberacoso"):
        doc.add_paragraph(linea)
    doc.add_paragraph(TEXTO)
    assert revisar(doc) == [
        "3 líneas empiezan con una viñeta escrita a mano («• Comparación social»). Además, se mezclan símbolos "
        "distintos («•» y «-»). Es mejor usar la lista con viñetas de Word (Inicio › Viñetas): así la sangría y "
        "el espaciado quedan bien. No se ha cambiado porque habría que modificar el texto."]


def test_ni_las_listas_numeradas_a_mano_ni_las_de_word():
    doc = documento_con_portada()
    doc.add_paragraph("1. ¿Desde cuándo funciona su negocio?")
    doc.add_paragraph("2. ¿Cómo ha evolucionado?")
    doc.add_paragraph("Primer elemento", style="List Bullet")
    doc.add_paragraph(TEXTO)
    assert revisar(doc) == []


# --- Dobles espacios --------------------------------------------------------------------

def test_dobles_espacios_entre_palabras():
    doc = documento_con_portada()
    doc.add_paragraph("Durante esta etapa, la identidad  todavía se está formando.  Además, cambia.")
    assert revisar(doc) == [
        "Hay 2 sitios con dos o más espacios seguidos entre palabras (entre «identidad» y «todavía» y entre "
        "«formando.» y «Además,»). APA usa un solo espacio, también después del punto. En Word puedes "
        "corregirlos todos a la vez con Ctrl+L: busca dos espacios y reemplázalos por uno."]


def test_no_cuentan_los_espacios_del_principio_ni_tras_una_vineta():
    doc = documento_con_portada()
    doc.add_paragraph("   " + TEXTO)
    doc.add_paragraph("-  Ciberacoso")
    doc.add_paragraph("1)   Primera idea")
    doc.add_paragraph(TEXTO + "   ")
    assert [a for a in revisar(doc) if "espacios" in a] == []


# --- Texto en otra fuente ---------------------------------------------------------------

def test_parrafo_en_otra_fuente():
    doc = documento_con_portada()
    for _ in range(3):
        en_fuente(doc, TEXTO, "Calibri")
    en_fuente(doc, "Las redes sociales son plataformas digitales que permiten crear perfiles.", "Segoe UI")
    assert revisar(doc) == [
        "1 párrafo estaba en otra fuente, distinta de la del resto del trabajo (Calibri): Segoe UI en «Las redes "
        "sociales son pl…». Suele pasar al pegar texto de una página web o de otro documento. La fuente ya se ha "
        "igualado; comprueba que ese texto está citado y que se ve bien."]


def test_varios_parrafos_en_otra_fuente():
    doc = documento_con_portada()
    for _ in range(4):
        en_fuente(doc, TEXTO, "Calibri")
    for nombre in ("Segoe UI", "Times New Roman", "Georgia"):
        en_fuente(doc, TEXTO, nombre)
    aviso = revisar(doc)[0]
    assert aviso.startswith("3 párrafos estaban en otra fuente")
    assert "Segoe UI en «Este es un párrafo normal…», Times New Roman en «Este es un párrafo normal…» y 1 más" in aviso


def test_lo_que_no_cuenta_como_otra_fuente():
    doc = documento_con_portada()
    for _ in range(3):
        en_fuente(doc, TEXTO, "Calibri")
    titulo = doc.add_heading("Método", level=1)
    titulo.runs[0].font.name = "Arial Black"                                  # un título
    en_fuente(doc, TEXTO, "Calibri Light")                                    # otra variante de la familia
    en_fuente(doc, "for i in range(10): print(i)", "Courier New")             # código
    corto = doc.add_paragraph()
    corto.add_run("Una palabra en ")
    corto.runs[0].font.name = "Calibri"
    corto.add_run("cursiva").font.name = "Georgia"                            # menos de 20 letras
    corto.add_run(" y el resto del párrafo normal, que acaba en punto.").font.name = "Calibri"
    celda = doc.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0]
    celda.add_run(TEXTO).font.name = "Times New Roman"                        # en una tabla
    assert revisar(doc) == []


def test_la_portada_no_cuenta_como_otra_fuente():
    doc = docx.Document()
    portada = en_fuente(doc, "REDES SOCIALES Y AUTOESTIMA EN LA ADOLESCENCIA", "Comic Sans MS")
    portada.runs[0].add_break(WD_BREAK.PAGE)
    for _ in range(3):
        en_fuente(doc, TEXTO, "Calibri")
    assert revisar(doc) == []


# --- Orden en el pipeline ---------------------------------------------------------------

def test_los_avisos_se_calculan_antes_de_igualar_la_fuente(tmp_path):
    """En el 04, después de la regla de fuente todo es Calibri: el aviso solo puede salir si va antes."""
    original = next(p for p in PRUEBAS if p.name.startswith("04"))
    assert pipeline.REGLAS.index(avisos) < pipeline.REGLAS.index(fuente)
    _, informe = pipeline.procesar(original, tmp_path)
    assert any(a.startswith("7 párrafos estaban en otra fuente") for a in informe.avisos)


def test_despues_de_la_regla_de_fuente_ya_no_se_veria(tmp_path, monkeypatch):
    """A la contra: con la regla de avisos detrás de la de fuente, el aviso desaparece."""
    original = next(p for p in PRUEBAS if p.name.startswith("04"))
    reglas = [r for r in pipeline.REGLAS if r is not avisos]
    reglas.insert(reglas.index(fuente) + 1, avisos)
    monkeypatch.setattr(pipeline, "REGLAS", reglas)
    _, informe = pipeline.procesar(original, tmp_path)
    assert not any("en otra fuente" in a for a in informe.avisos)


# --- Los 5 documentos de prueba ---------------------------------------------------------

ESPERADO = {   # el principio de cada aviso de esta regla, en orden
    "01": [],
    "02": ["1 párrafo estaba en otra fuente"],
    "03": [],
    "04": ["1 párrafo tiene saltos de línea manuales", "4 líneas empiezan con una viñeta escrita a mano",
           "Hay 1 sitio con dos o más espacios seguidos", "7 párrafos estaban en otra fuente"],
    "05": [],
}


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    _, informe = pipeline.procesar(original, tmp_path)
    de_esta_regla = [e.mensaje for e in informe.entradas if e.regla == "avisos"]
    esperado = ESPERADO[original.name[:2]]
    assert len(de_esta_regla) == len(esperado), de_esta_regla
    for mensaje, principio in zip(de_esta_regla, esperado):
        assert mensaje.startswith(principio)
    assert all(e.tipo == "aviso" for e in informe.entradas if e.regla == "avisos")
