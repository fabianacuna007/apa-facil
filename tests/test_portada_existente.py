"""
Paso 5.3: la portada que ya tenía el documento (la de la universidad).

Por defecto, con datos de portada, se sustituye por la portada APA de estudiante
(decisión del propietario), logo incluido. Se mantiene con --mantener-portada,
sin datos, sin título o si quitarla podría borrar algo que no es de la portada.
Si se mantiene, se deja con su diseño y el texto empieza en una página nueva.

tests/datos/05_datos.json lleva los datos de la portada del 05 (título, autora,
facultad, universidad, tutor y «Trabajo Fin de Grado») y, para lo que no tiene,
valores de prueba ficticios: el código «EJM 400» y el día de la fecha (la portada
solo dice «junio de 2026»).
"""

import copy
import json
from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from docx.oxml import OxmlElement

from apa_facil import documento, pipeline
from apa_facil.__main__ import main
from apa_facil.datos_portada import combinar, parrafos_de_portada
from apa_facil.documento import CambioPrevisto, comparar, radiografia, texto_parrafo, w
from apa_facil.estructura import clasificar
from apa_facil.informe import Informe
from apa_facil.reglas import portada as regla
from herramientas.verificar import PORTADA_PROPIA, verificar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = RAIZ / "pruebas"
DOC_01, DOC_05 = PRUEBAS / "01_ensayo_basico.docx", PRUEBAS / "05_tfg_completo.docx"
DATOS_01, DATOS_05 = RAIZ / "tests" / "datos" / "01_datos.json", RAIZ / "tests" / "datos" / "05_datos.json"


def primeros(ruta, cuantos):
    return [texto_parrafo(p) for p in list(docx.Document(str(ruta)).element.body.iter(w("p")))[:cuantos]]


@pytest.fixture(scope="module")
def sustituida_05(tmp_path_factory):
    salida, informe = pipeline.procesar(DOC_05, tmp_path_factory.mktemp("05"), {"portada": combinar(DATOS_05)})
    return salida, informe


# --- Sustituir (por defecto) --------------------------------------------------------------------

def test_el_01_recibe_una_portada_apa_completa(tmp_path):
    # Criterio de la fase 5: «el documento 01 recibe una portada APA completa a partir de datos.json».
    salida, informe = pipeline.procesar(DOC_01, tmp_path, {"portada": combinar(DATOS_01)})
    assert primeros(salida, 12) == [
        "", "", "", "La Revolución Industrial y Sus Efectos en el Mundo del Trabajo", "", "Javier Moreno Díaz",
        "Departamento de Ejemplo, Universidad de Ejemplo", "EJM 101: Historia Contemporánea", "Elena Sáez",
        "16 de octubre de 2026", "La Revolución Industrial y Sus Efectos en el Mundo del Trabajo", "Introducción"]
    todo = " ".join(primeros(salida, 400))
    assert "Alumno:" not in todo and "Curso 2025-2026" not in todo, "la portada antigua ya no está"
    assert any(c.startswith("Portada: se ha sustituido tu portada por la portada APA de estudiante. Se han quitado "
                            "sus 5 líneas («La Revolución Industrial y sus efectos e…», «Alumno: Javier Moreno "
                            "Díaz», «Asignatura: Historia Contemporánea»…)") for c in informe.cambios)
    resultado = verificar(str(salida))
    assert resultado.incumplimientos == []
    assert not [h for h in resultado.excepciones if h.regla == PORTADA_PROPIA]


def test_el_05_sin_logo_y_con_sus_secciones(sustituida_05):
    salida, informe = sustituida_05
    original, copia = radiografia(documento.abrir(DOC_05)), radiografia(documento.abrir(salida))
    assert copia.imagenes == original.imagenes - 1, "el logo, y solo el logo"
    assert copia.secciones == original.secciones == 3
    parrafos = list(docx.Document(str(salida)).element.body.iter(w("p")))
    assert texto_parrafo(parrafos[9]) == "15 de junio de 2026"
    assert parrafos[9].find(f"{w('pPr')}/{w('sectPr')}") is not None, "el salto de sección, al final de la portada"
    assert texto_parrafo(parrafos[10]) == "Resumen", "detrás, el resumen (en su página)"
    assert any("y la imagen (el logo)" in c for c in informe.cambios)
    assert not any("alineados a la derecha" in a for a in informe.avisos), "esas líneas ya no están"


def test_el_05_cumple_el_verificador(sustituida_05):
    salida, _ = sustituida_05
    resultado = verificar(str(salida))
    assert resultado.incumplimientos == []
    assert not [h for h in resultado.excepciones if h.regla in (PORTADA_PROPIA, "1.4 Alineación")]


def test_el_titulo_del_texto_sale_de_los_datos(sustituida_05):
    salida, _ = sustituida_05
    titulos = [texto_parrafo(p) for p in docx.Document(str(salida)).element.body.iter(w("p"))
               if texto_parrafo(p).startswith("Turismo Rural")]
    assert titulos == ["Turismo Rural y Despoblación en la España Interior: Estudio de Cuatro Municipios"] * 2


# --- Mantener -------------------------------------------------------------------------------------

def test_mantener_portada(tmp_path):
    salida, informe = pipeline.procesar(DOC_05, tmp_path, {"portada": combinar(DATOS_05), "mantener_portada": True})
    assert primeros(salida, 11)[1:4] == ["UNIVERSIDAD AUTÓNOMA DEL GUADIANA", "Facultad de Ciencias Sociales y Humanas",
                                         "Grado en Geografía y Ordenación del Territorio"]
    assert radiografia(documento.abrir(salida)).imagenes == radiografia(documento.abrir(DOC_05)).imagenes
    assert any(a.startswith("Se ha mantenido tu portada («UNIVERSIDAD AUTÓNOMA DEL GUADIANA»), como has pedido. "
                            "APA pide la portada de estudiante salvo que tu universidad o tu profesor exijan otra: "
                            "no se han usado tus datos de portada.") for a in informe.avisos)
    resultado = verificar(str(salida))
    assert resultado.incumplimientos == []
    assert [h.regla for h in resultado.excepciones if h.regla == PORTADA_PROPIA] == [PORTADA_PROPIA]


def test_sin_datos_se_mantiene_y_el_texto_empieza_pagina(tmp_path):
    # El 01 no acaba su portada en un salto de página: ahora el texto empieza en la página siguiente.
    salida, informe = pipeline.procesar(DOC_01, tmp_path)
    assert any(a.startswith("Se ha mantenido tu portada («La Revolución Industrial y sus efectos …») porque no "
                            "has dado los datos") for a in informe.avisos)
    assert "Portada: el texto empieza ahora en una página nueva, después de tu portada, como pide APA." \
           in informe.cambios
    parrafos = list(docx.Document(str(salida)).element.body.iter(w("p")))
    assert texto_parrafo(parrafos[5]) == "La Revolución Industrial y Sus Efectos en el Mundo del Trabajo"
    assert parrafos[5].find(f"{w('pPr')}/{w('pageBreakBefore')}") is not None
    assert verificar(str(salida)).incumplimientos == []


def test_una_portada_que_ya_acaba_en_salto_no_se_toca(tmp_path):
    # La del 05 acaba en un salto de sección: no hace falta nada más.
    _, informe = pipeline.procesar(DOC_05, tmp_path)
    assert not any(c.startswith("Portada: el texto empieza ahora") for c in informe.cambios)


def test_datos_sin_titulo(tmp_path):
    salida, informe = pipeline.procesar(DOC_01, tmp_path, {"portada": combinar(None, {"autores": ["Ana"]})})
    assert any(a.startswith("Se ha mantenido tu portada") and "falta el título del trabajo" in a for a in informe.avisos)
    assert primeros(salida, 1) == ["La Revolución Industrial y sus efectos en el mundo del trabajo"]


def _portada_del_01():
    """El 01 y los párrafos de su portada según el clasificador (que se pueden estropear después)."""
    doc = docx.Document(str(DOC_01))
    return doc, regla.portada_existente(clasificar(doc))


def test_no_se_quita_una_portada_con_una_tabla_en_medio():
    # Por si el clasificador se equivoca: quitar la portada borraría la tabla.
    _, portada = _portada_del_01()
    portada[1].elemento.addnext(OxmlElement("w:tbl"))
    assert regla._por_que_no_se_puede_quitar(portada) == "hay una tabla u otro elemento entre sus líneas"


def test_no_se_quita_una_portada_con_un_parrafo_largo():
    _, portada = _portada_del_01()
    portada[2] = SimpleNamespace(elemento=portada[2].elemento, texto="palabra " * 30)
    assert regla._por_que_no_se_puede_quitar(portada) == "tiene un párrafo largo, que podría ser parte del texto"


def test_no_se_quita_una_portada_con_una_nota():
    _, portada = _portada_del_01()
    p = portada[1].elemento
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:footnoteReference"))
    p.append(r)
    assert regla._por_que_no_se_puede_quitar(portada) == ("lleva una nota, un campo, un comentario o un cuadro de "
                                                          "texto")


def test_no_se_quita_una_portada_con_un_salto_de_seccion_en_medio():
    _, portada = _portada_del_01()
    portada[1].elemento.get_or_add_pPr().append(OxmlElement("w:sectPr"))
    assert regla._por_que_no_se_puede_quitar(portada) == "tiene un salto de sección entre sus líneas"


def test_si_no_se_puede_quitar_se_mantiene_y_se_avisa(tmp_path):
    doc = docx.Document(str(DOC_01))
    r = OxmlElement("w:r")
    r.append(OxmlElement("w:footnoteReference"))
    doc.paragraphs[1]._p.append(r)
    clasificados = clasificar(doc)
    assert regla.que_hacer({"portada": combinar(DATOS_01)}, clasificados) == (
        "mantener", "lleva una nota, un campo, un comentario o un cuadro de texto")
    informe = Informe()
    regla.aplicar(doc, {"portada": combinar(DATOS_01)}, informe, clasificados)
    assert any("no se ha sustituido por la de APA porque lleva una nota" in a for a in informe.avisos)


# --- Integridad --------------------------------------------------------------------------------------

def test_integridad_sin_anotar_que_se_quita_la_portada(sustituida_05):
    salida, informe = sustituida_05
    antes, despues = radiografia(documento.abrir(DOC_05)), radiografia(documento.abrir(salida))
    assert comparar(antes, despues, informe.cambios_previstos) == []
    sin_quitar = [c for c in informe.cambios_previstos if c.accion != "quitar_portada"]
    assert comparar(antes, despues, sin_quitar), "sin anotarlo, la integridad lo ve"


def _sustitucion(antes, quitados):
    """La copia que quedaría quitando los párrafos `quitados` y poniendo la portada APA, y lo anotado."""
    datos = combinar(DATOS_01)
    lineas = parrafos_de_portada(datos, datos.titulo)
    despues = copy.deepcopy(antes)
    despues.textos = [(f"Nuevo {n}", t) for n, t in enumerate(lineas)] + \
                     [t for n, t in enumerate(antes.textos, start=1) if n not in quitados]
    previstos = [CambioPrevisto(n, "quitar_portada") for n in quitados] + \
                [CambioPrevisto(1, "añadir_portada", parrafos=tuple(lineas), datos=datos)]
    return despues, previstos


def test_integridad_acepta_sustituir_la_portada():
    antes = radiografia(documento.abrir(DOC_01))
    assert comparar(antes, *_sustitucion(antes, range(1, 6))) == []


def test_integridad_solo_la_portada_del_principio():
    # Quitar del 2 al 6 («Introducción» incluida) no es quitar la portada.
    antes = radiografia(documento.abrir(DOC_01))
    assert comparar(antes, *_sustitucion(antes, range(2, 7)))


def test_integridad_no_quita_parrafos_largos():
    antes = radiografia(documento.abrir(DOC_01))
    antes.textos[1] = (antes.textos[1][0], "palabra " * 30)
    assert comparar(antes, *_sustitucion(antes, range(1, 6)))


def test_integridad_no_quita_portada_sin_poner_la_de_apa():
    antes = radiografia(documento.abrir(DOC_01))
    despues = copy.deepcopy(antes)
    del despues.textos[0:5]
    assert comparar(antes, despues, [CambioPrevisto(n, "quitar_portada") for n in range(1, 6)])


def test_integridad_cuenta_las_imagenes_del_logo(sustituida_05):
    salida, informe = sustituida_05
    antes, despues = radiografia(documento.abrir(DOC_05)), radiografia(documento.abrir(salida))
    despues.imagenes -= 1   # como si se hubiera borrado otra imagen además del logo
    assert any("quitando las 1 de la portada sustituida" in d for d in comparar(antes, despues,
                                                                                 informe.cambios_previstos))


# --- Verificador -----------------------------------------------------------------------------------------

def test_verificador_portada_propia_sin_pagina_nueva():
    # El original del 01: su portada no acaba en salto y el texto sigue en la misma página.
    fallos = [h.mensaje for h in verificar(str(DOC_01)).incumplimientos if h.regla == PORTADA_PROPIA]
    assert fallos == ["Después de la portada, el texto no empieza en una página nueva."]


# --- Línea de comandos ---------------------------------------------------------------------------------

def test_linea_de_comandos_mantener_portada(tmp_path, capsys):
    assert main([str(DOC_05), "--datos", str(DATOS_05), "--mantener-portada"], carpeta_salida=tmp_path) == 0
    assert "como has pedido" in capsys.readouterr().out
    informe = json.loads((tmp_path / "05_tfg_completo_APA.json").read_text(encoding="utf-8"))
    assert informe["opciones"]["mantener_portada"] is True


def test_la_ayuda_explica_mantener_portada(capsys):
    with pytest.raises(SystemExit):
        main(["-h"])
    assert "--mantener-portada" in capsys.readouterr().out
