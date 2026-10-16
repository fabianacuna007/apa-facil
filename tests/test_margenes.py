"""
Paso 1.1: márgenes de 2,54 cm en los cuatro lados, en todas las secciones.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

import base64
import io
from pathlib import Path

import docx
import pytest
from docx.enum.section import WD_ORIENT
from docx.shared import Cm

from apa_facil import pipeline
from apa_facil.__main__ import main
from apa_facil.documento import w
from apa_facil.informe import Informe
from apa_facil.reglas import margenes

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
APA = 1440   # 2,54 cm en twips

# Una imagen PNG de 1 × 1 píxeles, para poder insertar imágenes de cualquier tamaño.
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")


def margenes_de(seccion):
    return (seccion.top_margin.twips, seccion.bottom_margin.twips,
            seccion.left_margin.twips, seccion.right_margin.twips)


def poner(seccion, superior, inferior, izquierdo, derecho, encuadernacion=0):
    seccion.top_margin, seccion.bottom_margin = Cm(superior), Cm(inferior)
    seccion.left_margin, seccion.right_margin = Cm(izquierdo), Cm(derecho)
    seccion.gutter = Cm(encuadernacion)


def aplicar(doc):
    informe = Informe()
    margenes.aplicar(doc, {}, informe)
    return informe


def documento_dos_secciones():
    """Sección 1 vertical con encuadernación; sección 2 horizontal."""
    doc = docx.Document()
    poner(doc.sections[0], 1, 1, 2, 2, encuadernacion=1)
    doc.add_paragraph("Primera sección")
    horizontal = doc.add_section()
    horizontal.orientation = WD_ORIENT.LANDSCAPE
    horizontal.page_width, horizontal.page_height = Cm(29.7), Cm(21)
    poner(horizontal, 1.27, 1.27, 1.27, 1.27)
    doc.add_paragraph("Segunda sección")
    return doc


def tabla_de_ancho(doc, centimetros, columnas=3):
    tabla = doc.add_table(rows=1, cols=columnas)
    tabla.rows[0].cells[0].text = "Variable"
    for col in tabla._tbl.iter(w("gridCol")):
        col.set(w("w"), str(int(Cm(centimetros / columnas).twips)))
    return tabla


def test_cuatro_lados_en_todas_las_secciones():
    doc = documento_dos_secciones()
    aplicar(doc)
    for seccion in doc.sections:
        assert margenes_de(seccion) == (APA,) * 4
        assert seccion.gutter.twips == 0


def test_conserva_tamano_y_orientacion():
    doc = documento_dos_secciones()
    antes = [(s.orientation, s.page_width, s.page_height) for s in doc.sections]
    aplicar(doc)
    assert [(s.orientation, s.page_width, s.page_height) for s in doc.sections] == antes
    assert doc.sections[1].orientation == WD_ORIENT.LANDSCAPE


def test_no_toca_la_distancia_del_encabezado():
    doc = documento_dos_secciones()
    antes = [(s.header_distance, s.footer_distance) for s in doc.sections]
    aplicar(doc)
    assert [(s.header_distance, s.footer_distance) for s in doc.sections] == antes


def test_el_informe_explica_cada_seccion():
    informe = aplicar(documento_dos_secciones())
    assert len(informe.cambios) == 2
    primera, segunda = informe.cambios
    assert primera.startswith("Sección 1:")
    assert "superior 1,00 · inferior 1,00 · izquierdo 2,00 · derecho 2,00 cm" in primera
    assert "encuadernación de 1,00 cm" in primera
    assert segunda.startswith("Sección 2 (página horizontal, se mantiene así):")
    assert "antes: 1,27 cm en todos" in segunda
    assert informe.avisos == informe.errores == []


def test_secciones_iguales_van_en_un_solo_mensaje():
    doc = docx.Document()
    for _ in range(2):
        doc.add_section()
    for seccion in doc.sections:
        poner(seccion, 2.5, 2.5, 3, 3)
    assert aplicar(doc).cambios == [
        "Secciones 1, 2 y 3: márgenes cambiados a 2,54 cm en los cuatro lados "
        "(antes: superior 2,50 · inferior 2,50 · izquierdo 3,00 · derecho 3,00 cm)."]


def test_una_sola_seccion_no_dice_seccion():
    doc = docx.Document()
    poner(doc.sections[0], 1.5, 1.5, 1.5, 1.5)
    assert aplicar(doc).cambios == [
        "Márgenes cambiados a 2,54 cm en los cuatro lados (antes: 1,50 cm en todos)."]


def test_si_ya_cumple_no_cambia_nada():
    doc = docx.Document()
    poner(doc.sections[0], 2.54, 2.54, 2.54, 2.54)
    informe = aplicar(doc)
    assert informe.cambios == informe.avisos == []
    assert margenes_de(doc.sections[0]) == (APA,) * 4


def test_solo_encuadernacion():
    doc = docx.Document()
    poner(doc.sections[0], 2.54, 2.54, 2.54, 2.54, encuadernacion=0.5)
    informe = aplicar(doc)
    assert doc.sections[0].gutter.twips == 0
    assert informe.cambios == ["Márgenes de 2,54 cm en los cuatro lados (ya lo eran). Se ha quitado "
                               "el margen de encuadernación de 0,50 cm, que hacía más ancho uno de los márgenes."]


# --- Aviso de lo que no cabe ------------------------------------------------------
# En un A4 vertical con 2,54 cm de margen caben 21,00 − 5,08 = 15,92 cm.

def test_avisa_de_una_tabla_que_no_cabe():
    doc = docx.Document()
    poner(doc.sections[0], 1.5, 1.5, 1.5, 1.5)
    doc.sections[0].page_width = Cm(21)
    tabla_de_ancho(doc, 15)   # cabe
    tabla_de_ancho(doc, 17)   # no cabe
    informe = aplicar(doc)
    assert len(informe.avisos) == 1
    aviso = informe.avisos[0]
    assert aviso.startswith("Tabla 2 («Variable») mide 17,00 cm de ancho")
    assert "solo caben 15,92 cm" in aviso
    assert "Autoajustar a la ventana" in aviso


def test_avisa_de_una_imagen_que_no_cabe():
    doc = docx.Document()
    doc.sections[0].page_width = Cm(21)
    doc.add_picture(io.BytesIO(PNG), width=Cm(10))
    doc.add_picture(io.BytesIO(PNG), width=Cm(18))
    informe = aplicar(doc)
    assert len(informe.avisos) == 1
    assert informe.avisos[0].startswith("Imagen 2 mide 18,00 cm de ancho")


def test_no_avisa_de_imagenes_dentro_de_tablas():
    doc = docx.Document()
    doc.sections[0].page_width = Cm(21)
    tabla = tabla_de_ancho(doc, 15, columnas=1)
    tabla.rows[0].cells[0].paragraphs[0].add_run().add_picture(io.BytesIO(PNG), width=Cm(18))
    assert aplicar(doc).avisos == []


def test_tabla_ancha_cabe_en_seccion_horizontal():
    """24 cm no caben en vertical (15,92) pero sí en horizontal (24,62)."""
    doc = docx.Document()
    doc.sections[0].page_width = Cm(21)
    doc.add_paragraph("Vertical")
    horizontal = doc.add_section()
    horizontal.page_width, horizontal.page_height = Cm(29.7), Cm(21)
    tabla_de_ancho(doc, 24)
    assert aplicar(doc).avisos == []


def test_tabla_en_porcentaje_no_avisa():
    doc = docx.Document()
    doc.sections[0].page_width = Cm(21)
    tabla = tabla_de_ancho(doc, 30)
    ancho = tabla._tbl.tblPr.find(w("tblW"))
    ancho.set(w("type"), "pct")
    ancho.set(w("w"), "5000")   # 100 % del espacio disponible
    assert aplicar(doc).avisos == []


# --- Documentos de prueba -----------------------------------------------------------

@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    orig = docx.Document(str(original))
    salida, informe = pipeline.procesar(original, tmp_path)
    copia = docx.Document(str(salida))

    assert len(copia.sections) == len(orig.sections)
    for antes, despues in zip(orig.sections, copia.sections):
        assert margenes_de(despues) == (APA,) * 4
        assert despues.gutter.twips == 0
        assert (despues.orientation, despues.page_width, despues.page_height) == \
               (antes.orientation, antes.page_width, antes.page_height)
    assert informe.cambios, "ningún documento de prueba cumplía los márgenes"
    assert informe.avisos == [], "en los documentos de prueba todo cabe entre los márgenes"
    assert informe.errores == []


def test_la_seccion_horizontal_del_04_sigue_horizontal(tmp_path):
    original = next(p for p in PRUEBAS if p.name.startswith("04"))
    salida, informe = pipeline.procesar(original, tmp_path)
    assert docx.Document(str(salida)).sections[1].orientation == WD_ORIENT.LANDSCAPE
    assert any(c.startswith("Sección 2 (página horizontal") for c in informe.cambios)


def test_la_linea_de_comandos_muestra_los_cambios(tmp_path, capsys):
    assert main([str(PRUEBAS[0])], carpeta_salida=tmp_path) == 0
    salida = capsys.readouterr().out
    assert "Cambios:" in salida
    assert "Márgenes cambiados a 2,54 cm" in salida
