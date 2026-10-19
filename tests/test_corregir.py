"""
Paso 7.7: los errores de forma de las citas y las referencias que se pueden
corregir sin inventar nada se corrigen en la copia (salvo con --citas-tal-cual).
"""

import docx
import pytest

from apa_facil import documento, pipeline
from apa_facil.documento import CambioPrevisto, correccion_de_forma_permitida
from tests.test_revision import TEXTO, trabajo

ORTEGA = "Ortega, L. y Blanco, R. (2018). El tiempo de la fábrica. Revista de Historia Social, 34(2), 45-67."
MOLINA = "Molina, R. (2017). Turismo y territorio. Editorial Ejemplo."


def procesar(doc, tmp_path, opciones=None):
    entrada = tmp_path / "trabajo.docx"
    doc.save(str(entrada))
    salida, informe = pipeline.procesar(entrada, tmp_path / "salida", opciones or {})   # con la integridad
    return [p.text for p in docx.Document(str(salida)).paragraphs], informe, salida


def avisos_de_revision(informe):
    return [e.mensaje for e in informe.entradas if e.regla == "revision" and e.tipo == "aviso"]


@pytest.mark.parametrize("cita, corregida", [
    ("(Molina 2017)", "(Molina, 2017)"),
    ("(Ortega et. al, 2018)", "(Ortega et al., 2018)"),
    ("(Molina, s.f)", "(Molina, s. f.)"),
    ("(Ortega & Blanco, 2018)", "(Ortega y Blanco, 2018)"),
    ("(Molina, 2017, p.45)", "(Molina, 2017, p. 45)"),
    ("(Molina, 2017, p. 45-47)", "(Molina, 2017, pp. 45-47)"),
    ("(Ortega et. al 2018)", "(Ortega et al., 2018)"),    # dos errores en la misma cita
])
def test_citas(cita, corregida, tmp_path):
    textos, informe, _ = procesar(trabajo([f"{TEXTO} {cita}."], [MOLINA, ORTEGA]), tmp_path)
    assert f"{TEXTO} {corregida}." in textos
    assert not [m for m in avisos_de_revision(informe) if "citas del texto" in m]
    cambios = [e.mensaje for e in informe.entradas if e.regla == "revision" and e.tipo == "cambio"]
    assert all(f"«{cita}» → «{corregida}»" in cambio for cambio in cambios)   # uno por cada tipo de error


@pytest.mark.parametrize("referencia, corregida", [
    ("Ortega, L.; Blanco, R. (2018). El tiempo. Editorial Ejemplo.",
     "Ortega, L. y Blanco, R. (2018). El tiempo. Editorial Ejemplo."),
    ("Ortega, L., & Blanco, R. (2018). El tiempo. Editorial Ejemplo.",
     "Ortega, L. y Blanco, R. (2018). El tiempo. Editorial Ejemplo."),
    ("Molina, R. 2017. Turismo y territorio. Editorial Ejemplo.", MOLINA),
    ("Molina, R. (2017) Turismo y territorio. Editorial Ejemplo.", MOLINA),
    ("Molina, R. (2017), Turismo y territorio. Editorial Ejemplo.", MOLINA),
    ("Molina, R. (2017). Turismo y territorio. Madrid: Editorial Ejemplo.", MOLINA),
    ("Molina, R. (2017). Turismo y territorio. Editorial Ejemplo, Madrid.", MOLINA),
    ("Ortega, L. (2018). El tiempo. Revista Social, 34, nº 2, 45-67. doi:10.1234/rs.2018.2",
     "Ortega, L. (2018). El tiempo. Revista Social, 34(2), 45-67. https://doi.org/10.1234/rs.2018.2"),
    ("Ortega, L. (2018). El tiempo. Revista Social, 34(2), 45-67. Recuperado de https://ejemplo.org/tiempo",
     "Ortega, L. (2018). El tiempo. Revista Social, 34(2), 45-67. https://ejemplo.org/tiempo"),
])
def test_referencias(referencia, corregida, tmp_path):
    cita = "(Ortega y Blanco, 2018)" if "Blanco" in referencia else "(Ortega, 2018)" if "Ortega" in referencia \
        else "(Molina, 2017)"
    textos, informe, _ = procesar(trabajo([f"{TEXTO} {cita}."], [referencia]), tmp_path)
    assert corregida in textos, textos
    assert not avisos_de_revision(informe)


def test_lo_que_obligaria_a_inventar_solo_se_avisa(tmp_path):
    referencias = ["Molina, R. Turismo y territorio. Editorial Ejemplo.",                 # falta el año
                   "Ortega, L. (2018). El tiempo de la fábrica. Revista de Historia Social, 34."]   # sin páginas
    textos, informe, _ = procesar(trabajo([f"{TEXTO} (Molina, s. f.) (Ortega, 2018)."], referencias), tmp_path)
    for referencia in referencias:
        assert referencia in textos
    avisos = avisos_de_revision(informe)
    assert any(m.startswith("Falta el año") for m in avisos)
    assert any(m.startswith("Referencia incompleta") for m in avisos)


def test_con_citas_tal_cual_solo_se_avisa(tmp_path):
    doc = trabajo([f"{TEXTO} (Molina 2017)."], ["Molina, R. (2017) Turismo y territorio. Madrid: Editorial Ejemplo."])
    textos, informe, _ = procesar(doc, tmp_path, {"citas_tal_cual": True})
    assert f"{TEXTO} (Molina 2017)." in textos
    assert "Molina, R. (2017) Turismo y territorio. Madrid: Editorial Ejemplo." in textos
    assert not [e for e in informe.entradas if e.regla == "revision" and e.tipo == "cambio"]
    assert len(avisos_de_revision(informe)) == 3


def test_se_conserva_la_cursiva(tmp_path):
    doc = trabajo([f"{TEXTO} (Ortega, 2018)."], [])
    p = doc.add_paragraph("Ortega, L. (2018). El tiempo. ")
    p.add_run("Revista Social, 34").italic = True
    p.add_run(", nº 2, 45-67.")
    _, _, salida = procesar(doc, tmp_path)
    (corregida,) = [p for p in docx.Document(str(salida)).paragraphs if p.text.startswith("Ortega")]
    assert corregida.text == "Ortega, L. (2018). El tiempo. Revista Social, 34(2), 45-67."
    cursiva = "".join(r.text for r in corregida.runs if r.italic)
    assert cursiva == "Revista Social, 34"   # el número de la revista no va en cursiva


def test_con_comentarios_solo_lo_que_queda_por_revisar(tmp_path):
    doc = trabajo([f"{TEXTO} (Molina 2017) (Castell, 2011)."], [MOLINA])
    textos, informe, salida = procesar(doc, tmp_path, {"comentarios": True})
    assert f"{TEXTO} (Molina, 2017) (Castell, 2011)." in textos
    comentarios = [c.text for c in docx.Document(str(salida)).comments]
    assert len(comentarios) == 1 and "no está en la lista de referencias" in comentarios[0]


@pytest.mark.parametrize("quitado, nuevo, vale", [
    ("(Martínez 2019)", "(Martínez, 2019)", True),
    ("et. al", "et al.", True),
    ("s.f", "s. f.", True),
    ("n.d.", "s. f.", True),
    ("Pérez, L.; Romero, A.", "Pérez, L. y Romero, A.", True),
    ("doi:10.1/x", "https://doi.org/10.1/x", True),
    ("Madrid: ", "", True),
    ("Recuperado de ", "", True),
    (", nº 2", "(2)", True),
    ("(Martínez 2019)", "(Martín, 2019)", False),            # cambia una palabra
    ("Santa Cruz de Tenerife: ", "", True),                  # una ciudad: hasta cuatro palabras
    ("uno dos tres cuatro cinco", "", False),
    ("", "texto", False),
    ("igual", "igual", False),
])
def test_que_correcciones_acepta_la_integridad(quitado, nuevo, vale):
    assert correccion_de_forma_permitida(quitado, nuevo) is vale


def test_la_integridad_rechaza_una_correccion_no_permitida(tmp_path):
    doc = trabajo([f"{TEXTO} (Molina 2017)."], [MOLINA])
    antes = documento.radiografia(doc)
    doc.paragraphs[1].runs[0].text = f"{TEXTO} (Molina, 2018)."
    despues = documento.radiografia(doc)
    mal = CambioPrevisto(2, "corregir_forma", quitado="(Molina 2017)", nuevo="(Molina, 2018)", datos=0)
    assert documento.comparar(antes, despues, [mal])
    bien = CambioPrevisto(2, "corregir_forma", quitado="(Molina 2017)", nuevo="(Molina, 2017)", datos=0)
    doc.paragraphs[1].runs[0].text = f"{TEXTO} (Molina, 2017)."
    assert documento.comparar(antes, documento.radiografia(doc), [bien]) == []
