"""
Paso 5.4: el índice. APA no lo usa en los trabajos de estudiante (Order of pages),
pero muchas universidades lo exigen: por defecto se deja y se avisa; con
--quitar-indice se quita, sin dejar páginas en blanco ni perder secciones.

El 05 tiene el índice automático de Word (un campo TOC en un control de contenido
«Tabla de contenido», vacío: Word lo rehace al abrir). El índice escrito a mano se
prueba con documentos hechos aquí.
"""

import copy
import json
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK

from apa_facil import documento, pipeline
from apa_facil.__main__ import main
from apa_facil.documento import CambioPrevisto, comparar, es_parte_de_un_indice, radiografia, texto_parrafo, w
from herramientas import a_pdf
from herramientas.verificar import INDICE, verificar

RAIZ = Path(__file__).resolve().parent.parent
TFG = RAIZ / "pruebas" / "05_tfg_completo.docx"
TEXTO = ("La despoblación del medio rural es uno de los problemas más graves de la España interior, y el turismo "
         "rural se ha presentado como una de las respuestas posibles para frenarla en muchos municipios.")


@pytest.fixture(scope="module")
def con_y_sin(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("indice")
    # «con» es un nombre reservado en Windows: Word no abre lo que hay en una carpeta así.
    con, informe_con = pipeline.procesar(TFG, carpeta / "con_indice")
    sin, informe_sin = pipeline.procesar(TFG, carpeta / "sin_indice", {"quitar_indice": True})
    return (con, informe_con), (sin, informe_sin)


def parrafos(ruta):
    return list(docx.Document(str(ruta)).element.body.iter(w("p")))


def tiene_toc(ruta):
    return any("TOC" in (t.text or "") for t in docx.Document(str(ruta)).element.body.iter(w("instrText")))


# --- Qué puede ser parte de un índice -------------------------------------------------------------

@pytest.mark.parametrize("texto", ["", "  ", "\n", "Índice", "ÍNDICE:", "Tabla de contenido", "Contenidos",
                                   "1. Introducción\t3", "Método ........ 5", "Resultados…12",
                                   "Anexo A    25", "Resumen\tiv"])
def test_lineas_de_indice(texto):
    assert es_parte_de_un_indice(texto)


@pytest.mark.parametrize("texto", ["En 2020 crecimos 3", "Capítulo 2", TEXTO, "Introducción", "Método ...."])
def test_lineas_que_no_son_de_indice(texto):
    assert not es_parte_de_un_indice(texto)


# --- El índice del 05 ---------------------------------------------------------------------------------

def test_por_defecto_se_mantiene_y_se_avisa(con_y_sin):
    (con, informe), _ = con_y_sin
    assert tiene_toc(con)
    assert "Tu trabajo tiene un índice. APA no lo pide en los trabajos de estudiante, pero muchas universidades lo " \
           "exigen, así que se ha dejado. Si la tuya no lo pide, usa la opción --quitar-indice." in informe.avisos
    assert docx.Document(str(con)).settings.element.find(w("updateFields")) is not None
    assert [h.regla for h in verificar(str(con)).excepciones if h.regla == INDICE] == [INDICE]


def test_con_la_opcion_se_quita(con_y_sin):
    _, (sin, informe) = con_y_sin
    assert not tiene_toc(sin)
    assert next(docx.Document(str(sin)).element.body.iter(w("sdt")), None) is None, "sin el control de contenido"
    assert ("Índice: se ha quitado el índice (el índice automático de Word), como has pedido: APA no lo usa en los "
            "trabajos de estudiante. También se ha quitado el salto de página que había antes, para que no quede "
            "una página en blanco.") in informe.cambios
    assert not any("--quitar-indice" in a for a in informe.avisos)


def test_sin_perder_secciones_ni_dejar_paginas_en_blanco(con_y_sin):
    _, (sin, _) = con_y_sin
    assert radiografia(documento.abrir(sin)).secciones == radiografia(documento.abrir(TFG)).secciones == 3
    ps = parrafos(sin)
    claves = next(n for n, p in enumerate(ps) if texto_parrafo(p).startswith("Palabras clave"))
    # Detrás del resumen, solo el párrafo vacío con el salto de sección (sin el salto de página de antes).
    assert texto_parrafo(ps[claves + 1]) == "" and ps[claves + 1].find(f"{w('pPr')}/{w('sectPr')}") is not None
    assert texto_parrafo(ps[claves + 2]).startswith("Turismo Rural y Despoblación")


def test_ya_no_pregunta_por_los_campos(con_y_sin):
    _, (sin, informe) = con_y_sin
    assert docx.Document(str(sin)).settings.element.find(w("updateFields")) is None
    assert any(c.startswith("Índice: Word ya no preguntará al abrir el trabajo") for c in informe.cambios)


def test_la_copia_sin_indice_cumple(con_y_sin):
    _, (sin, _) = con_y_sin
    resultado = verificar(str(sin))
    assert resultado.incumplimientos == []
    assert not [h for h in resultado.excepciones if h.regla == INDICE]


def test_json_de_la_opcion(con_y_sin):
    _, (sin, _) = con_y_sin
    assert json.loads(sin.with_suffix(".json").read_text(encoding="utf-8"))["opciones"]["quitar_indice"] is True


# --- Índice escrito a mano -----------------------------------------------------------------------------

def a_mano(lineas, tmp_path):
    doc = docx.Document()
    doc.add_paragraph("Índice")
    for linea in lineas:
        p = doc.add_paragraph(linea)
    p.runs[-1].add_break(WD_BREAK.PAGE)   # el índice acaba en un salto de página
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(TEXTO)
    doc.add_heading("Método", level=1)
    doc.add_paragraph(TEXTO)
    ruta = tmp_path / "a_mano.docx"
    doc.save(str(ruta))
    return ruta


def test_indice_a_mano_se_quita(tmp_path):
    ruta = a_mano(["Introducción ........ 1", "Método ........ 2"], tmp_path)
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", {"quitar_indice": True})
    textos = [texto_parrafo(p) for p in parrafos(salida)]
    assert "Índice" not in textos and not any(t.startswith("Método ....") for t in textos)
    assert "Introducción" in textos and "Método" in textos, "los títulos no se tocan"
    assert any(c.startswith("Índice: se ha quitado el índice: 3 líneas («Índice»…)") for c in informe.cambios)


def test_indice_a_mano_con_algo_que_no_es_del_indice(tmp_path):
    ruta = a_mano(["Introducción ........ 1", "Este párrafo explica cómo leer el trabajo y no es del índice."],
                  tmp_path)
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", {"quitar_indice": True})
    assert "Índice" in [texto_parrafo(p) for p in parrafos(salida)]
    assert any(a.startswith("No se ha quitado el índice porque hay líneas que no parecen de un índice") for a in
               informe.avisos)


# --- Integridad ------------------------------------------------------------------------------------------

def test_integridad_solo_quita_lineas_de_indice(tmp_path):
    ruta = a_mano(["Introducción ........ 1", "Método ........ 2"], tmp_path)
    antes = radiografia(documento.abrir(ruta))
    despues = copy.deepcopy(antes)
    del despues.textos[0:3]
    assert comparar(antes, despues, [CambioPrevisto(n, "quitar_indice") for n in (1, 2, 3)]) == []
    despues = copy.deepcopy(antes)
    del despues.textos[3:5]   # «Introducción» y su texto
    assert comparar(antes, despues, [CambioPrevisto(n, "quitar_indice") for n in (4, 5)])


def test_integridad_el_indice_entero_y_seguido(tmp_path):
    ruta = a_mano(["Introducción ........ 1", "Método ........ 2"], tmp_path)
    antes = radiografia(documento.abrir(ruta))
    despues = copy.deepcopy(antes)
    del despues.textos[2]
    del despues.textos[0]
    assert comparar(antes, despues, [CambioPrevisto(n, "quitar_indice") for n in (1, 3)])


# --- Línea de comandos y Word ---------------------------------------------------------------------------

def test_linea_de_comandos(tmp_path, capsys):
    assert main([str(TFG), "--quitar-indice"], carpeta_salida=tmp_path) == 0
    assert "Índice: se ha quitado el índice" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        main(["-h"])
    assert "--quitar-indice" in capsys.readouterr().out


@pytest.mark.skipif(not a_pdf.hay_word(), reason="Word no está instalado")
def test_word_sin_indice_no_deja_paginas_en_blanco(con_y_sin, tmp_path):
    # El índice del 05 ocupa su propia página: sin él, el PDF tiene una página menos. Si quedara el salto de
    # página de antes del índice, su página seguiría ahí, en blanco, y tendría las mismas.
    (con, _), (sin, _) = con_y_sin
    # Las dos copias se llaman igual: cada PDF, en su carpeta.
    (_, [con_pdf]), (_, [sin_pdf]) = a_pdf.convertir([con], tmp_path / "pdf_con"), a_pdf.convertir([sin], tmp_path / "pdf_sin")
    assert con_pdf.error is None and sin_pdf.error is None, (con_pdf.error, sin_pdf.error)
    assert sin_pdf.paginas == con_pdf.paginas - 1
