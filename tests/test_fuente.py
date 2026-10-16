"""
Paso 1.2: una sola fuente para todo el documento, con su tamaño.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

from collections import Counter
from pathlib import Path

import docx
import pytest
from docx.oxml import OxmlElement
from docx.shared import Pt

from apa_facil import pipeline
from apa_facil.documento import Estilos, w
from apa_facil.informe import Informe
from apa_facil.reglas import fuente

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))


def aplicar(doc, **opciones):
    informe = Informe()
    fuente.aplicar(doc, opciones, informe)
    return informe


def documento(*fragmentos):
    """Un documento con un párrafo por fragmento: (texto, fuente, tamaño)."""
    doc = docx.Document()
    for texto, nombre, tamano in fragmentos:
        run = doc.add_paragraph().add_run(texto)
        run.font.name, run.font.size = nombre, Pt(tamano)
    return doc


def fuentes_que_se_ven(doc):
    """(fuente, tamaño) → letras, con el formato efectivo de cada fragmento del cuerpo."""
    estilos, cuenta = Estilos(doc), Counter()
    for p in doc.element.body.iter(w("p")):
        for r in p.iter(w("r")):
            letras = sum(len(t.text or "") for t in r.iter(w("t")))
            if letras:
                cuenta[(estilos.fuente(r, p), estilos.tamano(r, p))] += letras
    return cuenta


def texto_largo(n=200):
    return "palabra " * (n // 8)


# --- Elegir la fuente ---------------------------------------------------------------

@pytest.mark.parametrize("nombre, esperada", [
    ("Calibri", "Calibri"), ("Calibri Light", "Calibri"), ("Arial Black", "Arial"),
    ("Aptos Display", "Aptos"), ("Times New Roman", "Times New Roman"),
    ("CMU Serif", "Computer Modern"), ("Latin Modern Roman", "Computer Modern"),
    ("Lucida Sans Unicode", "Lucida Sans Unicode"), ("Verdana", None), ("Lucida Sans", None),
])
def test_familia(nombre, esperada):
    assert fuente.familia(nombre) == esperada


@pytest.mark.parametrize("nombre, esperada, conocida", [
    ("Verdana", "Arial", True), ("Segoe UI", "Calibri", True), ("Comic Sans MS", "Calibri", True),
    ("Cambria", "Times New Roman", True), ("Constantia", "Georgia", True),
    ("Lucida Sans", "Lucida Sans Unicode", True), ("Courier New", "Times New Roman", True),
    ("Fuente Rara Sans", "Arial", False), ("Fuente Rarísima", "Times New Roman", False),
])
def test_mas_parecida(nombre, esperada, conocida):
    assert fuente.mas_parecida(nombre) == (esperada, conocida)


def test_recomendada_se_mantiene_y_cambia_el_tamano():
    doc = documento((texto_largo(), "Georgia", 12), ("Título", "Georgia", 16))
    informe = aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Georgia", 11)}
    assert informe.cambios == ["Fuente: se mantiene Georgia, recomendada por APA, y todo el texto "
                               "pasa a 11 pt (antes había texto a 12 y 16 pt)."]


def test_aptos_es_recomendada():
    doc = documento((texto_largo(), "Aptos", 11))
    aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Aptos", 12)}


def test_no_recomendada_se_cambia_por_la_mas_parecida():
    doc = documento((texto_largo(), "Verdana", 9))
    informe = aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Arial", 11)}
    assert informe.cambios == ["Fuente: Verdana se ha cambiado por Arial 11 pt, la fuente "
                               "recomendada por APA más parecida."]
    assert informe.avisos == []


def test_fuente_desconocida_avisa():
    doc = documento((texto_largo(), "Fuente Rarísima", 12))
    informe = aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Times New Roman", 12)}
    assert len(informe.avisos) == 1 and "No conocemos la fuente «Fuente Rarísima»" in informe.avisos[0]


def test_mezcla_se_unifica_en_la_principal():
    doc = documento((texto_largo(400), "Calibri", 11), (texto_largo(100), "Verdana", 8),
                    ("Título", "Arial Black", 14))
    informe = aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Calibri", 11)}
    assert informe.cambios == ["Tu trabajo mezclaba 3 fuentes (Calibri, Verdana y Arial). Se ha unificado "
                               "todo en Calibri 11 pt, la fuente recomendada por APA que más usabas."]


def test_principal_no_recomendada_en_una_mezcla():
    doc = documento((texto_largo(400), "Segoe UI", 10), (texto_largo(100), "Georgia", 12))
    informe = aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Calibri", 11)}
    assert "la que más usabas, Segoe UI, se ha cambiado por la fuente recomendada" in informe.cambios[0]


def test_si_ya_cumple_no_hay_cambio():
    doc = documento((texto_largo(), "Times New Roman", 12))
    assert aplicar(doc).cambios == []


def test_la_opcion_fuente_manda():
    doc = documento((texto_largo(), "Georgia", 11))
    informe = aplicar(doc, fuente="Arial")
    assert set(fuentes_que_se_ven(doc)) == {("Arial", 11)}
    assert informe.cambios == ["Fuente: se ha aplicado Arial 11 pt en todo el documento, como has pedido."]


def test_opcion_fuente_no_recomendada_da_error_y_elige_sola():
    doc = documento((texto_largo(), "Georgia", 12))
    informe = aplicar(doc, fuente="Papyrus")
    assert "«Papyrus» no es una de las fuentes recomendadas" in informe.errores[0]
    assert set(fuentes_que_se_ven(doc)) == {("Georgia", 11)}


def test_computer_modern_conserva_su_nombre_en_word():
    doc = documento((texto_largo(), "CMU Serif", 12))
    aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("CMU Serif", 10)}


# --- Aplicarla en todas partes --------------------------------------------------------

def test_quita_la_fuente_del_tema():
    """El documento en blanco de python-docx usa las fuentes del tema (Calibri)."""
    doc = docx.Document()
    doc.add_heading("Un título", level=1)   # el estilo de título usa la fuente "major" del tema
    doc.add_paragraph(texto_largo())
    aplicar(doc, fuente="Georgia")
    estilos = doc.styles.element
    for atributo in ("asciiTheme", "hAnsiTheme"):
        assert not [f for f in estilos.iter(w("rFonts")) if f.get(w(atributo))]
    assert set(fuentes_que_se_ven(doc)) == {("Georgia", 11)}


def test_tablas_misma_fuente_y_tamano():
    doc = documento((texto_largo(), "Arial", 12))
    celda = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    run = celda.paragraphs[0].add_run("dato")
    run.font.name, run.font.size = "Verdana", Pt(8)
    aplicar(doc)
    assert set(fuentes_que_se_ven(doc)) == {("Arial", 11)}


def test_encabezado_y_pie():
    doc = documento((texto_largo(), "Georgia", 12))
    run = doc.sections[0].header.paragraphs[0].add_run("1")
    run.font.name, run.font.size = "Verdana", Pt(9)
    aplicar(doc)
    rpr = run._r.rPr
    assert rpr.rFonts.get(w("ascii")) == "Georgia" and rpr.sz.val == Pt(11)


def test_simbolos_no_se_tocan():
    doc = documento((texto_largo(), "Georgia", 12), ("", "Georgia", 12))
    simbolo = doc.paragraphs[-1].add_run("")
    simbolo.font.name = "Symbol"
    aplicar(doc)
    assert simbolo.font.name == "Symbol"


def test_ecuaciones_no_se_tocan():
    doc = documento((texto_largo(), "Georgia", 12))
    m = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    ecuacion = OxmlElement("m:oMath")
    run = OxmlElement("w:r")
    rpr = run.get_or_add_rPr()
    rpr.get_or_add_rFonts().set(w("ascii"), "Cambria Math")
    ecuacion.append(run)
    doc.paragraphs[0]._p.append(ecuacion)
    aplicar(doc)
    assert ecuacion.tag == f"{{{m}}}oMath"
    assert run.rPr.rFonts.get(w("ascii")) == "Cambria Math"


# --- Código de programación ---------------------------------------------------------------

@pytest.mark.parametrize("texto", [
    "def media(datos):",
    "print(resultado)",
    "for (int i = 0; i < n; i++) {",
    "x = 5",
    "if edad >= 18:",
    "// comentario",
])
def test_codigo_se_respeta(texto):
    doc = documento((texto_largo(), "Georgia", 12), (texto, "Courier New", 9))
    informe = aplicar(doc)
    codigo = doc.paragraphs[-1].runs[0]
    assert codigo.font.name == "Courier New" and codigo.font.size == Pt(10)
    assert len(informe.avisos) == 1 and "parece código de programación" in informe.avisos[0]


@pytest.mark.parametrize("texto", [
    "La tabla de la página siguiente resume los datos.",
    "Pérez, L.; Romero, A. (2020) Autoestima digital. https://www.example.org/datos?id=3",
    "Webgrafía:",
])
def test_courier_que_no_es_codigo_se_unifica(texto):
    doc = documento((texto_largo(), "Georgia", 12), (texto, "Courier New", 10))
    informe = aplicar(doc)
    assert doc.paragraphs[-1].runs[0].font.name == "Georgia"
    assert informe.avisos == []


def test_estilo_de_codigo_se_respeta():
    doc = documento((texto_largo(), "Georgia", 12))
    estilo = doc.styles.add_style("Code", 1)
    estilo.font.name = "Consolas"
    p = doc.add_paragraph("resultado", style="Code")
    aplicar(doc)
    assert doc.styles["Code"].font.name == "Consolas"
    assert p.runs[0].font.size == Pt(10)


# --- Notas al pie y documentos de prueba ------------------------------------------------------

def test_notas_al_pie_conservan_su_tamano(tmp_path):
    original = next(p for p in PRUEBAS if p.name.startswith("05"))
    salida, _ = pipeline.procesar(original, tmp_path)
    copia = docx.Document(str(salida))
    notas = next(p for p in copia.part.package.iter_parts() if str(p.partname) == "/word/footnotes.xml")
    xml = notas.blob.decode("utf-8")
    assert 'w:sz w:val="20"' in xml, "las notas del 05 estaban a 10 pt y deben seguir así"
    assert "Theme" not in xml


ESPERADO = {
    "01": ("Arial", 11),
    "02": ("Calibri", 11),
    "03": ("Georgia", 11),
    "04": ("Calibri", 11),
    "05": ("Times New Roman", 12),
}


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    salida, informe = pipeline.procesar(original, tmp_path)
    assert set(fuentes_que_se_ven(docx.Document(str(salida)))) == {ESPERADO[original.name[:2]]}
    assert any("fuente" in c.lower() for c in informe.cambios)
    assert informe.errores == []
    assert not any("parece código" in a for a in informe.avisos), "ningún documento de prueba tiene código"
