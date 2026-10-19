"""
Paso 6.5: el orden alfabético de las referencias, según la web oficial de APA
(Reference list setup: por la primera palabra, normalmente el apellido del primer
autor; sin autor, por el título). Por defecto se reordenan (decisión del
propietario) y --orden-tal-cual deja el orden con aviso. La integridad solo admite
cambiar de sitio los párrafos de la lista, sin cambiar ninguna letra.
"""

from pathlib import Path

import docx
import pytest

from apa_facil import estructura, pipeline
from apa_facil.documento import CambioPrevisto, comparar, radiografia, texto_parrafo
from apa_facil.informe import Informe
from apa_facil.reglas import orden as regla
from apa_facil.reglas.referencias import listas
from herramientas.verificar import ORDEN, verificar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {p.name[:2]: p for p in (RAIZ / "pruebas").glob("*.docx")}
TEXTO = ("El sueño cumple funciones esenciales en la consolidación de la memoria y en la regulación emocional, según "
         "los estudios revisados (Navarro y Gil, 2020).")


def trabajo(*obras, rotulo="Referencias"):
    doc = docx.Document()
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(TEXTO)
    doc.add_heading(rotulo, level=1)
    for obra in obras:
        doc.add_paragraph(obra)
    return doc


def lista(doc):
    return [texto_parrafo(c.elemento) for c in listas(estructura.clasificar(doc))[0][1]]


def procesar(doc, tmp_path, opciones=None):
    ruta = tmp_path / "trabajo.docx"
    doc.save(str(ruta))
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", opciones)
    return docx.Document(str(salida)), informe


def ordenadas(*obras):
    return sorted(obras, key=regla.clave)


# --- Los documentos de prueba ---------------------------------------------------------------------

@pytest.fixture(scope="module")
def copias(tmp_path_factory):
    carpeta = tmp_path_factory.mktemp("orden")
    return {clave: pipeline.procesar(ruta, carpeta) for clave, ruta in PRUEBAS.items()}


@pytest.mark.parametrize("clave, primeras_palabras, cambia", [
    ("01", ["Lefèvre", "Ortega", "Ramírez", "https://www.historiaejemplo.es/revolucion-industrial-trabajo"], True),
    ("02", ["Castro", "Navarro"], False),
    ("03", ["Campos", "Herrero"], True),
    ("04", ["Martínez", "Pérez", "www.example.org/datos-uso-redes-2024"], False),
    ("05", ["Collantes", "Consejo", "Herranz", "Molina", "Sáenz", "Torres"], True),
])
def test_las_listas_de_los_cinco(copias, clave, primeras_palabras, cambia):
    salida, informe = copias[clave]
    doc = docx.Document(str(salida))
    assert [t.split(",")[0].split()[0] for t in lista(doc)] == primeras_palabras
    assert any(c.startswith("Referencias: se han ordenado alfabéticamente") for c in informe.cambios) == cambia
    assert [h for h in verificar(doc).incumplimientos if h.regla == ORDEN] == []
    original = [h for h in verificar(str(PRUEBAS[clave])).incumplimientos if h.regla == ORDEN]
    assert bool(original) == cambia


def test_la_url_suelta_del_01_se_avisa(copias):
    _, informe = copias["01"]
    assert any(a.startswith("«https://www.historiaejemplo.es/revoluci…» es solo una dirección web") for a in
               informe.avisos)


# --- La clave de orden ------------------------------------------------------------------------------

def test_nada_va_antes_que_algo():
    assert ordenadas("Browning, A. (2010). Uno.", "Brown, J. (2015). Dos.") == [
        "Brown, J. (2015). Dos.", "Browning, A. (2010). Uno."]


def test_un_solo_autor_antes_que_con_otros():
    assert ordenadas("Castro, P. y Gil, S. (2015). Uno.", "Castro, P. (2020). Dos.") == [
        "Castro, P. (2020). Dos.", "Castro, P. y Gil, S. (2015). Uno."]


def test_mismos_autores_por_año_y_sin_fecha_primero():
    assert ordenadas("Molina, R. (2019). B.", "Molina, R. (s. f.). C.", "Molina, R. (2017). A.") == [
        "Molina, R. (s. f.). C.", "Molina, R. (2017). A.", "Molina, R. (2019). B."]


def test_mismo_año_con_letra():
    assert ordenadas("Gil, S. (2020b). Zeta.", "Gil, S. (2020a). Alfa.") == [
        "Gil, S. (2020a). Alfa.", "Gil, S. (2020b). Zeta."]


def test_tildes_y_eñe():
    assert ordenadas("Álvarez, M. (2020). Uno.", "Alonso, P. (2020). Dos.") == [
        "Alonso, P. (2020). Dos.", "Álvarez, M. (2020). Uno."]
    assert ordenadas("Lorenzo, P. (2020). Dos.", "Álvarez, M. (2020). Uno.") == [
        "Álvarez, M. (2020). Uno.", "Lorenzo, P. (2020). Dos."], "la Á va con la A"
    assert ordenadas("Peña, L. (2020). Uno.", "Penas, R. (2020). Dos.", "Pereira, A. (2020). Tres.") == [
        "Penas, R. (2020). Dos.", "Peña, L. (2020). Uno.", "Pereira, A. (2020). Tres."]


def test_sin_autor_por_el_titulo():
    assert ordenadas("Martín, A. (2020). Uno.", "Manual de estilo (2019). Editorial.") == [
        "Manual de estilo (2019). Editorial.", "Martín, A. (2020). Uno."]


# --- Casos hechos a mano ------------------------------------------------------------------------------

OBRAS = ["Torres, E. (2019). Agroturismo. Editorial Ejemplo.", "Collantes, F. (2011). Lugares. Prensas.",
         "Molina, R. (2017). Turismo y territorio. Editorial Ejemplo."]


def test_con_orden_tal_cual(tmp_path):
    doc, informe = procesar(trabajo(*OBRAS), tmp_path, {"orden_tal_cual": True})
    assert lista(doc) == OBRAS
    assert any(a.startswith("Las referencias no están en orden alfabético") for a in informe.avisos)
    assert [h.regla for h in verificar(doc).incumplimientos if h.regla == ORDEN] == [ORDEN]


def test_varias_listas_no_se_ordenan(tmp_path):
    original = trabajo(*OBRAS, rotulo="Bibliografía")
    original.add_heading("Webgrafía", level=1)
    original.add_paragraph("INE (2024). Padrón municipal. https://www.ine.es")
    doc, informe = procesar(original, tmp_path)
    assert lista(doc) == OBRAS
    assert any(a.startswith("Como hay 2 listas de referencias, no se han ordenado") for a in informe.avisos)


def test_mismo_autor_y_año_sin_letra():
    informe = Informe()
    regla.aplicar(trabajo("Gil, S. (2020). Alfa.", "Gil, S. (2020). Beta."), {}, informe)
    assert any(a.startswith("Hay obras de los mismos autores y el mismo año") for a in informe.avisos)


def test_el_formato_va_con_su_referencia(tmp_path):
    original = trabajo(*OBRAS)
    original.paragraphs[-1].runs[0].italic = True   # «Molina…», en cursiva
    doc, _ = procesar(original, tmp_path)
    molina = next(p for p in doc.paragraphs if p.text.startswith("Molina"))
    assert molina.runs[0].italic


# --- Integridad --------------------------------------------------------------------------------------

def _radiografias(antes, despues):
    def radio(textos):
        doc = docx.Document()
        for texto in textos:
            if texto == "TABLA":
                doc.add_table(rows=1, cols=1).cell(0, 0).text = "dato"
            else:
                doc.add_paragraph(texto)
        return radiografia(doc)
    return radio(antes), radio(despues)


def test_integridad_admite_reordenar():
    a, d = _radiografias(["Referencias", "B", "", "A"], ["Referencias", "A", "", "B"])
    assert comparar(a, d, [CambioPrevisto(2, "reordenar_referencias", desde=(4, 2))]) == []


@pytest.mark.parametrize("antes, despues, cambio", [
    # con texto en medio que no se reordena: no es una lista seguida
    (["B", "Texto.", "A"], ["A", "Texto.", "B"], CambioPrevisto(1, "reordenar_referencias", desde=(3, 1))),
    # un párrafo vacío no es una referencia
    (["B", ""], ["", "B"], CambioPrevisto(1, "reordenar_referencias", desde=(2, 1))),
    # el texto cambia, además del orden
    (["B", "A"], ["A", "C"], CambioPrevisto(1, "reordenar_referencias", desde=(2, 1))),
    # un párrafo de una tabla
    (["B", "TABLA"], ["TABLA", "B"], CambioPrevisto(1, "reordenar_referencias", desde=(2, 1))),
])
def test_integridad_rechaza_lo_demas(antes, despues, cambio):
    a, d = _radiografias(antes, despues)
    assert comparar(a, d, [cambio])
