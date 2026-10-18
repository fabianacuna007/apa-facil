"""
Paso 5.5: el resumen y las palabras clave, según la guía oficial de APA (Abstract
and Keywords Guide): el rótulo «Palabras clave:» en cursiva y sin negrita, las
palabras sin cursiva ni negrita y sin punto al final, y lo que sigue al resumen en
una página nueva. Lo demás (más de 250 palabras, varios párrafos, número de
palabras clave, mayúsculas) solo se avisa.
"""

import copy
from pathlib import Path

import docx
import pytest

from apa_facil import documento, pipeline
from apa_facil.documento import CambioPrevisto, Estilos, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import resumen as regla
from apa_facil.reglas.fuente import _runs_propios
from herramientas import a_pdf
from herramientas.verificar import RESUMEN, verificar
from test_portada_con_word import posiciones_segun_word

RAIZ = Path(__file__).resolve().parent.parent
TFG, INFORME = RAIZ / "pruebas" / "05_tfg_completo.docx", RAIZ / "pruebas" / "02_informe_tablas.docx"
TEXTO = ("Este trabajo analiza el turismo rural en la España interior y su papel frente a la despoblación, a partir "
         "de cuatro municipios y de entrevistas con empresarios del sector.")


def trabajo(claves, resumen=(TEXTO,), estilo_clave=None):
    """Un trabajo con resumen (y, si se dan, palabras clave), y un apartado detrás."""
    doc = docx.Document()
    doc.add_heading("Resumen", level=1)
    for parrafo in resumen:
        doc.add_paragraph(parrafo)
    if claves is not None:
        p = doc.add_paragraph()
        r = p.add_run(claves)
        if estilo_clave == "cursiva":
            r.italic = True
        elif estilo_clave == "negrita":
            r.bold = True
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(TEXTO + " La España interior pierde población desde hace décadas.")
    return doc


def aplicar(doc):
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    return informe


def parrafo_de_claves(doc):
    return next(p for p in doc.element.body.iter(w("p")) if texto_parrafo(p).startswith(("Palabras", "Keywords")))


def formato(doc, p):
    """[(texto, cursiva, negrita)] de cada fragmento con texto."""
    estilos = Estilos(doc)
    return [("".join(t.text or "" for t in r.iter(w("t"))), regla._activado(estilos, r, p, "i"),
             regla._activado(estilos, r, p, "b")) for r in _runs_propios(p)
            if "".join(t.text or "" for t in r.iter(w("t")))]


# --- El 05 y el 02 -------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def tfg(tmp_path_factory):
    return pipeline.procesar(TFG, tmp_path_factory.mktemp("tfg"))


def test_palabras_clave_del_05(tfg):
    salida, informe = tfg
    doc = docx.Document(str(salida))
    p = parrafo_de_claves(doc)
    assert texto_parrafo(p) == "Palabras clave: turismo rural, despoblación, desarrollo local, España interior"
    assert formato(doc, p) == [("Palabras clave:", True, False), (" ", False, False),
                               ("turismo rural, despoblación, desarrollo local, España interior", False, False)]
    assert "Resumen: se ha quitado el punto del final de las palabras clave, como pide APA." in informe.cambios
    assert "Resumen: el rótulo «Palabras clave:» va ahora en cursiva y sin negrita, y las palabras clave sin " \
           "cursiva ni negrita, como pide APA." in informe.cambios
    assert not any("palabras clave empiezan con mayúscula" in a for a in informe.avisos), "«España» es nombre propio"


def test_el_05_cumple_y_el_original_no(tfg):
    salida, _ = tfg
    assert [h for h in verificar(str(salida)).incumplimientos if h.regla == RESUMEN] == []
    assert [h.mensaje for h in verificar(str(TFG)).incumplimientos if h.regla == RESUMEN] == [
        "Palabras clave: el rótulo tiene que ir en cursiva y sin negrita.", "Palabras clave: terminan en punto."]


def test_el_texto_del_02_empieza_en_pagina_nueva(tmp_path):
    salida, informe = pipeline.procesar(INFORME, tmp_path)
    assert "Resumen: lo que sigue al resumen empieza ahora en una página nueva, como pide APA." in informe.cambios
    ps = list(docx.Document(str(salida)).element.body.iter(w("p")))
    resumen = next(n for n, p in enumerate(ps) if texto_parrafo(p).startswith("Este informe analiza"))
    assert ps[resumen + 1].find(f"{w('pPr')}/{w('pageBreakBefore')}") is not None
    assert [h.mensaje for h in verificar(str(INFORME)).incumplimientos if h.regla == RESUMEN] == [
        "Lo que sigue al resumen no empieza en una página nueva."]
    assert [h for h in verificar(str(salida)).incumplimientos if h.regla == RESUMEN] == []


# --- Casos hechos a mano ------------------------------------------------------------------------------

def test_rotulo_y_palabras_en_un_solo_fragmento_en_cursiva():
    doc = trabajo("Palabras clave: turismo, despoblación, empleo", estilo_clave="cursiva")
    antes = radiografia(doc)
    aplicar(doc)
    p = parrafo_de_claves(doc)
    assert formato(doc, p) == [("Palabras clave:", True, False), (" turismo, despoblación, empleo", False, False)]
    assert comparar(antes, radiografia(doc)) == [], "el texto no cambia al partir el fragmento"


def test_si_ya_estaba_bien_no_cambia_nada():
    doc = trabajo("Palabras clave:")
    p = doc.paragraphs[2]
    p.runs[0].italic = True
    p.add_run(" turismo, despoblación, empleo")
    informe = aplicar(doc)
    assert not any(c.startswith("Resumen: el rótulo") or "punto" in c for c in informe.cambios)


def test_punto_de_una_abreviatura_se_deja():
    doc = trabajo("Palabras clave: turismo, despoblación, EE. UU.")
    informe = aplicar(doc)
    assert texto_parrafo(parrafo_de_claves(doc)).endswith("EE. UU.")
    assert any(a.startswith("Las palabras clave terminan en un punto que parece de una abreviatura") for a in
               informe.avisos)
    assert [h.regla for h in verificar(doc).excepciones if h.regla == RESUMEN] == [RESUMEN]


def test_sin_dos_puntos_se_avisa():
    informe = aplicar(trabajo("Palabras clave turismo, despoblación, empleo"))
    assert any(a.startswith("El rótulo de las palabras clave no lleva dos puntos") for a in informe.avisos)


@pytest.mark.parametrize("claves, cuantas", [("Palabras clave: turismo, empleo", 2),
                                             ("Palabras clave: a, b, c, d, e, f", 6)])
def test_numero_de_palabras_clave(claves, cuantas):
    informe = aplicar(trabajo(claves))
    assert f"Hay {cuantas} palabras clave: APA recomienda de 3 a 5." in informe.avisos


def test_mayusculas_en_las_palabras_clave():
    informe = aplicar(trabajo("Palabras clave: Turismo, Despoblación, España interior"))
    aviso = next(a for a in informe.avisos if a.startswith("Algunas palabras clave empiezan con mayúscula"))
    assert "«Turismo», «Despoblación»" in aviso and "España" not in aviso, "«España» es un nombre propio del texto"


def test_resumen_largo_y_en_varios_parrafos():
    informe = aplicar(trabajo(None, resumen=(TEXTO * 6, TEXTO * 6)))
    assert "El resumen tiene 2 párrafos: APA lo escribe en uno solo. Únelos tú." in informe.avisos
    assert any(a.startswith("El resumen tiene ") and "máximo de 250" in a for a in informe.avisos)


def test_sin_resumen_no_hace_nada():
    doc = docx.Document()
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(TEXTO)
    informe = aplicar(doc)
    assert informe.entradas == []


# --- Integridad --------------------------------------------------------------------------------------

def _con_punto_quitado(doc, indice):
    antes = radiografia(doc)
    despues = copy.deepcopy(antes)
    donde, texto = despues.textos[indice]
    despues.textos[indice] = (donde, texto[:-1])
    return antes, despues


def test_integridad_solo_el_punto_de_las_palabras_clave():
    doc = trabajo("Palabras clave: turismo, despoblación, empleo.")
    antes, despues = _con_punto_quitado(doc, 2)   # 1 «Resumen», 2 el resumen, 3 las palabras clave
    assert comparar(antes, despues, [CambioPrevisto(3, "quitar_punto_final", quitado=".")]) == []
    antes, despues = _con_punto_quitado(doc, 1)   # el punto del resumen
    assert comparar(antes, despues, [CambioPrevisto(2, "quitar_punto_final", quitado=".")])


def test_integridad_solo_un_punto():
    doc = trabajo("Palabras clave: turismo, despoblación, empleo;")
    antes, despues = _con_punto_quitado(doc, 2)
    assert comparar(antes, despues, [CambioPrevisto(3, "quitar_punto_final", quitado=";")])


# --- Word ---------------------------------------------------------------------------------------------

@pytest.mark.skipif(not a_pdf.hay_word(), reason="Word no está instalado")
def test_word_el_texto_del_02_empieza_en_la_pagina_siguiente(tmp_path):
    salida, _ = pipeline.procesar(INFORME, tmp_path / "salida_02")
    ps = list(docx.Document(str(salida)).element.body.iter(w("p")))
    resumen = next(n for n, p in enumerate(ps, start=1) if texto_parrafo(p) == "Resumen")
    filas = posiciones_segun_word(salida, resumen + 2)
    assert filas[resumen - 1][0] == 2, "el resumen, en la página 2"
    assert filas[resumen + 1] == (3, pytest.approx(72, abs=1)), "lo que sigue, arriba de la página 3"
