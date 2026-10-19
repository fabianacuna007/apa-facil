"""
Paso 6.3: las citas en bloque, según la web oficial de APA (Quotations): todo el
bloque sangrado 1,27 cm desde el margen izquierdo, sin sangría a la derecha ni en
la primera línea (en los párrafos segundo y siguientes, 1,27 cm más), a doble
espacio y sin la cursiva que va en todo el texto. Las comillas, el punto detrás de
la fuente, los párrafos con sangría de menos de 40 palabras y las citas largas
dentro de un párrafo solo se avisan: el texto no cambia.
"""

from pathlib import Path

import docx
import pytest
from docx.shared import Cm

from apa_facil import pipeline
from apa_facil.documento import Estilos, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import citas as regla
from apa_facil.reglas import titulos
from apa_facil.reglas.resumen import _poner
from herramientas.verificar import CITAS, SANGRIA, verificar

RAIZ = Path(__file__).resolve().parent.parent
TFG = RAIZ / "pruebas" / "05_tfg_completo.docx"
TEXTO = ("La despoblación de la España interior es un proceso largo que empezó con el éxodo rural de mediados del "
         "siglo pasado y que todavía no se ha detenido en muchas comarcas.")
CITA = ("El turismo rural puede complementar las rentas agrarias y crear oportunidades para las mujeres y los "
        "jóvenes, pero difícilmente sustituirá a otros sectores como base económica de un territorio. Sin servicios "
        "públicos, transporte y acceso a internet de calidad, los visitantes llegan, pero los nuevos residentes no se "
        "quedan.")
FUENTE = " (Consejo Económico de Ejemplo, 2022, p. 45)"


def trabajo(*citas, cursiva=False, derecha=False):
    """Un apartado con texto y, detrás, los párrafos `citas` con sangría a la izquierda de 1 cm."""
    doc = docx.Document()
    doc.add_heading("Marco teórico", level=1)
    doc.add_paragraph(TEXTO)
    for texto in citas:
        p = doc.add_paragraph()
        p.add_run(texto).italic = cursiva or None
        p.paragraph_format.left_indent = Cm(1)
        if derecha:
            p.paragraph_format.right_indent = Cm(1)
    doc.add_paragraph(TEXTO)
    return doc


def aplicar(doc):
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    return informe


def parrafo(doc, inicio):
    return next(p for p in doc.element.body.iter(w("p")) if texto_parrafo(p).startswith(inicio))


def sangrias(doc, p):
    estilos = Estilos(doc)
    return {a: estilos.parrafo(p, "ind", a) for a in ("left", "right", "firstLine")}


def incumplimientos(doc):
    """Los de 6.3 y los de sangría (1.5) en los párrafos de la cita (en los casos hechos a mano, el resto del texto
    no ha pasado por la regla de sangría)."""
    return [h.mensaje for h in verificar(doc).incumplimientos
            if h.regla == CITAS or (h.regla == SANGRIA and not h.extracto.startswith(TEXTO[:30]))]


# --- El 05 -----------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def tfg(tmp_path_factory):
    return pipeline.procesar(TFG, tmp_path_factory.mktemp("citas"))


def test_la_cita_del_05(tfg):
    salida, informe = tfg
    doc = docx.Document(str(salida))
    p = parrafo(doc, '"El turismo rural')
    assert sangrias(doc, p) == {"left": "720", "right": None, "firstLine": "0"}
    estilos = Estilos(doc)
    assert not any(titulos.cursiva_efectiva(estilos, r, p) for r in titulos.fragmentos_con_letras(p))
    assert any(c.startswith("Citas en bloque: 1 cita") and "1,27 cm" in c for c in informe.cambios)
    assert any(c.startswith("Citas en bloque: se ha quitado la cursiva de 1 cita") for c in informe.cambios)
    assert any(a.startswith("1 cita en bloque va entre comillas") for a in informe.avisos)
    assert texto_parrafo(p).startswith('"El turismo rural'), "las comillas se quedan"


def test_la_copia_del_05_cumple_y_el_original_no(tfg):
    salida, _ = tfg
    assert incumplimientos(docx.Document(str(salida))) == []
    assert [h.mensaje for h in verificar(str(TFG)).incumplimientos if h.regla == CITAS] == [
        "Cita en bloque: todo el bloque no está sangrado 1,27 cm a la izquierda.",
        "Cita en bloque: tiene sangría a la derecha.", "Cita en bloque: va toda en cursiva."]


# --- Casos hechos a mano ----------------------------------------------------------------------------

def test_cita_de_varios_parrafos():
    mitad = CITA.split(". ")
    doc = trabajo(mitad[0] + ".", ". ".join(mitad[1:]) + FUENTE)
    antes = radiografia(doc)
    aplicar(doc)
    primero, segundo = parrafo(doc, "El turismo"), parrafo(doc, "Sin servicios")
    assert sangrias(doc, primero) == {"left": "720", "right": None, "firstLine": "0"}
    assert sangrias(doc, segundo) == {"left": "720", "right": None, "firstLine": "720"}
    assert comparar(antes, radiografia(doc)) == [], "el texto no cambia"
    assert incumplimientos(doc) == []


def test_sangria_derecha_se_quita():
    doc = trabajo(CITA + FUENTE, derecha=True)
    aplicar(doc)
    assert sangrias(doc, parrafo(doc, "El turismo"))["right"] is None


def test_cursiva_de_una_parte_se_respeta():
    doc = trabajo()
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.add_run("Como se lee en ")
    p.add_run("El campo vacío").italic = True
    p.add_run(", " + CITA + FUENTE)
    informe = aplicar(doc)
    assert not any("cursiva" in c for c in informe.cambios)
    estilos = Estilos(doc)
    p = parrafo(doc, "Como se lee")
    assert [titulos.cursiva_efectiva(estilos, r, p) for r in titulos.fragmentos_con_letras(p)] == [False, True, False]


def test_si_ya_estaba_bien_no_cambia_nada():
    doc = trabajo(CITA + FUENTE)
    p = parrafo(doc, "El turismo")
    p.find(f"{w('pPr')}/{w('ind')}").set(w("left"), "720")
    p.find(f"{w('pPr')}/{w('ind')}").set(w("firstLine"), "0")
    assert aplicar(doc).cambios == []


def test_punto_detras_de_la_fuente():
    informe = aplicar(trabajo(CITA + FUENTE + "."))
    assert any(a.startswith("1 cita en bloque acaba con un punto después del paréntesis") for a in informe.avisos)


def test_parrafo_con_sangria_de_menos_de_40_palabras():
    doc = trabajo("Una frase corta con sangría a la izquierda.")
    informe = aplicar(doc)
    assert informe.cambios == []
    assert any(a.startswith("1 párrafo con sangría a la izquierda tiene menos de 40 palabras") for a in informe.avisos)
    assert sangrias(doc, parrafo(doc, "Una frase corta"))["left"] == "567", "se deja como estaba"


def test_cita_larga_dentro_de_un_parrafo():
    doc = trabajo()
    doc.add_paragraph(f"Según el informe, «{CITA}»{FUENTE}.")
    informe = aplicar(doc)
    assert any(a.startswith("1 cita entre comillas dentro de un párrafo tiene 40 palabras o más") for a in informe.avisos)


def test_el_verificador_ve_la_cursiva_y_la_sangria(tmp_path):
    doc = trabajo(CITA + FUENTE)
    aplicar(doc)
    p = parrafo(doc, "El turismo")
    for r in titulos.fragmentos_con_letras(p):
        _poner(r, "i", True)
    p.find(f"{w('pPr')}/{w('ind')}").set(w("right"), "567")
    assert incumplimientos(doc) == ["Cita en bloque: tiene sangría a la derecha.", "Cita en bloque: va toda en cursiva."]
