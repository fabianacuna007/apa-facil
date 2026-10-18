"""
Paso 1.5 (versión simple): sangría de 1,27 cm en la primera línea de los
párrafos de texto del cuerpo, y en ningún otro.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

from pathlib import Path
from types import SimpleNamespace

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Cm, Pt

from apa_facil import estructura, pipeline
from apa_facil.documento import Estilos, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import sangria, titulos
from apa_facil.reglas.sangria import motivos_sin_sangria

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
TEXTO = ("Este es un párrafo normal de texto con bastantes palabras para que no parezca un título "
         "ni una línea de portada, y que termina en punto como cualquier frase.")


def aplicar(doc):
    informe = Informe()
    sangria.aplicar(doc, {}, informe)
    return informe


def primera_linea(doc, p):
    estilos = Estilos(doc)
    p = p._p if hasattr(p, "_p") else p
    return int(estilos.parrafo(p, "ind", "firstLine") or 0), int(estilos.parrafo(p, "ind", "hanging") or 0)


def documento_con_portada():
    """Un documento cuya primera página es una portada (para que no estorbe en las pruebas)."""
    doc = docx.Document()
    doc.add_paragraph("Portada").runs[0].add_break(WD_BREAK.PAGE)
    return doc


def motivo(doc, p):
    return dict(motivos_sin_sangria(doc))[p._p]


# --- Lo que lleva sangría -------------------------------------------------------------

def test_parrafo_de_texto_lleva_sangria():
    doc = documento_con_portada()
    p = doc.add_paragraph(TEXTO)
    informe = aplicar(doc)
    assert primera_linea(doc, p) == (720, 0)
    assert informe.cambios == ["Sangría: 1 párrafo de texto lleva ahora 1,27 cm en la primera línea."]


def test_sangria_de_otro_tamano_se_corrige():
    doc = documento_con_portada()
    p = doc.add_paragraph(TEXTO)
    p.paragraph_format.first_line_indent = Cm(1)
    aplicar(doc)
    assert primera_linea(doc, p) == (720, 0)


def test_sangria_del_estilo_pasa_al_parrafo():
    """Si el estilo Normal la traía, se quita del estilo: los títulos y tablas no deben heredarla."""
    doc = documento_con_portada()
    doc.styles["Normal"].paragraph_format.first_line_indent = Cm(1)
    p = doc.add_paragraph(TEXTO)
    titulo = doc.add_heading("Método", level=1)
    celda = doc.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0]
    aplicar(doc)
    assert doc.styles["Normal"].paragraph_format.first_line_indent is None
    assert primera_linea(doc, p) == (720, 0)
    assert primera_linea(doc, titulo)[0] == primera_linea(doc, celda)[0] == 0


def test_frase_que_acaba_en_dos_puntos_lleva_sangria():
    doc = documento_con_portada()
    p = doc.add_paragraph("La tabla siguiente muestra la evolución de la población:")
    aplicar(doc)
    assert primera_linea(doc, p)[0] == 720


def test_si_ya_la_tiene_no_hay_cambio():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO).paragraph_format.first_line_indent = Cm(1.27)
    assert aplicar(doc).cambios == []


# --- Lo que no lleva sangría ------------------------------------------------------------

@pytest.mark.parametrize("texto, esperado", [
    ("Introducción", "título"),                                 # corto, sin punto y seguido de texto
    ("1. Cambios en la organización del trabajo", "título"),    # numerado y seguido de texto (paso 3.2)
    ("Datos de uso por red social", "título"),
    ("•\tComparación social con imágenes retocadas", "lista"),
    ("-  Ciberacoso", "lista"),
    ("a) Primera condición del experimento, que se describe con detalle.", "lista"),
    ("Tabla 1. Evolución de la población", "leyenda de tabla"),
    ("Fuente: elaboración propia a partir del padrón municipal.", "leyenda de tabla"),
    ("Nota. Los datos son provisionales.", "leyenda de tabla"),
    ("\t" + TEXTO, "tabulador"),
    ("   " + TEXTO, "tabulador"),
    ("", "vacío"),
])
def test_exclusiones_por_el_texto(texto, esperado):
    """Cada caso va seguido de un párrafo de texto, como en un trabajo de verdad."""
    doc = documento_con_portada()
    p = doc.add_paragraph(texto)
    doc.add_paragraph(TEXTO)
    aplicar(doc)
    assert motivo(doc, p) == esperado
    assert primera_linea(doc, p)[0] == 0


def test_titulo_corto_en_negrita_aunque_acabe_en_dos_puntos():
    doc = documento_con_portada()
    p = doc.add_paragraph()
    p.add_run("Factores de riesgo:").bold = True
    assert motivo(doc, p) == "título"


def test_titulo_numerado_en_negrita():
    doc = documento_con_portada()
    p = doc.add_paragraph()
    p.add_run("1. Cambios en la organización del trabajo").bold = True
    assert motivo(doc, p) == "título"


def test_titulo_de_word():
    doc = documento_con_portada()
    p = doc.add_heading("Un título largo que tiene más de quince palabras para comprobar que el estilo manda "
                        "sobre la longitud.", level=2)
    assert motivo(doc, p) == "título"


def test_lista_de_word():
    doc = documento_con_portada()
    p = doc.add_paragraph(TEXTO, style="List Bullet")
    aplicar(doc)
    assert motivo(doc, p) == "lista"
    assert primera_linea(doc, p)[0] == 0


def test_centrado():
    doc = documento_con_portada()
    p = doc.add_paragraph(TEXTO)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    assert motivo(doc, p) == "centrado o a la derecha"


def test_cita_en_bloque():
    doc = documento_con_portada()
    p = doc.add_paragraph(TEXTO)
    p.paragraph_format.left_indent = Cm(1.27)
    aplicar(doc)
    assert motivo(doc, p) == "cita en bloque"
    assert primera_linea(doc, p)[0] == 0


def test_referencias_despues_del_rotulo_hasta_el_siguiente_titulo():
    doc = documento_con_portada()
    doc.add_paragraph("Webgrafía:")
    referencias = [doc.add_paragraph(f"Autor, A. ({2000 + i}). Título de una obra con su editorial. Editorial.")
                   for i in range(2)]
    doc.add_heading("Anexo", level=1)
    despues = doc.add_paragraph(TEXTO)
    aplicar(doc)
    assert [motivo(doc, r) for r in referencias] == ["referencia", "referencia"]
    assert primera_linea(doc, despues)[0] == 720


def test_sangria_francesa_con_autor_y_ano_es_referencia():
    doc = documento_con_portada()
    p = doc.add_paragraph("Molina, R. (2017). Turismo y territorio: límites del desarrollo rural. Editorial.")
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.27), Cm(-1.27)
    aplicar(doc)
    assert motivo(doc, p) == "referencia"
    assert primera_linea(doc, p)[1] > 0, "la sangría francesa se conserva"


def test_sangria_francesa_sin_autor_ni_ano_se_conserva():
    """Fuera de la lista y sin autor y año, es una referencia dudosa: mejor no tocarla."""
    doc = documento_con_portada()
    p = doc.add_paragraph(TEXTO)
    p.paragraph_format.left_indent, p.paragraph_format.first_line_indent = Cm(1.27), Cm(-1.27)
    aplicar(doc)
    assert motivo(doc, p) == "referencia"
    assert primera_linea(doc, p) == (0, 720)


def test_resumen():
    doc = documento_con_portada()
    doc.add_paragraph("Resumen")
    resumen = doc.add_paragraph(TEXTO)
    siguiente = doc.add_paragraph(TEXTO)
    aplicar(doc)
    assert motivo(doc, resumen) == "primer párrafo del resumen"
    assert primera_linea(doc, siguiente)[0] == 720


def test_portada():
    doc = docx.Document()
    portada = doc.add_paragraph("Trabajo de Psicología del Desarrollo, curso 2025-2026.")
    doc.add_paragraph(TEXTO * 2)
    aplicar(doc)
    assert motivo(doc, portada) == "portada"


def test_resumen_con_palabras_clave():
    doc = documento_con_portada()
    doc.add_paragraph().add_run("Resumen").bold = True
    resumen = doc.add_paragraph(TEXTO)
    claves = doc.add_paragraph("Palabras clave: sueño, memoria, rendimiento académico.")
    aplicar(doc)
    assert primera_linea(doc, resumen)[0] == 0
    assert primera_linea(doc, claves)[0] == 720, "APA: la línea de palabras clave lleva sangría"


# --- La estructura se decide sobre el original (paso 3.4) --------------------------------

def titulo_solo_por_tamano(doc):
    """«Resultados:» acaba en dos puntos y no va en negrita: solo su tamaño (14 pt) dice que es un título."""
    titulo = doc.add_paragraph()
    titulo.add_run("Resultados:").font.size = Pt(14)
    doc.add_paragraph().add_run(TEXTO).font.size = Pt(11)
    return titulo


def igualar_tamanos(doc, puntos=11):
    """Lo que hace la regla de fuente: el mismo tamaño en todo el texto."""
    for p in doc.paragraphs:
        for r in p.runs:
            r.font.size = Pt(puntos)


def test_con_la_estructura_del_original_el_titulo_no_lleva_sangria():
    doc = documento_con_portada()
    titulo = titulo_solo_por_tamano(doc)
    del_original = estructura.clasificar(doc)
    igualar_tamanos(doc)
    sangria.aplicar(doc, {}, Informe(), del_original)
    assert primera_linea(doc, titulo)[0] == 0


def test_sin_la_estructura_del_original_el_titulo_se_tomaria_por_texto():
    """Por eso el pipeline clasifica antes de aplicar ninguna regla."""
    doc = documento_con_portada()
    titulo = titulo_solo_por_tamano(doc)
    igualar_tamanos(doc)
    sangria.aplicar(doc, {}, Informe())
    assert primera_linea(doc, titulo)[0] == 720


def test_las_reglas_reciben_la_estructura_del_original(tmp_path, monkeypatch):
    """
    En el 04, la copia formateada se clasificaría distinto (sin la letra más grande de sus títulos)
    si sus títulos no recibieran su estilo de Word (paso 4.1). Con los estilos, la copia lleva su
    estructura escrita y se clasifica igual que el original.
    """
    original = next(p for p in PRUEBAS if p.name.startswith("04"))
    recibidas = []
    espia = SimpleNamespace(__name__="espia", aplicar=lambda doc, opciones, informe, estructura: recibidas.append(
        [(c.numero, c.tipo, c.nivel) for c in estructura]))
    monkeypatch.setattr(pipeline, "REGLAS", [espia] + pipeline.REGLAS + [espia])
    _, informe = pipeline.procesar(original, tmp_path)
    del_original = [(c.numero, c.tipo, c.nivel) for c in estructura.clasificar(docx.Document(str(original)))]
    borrados = {c.parrafo for c in informe.cambios_previstos if c.accion == "borrar"}
    assert recibidas[0] == del_original
    assert recibidas[1] == [c for c in del_original if c[0] not in borrados],         "la misma estructura después de las reglas, sin los párrafos vacíos borrados"
    copia = docx.Document(str(tmp_path / "04_formato_caotico_APA.docx"))
    # Sin el título del trabajo, que el paso 4.4 añade al principio del texto.
    de_la_copia = [(c.tipo, c.nivel) for c in estructura.clasificar(copia)
                   if Estilos(copia).nombre(Estilos(copia).estilo_de(c.elemento)).lower() != "title"]
    assert de_la_copia == [(c[1], c[2]) for c in recibidas[1]]

    sin_titulos = [r for r in pipeline.REGLAS if r is not titulos and r is not espia]
    monkeypatch.setattr(pipeline, "REGLAS", sin_titulos)
    otra = tmp_path / "sin_titulos"
    otra.mkdir()
    pipeline.procesar(original, otra)
    sin_estilos = [(c.tipo, c.nivel) for c in estructura.clasificar(docx.Document(str(otra / original.name.replace(".docx", "_APA.docx"))))]
    assert sin_estilos != [(c[1], c[2]) for c in recibidas[1]], "sin los estilos de título, la copia se clasifica distinto"


def test_aviso_de_tabulador():
    doc = documento_con_portada()
    doc.add_paragraph("\t" + TEXTO)
    informe = aplicar(doc)
    assert informe.avisos[0].startswith("1 párrafo empieza con un tabulador o espacios usados como sangría "
                                        "(«Este es un párrafo normal de texto")


# --- Documentos de prueba --------------------------------------------------------------

# 04: «Más información en…» es texto (paso 3.4) y sus 3 párrafos con tabulador se limpian (paso 3.5).
CON_SANGRIA = {"01": 8, "02": 8, "03": 6, "04": 7, "05": 13}
EMPIEZAN_POR = {   # algunos párrafos que NO deben llevar sangría, con el motivo
    "01": [("BIBLIOGRAFÍA", "rótulo de referencias"), ("Ramírez Soto", "referencia"), ("1. Cambios", "título")],
    "02": [("Este informe analiza", "primer párrafo del resumen"), ("Castro, P.", "referencia")],
    "03": [("Mejorar el aislamiento", "lista"), ("Herrero, C.", "referencia")],
    "04": [("•", "lista"), ("Webgrafía", "rótulo de referencias"), ("Martínez, J.", "referencia"),
           ("Efectos sobre", "título")],
    "05": [("Este trabajo analiza", "primer párrafo del resumen"), ('"El turismo rural', "cita en bloque"),
           ("Fuente:", "leyenda de tabla"), ("1. ¿Desde", "lista"), ("Molina, R.", "referencia")],
}


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    clave = original.name[:2]
    salida, informe = pipeline.procesar(original, tmp_path)
    copia = docx.Document(str(salida))
    estilos = Estilos(copia)
    clasificados = motivos_sin_sangria(copia)
    # Los títulos de nivel 4 y 5 llevan sangría de primera línea: la pone su estilo (paso 4.1).
    con_sangria_de_titulo = {c.elemento for c in estructura.clasificar(copia)
                             if c.tipo == "titulo" and c.nivel in (4, 5)}

    sangrados = [p for p, m in clasificados if m is None]
    assert len(sangrados) == CON_SANGRIA[clave]
    for p, m in clasificados:
        esperada = 720 if m is None else None
        primera = int(estilos.parrafo(p, "ind", "firstLine") or 0)
        if esperada or p in con_sangria_de_titulo:
            assert primera == 720
        elif m != "tabulador":
            assert primera == 0, f"«{texto_parrafo(p)[:30]}» ({m}) no debe llevar sangría"
    for inicio, esperado in EMPIEZAN_POR[clave]:
        p = next(p for p, _ in clasificados if texto_parrafo(p).strip().startswith(inicio))
        assert dict(clasificados)[p] == esperado, inicio
