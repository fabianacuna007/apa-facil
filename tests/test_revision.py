"""
Paso 7.5: la revisión de citas y referencias en el informe
(apa_facil/revision/revisar.py). Cada aviso es un tipo de problema y lleva sus
casos: dónde está cada uno (en palabras: Word no numera los párrafos), lo escrito
y cómo se corrige. El documento no cambia, y no se repite lo que ya avisa el
paso 6.5 (orden alfabético, referencias que son solo una URL, obras con los
mismos autores y año sin letra).
"""

import hashlib
import json
from pathlib import Path

import docx
import pytest
from lxml import etree

from apa_facil import documento, estructura, pipeline
from apa_facil.informe import MAXIMO_DE_CASOS_EN_LA_TERMINAL, Informe, caso_como_texto
from apa_facil.revision import revisar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {ruta.name[:2]: ruta for ruta in sorted((RAIZ / "pruebas").glob("*.docx"))}
TEXTO = ("El turismo rural es un fenómeno estudiado desde hace décadas por la geografía y la economía, y su relación "
         "con la despoblación sigue abierta en muchas comarcas del interior")


def trabajo(parrafos, referencias, apartado="Marco teórico"):
    """Un trabajo con un apartado, sus párrafos y una lista de referencias."""
    doc = docx.Document()
    doc.add_heading(apartado, level=1)
    for texto in parrafos:
        doc.add_paragraph(texto)
    doc.add_heading("Referencias", level=1)
    for texto in referencias:
        doc.add_paragraph(texto)
    return doc


def avisos(doc):
    informe = Informe()
    informe.regla = "revision"
    revisar.revisar(doc, estructura.clasificar(doc), informe)
    return informe.entradas


MOLINA = "Molina, R. (2017). Turismo y territorio. Editorial Ejemplo."


# --- Dónde está cada caso -------------------------------------------------------------

def test_en_el_texto_el_apartado_y_el_principio_del_parrafo():
    (aviso,) = avisos(trabajo([TEXTO + " (Molina 2017)."], [MOLINA]))
    (caso,) = aviso.casos
    assert caso == {"donde": "En «Marco teórico», en el párrafo que empieza «El turismo rural es un fenómeno…»",
                    "extracto": "(Molina 2017)", "correccion": "(Molina, 2017)", "explicacion": None, "parrafo": 2}
    assert aviso.mensaje == "Falta la coma antes del año en las citas del texto. Entre el autor y el año va una coma."


def test_un_parrafo_corto_sale_entero():
    (aviso,) = avisos(trabajo(["Lo dice Molina (2017) (Molina 2017)."], [MOLINA]))
    assert aviso.casos[0]["donde"] == "En «Marco teórico», en el párrafo que empieza «Lo dice Molina (2017) (Molina 2017).»"


def test_en_una_tabla_y_en_la_lista_de_referencias():
    """En el 04: la columna «Fuente» de su tabla, y la Webgrafía (por el principio de cada referencia)."""
    doc = documento.abrir(PRUEBAS["04"])
    entradas = avisos(doc)
    lugares = {c["donde"] for e in entradas for c in e.casos}
    assert "En la 1.ª tabla (apartado «Factores de riesgo»)" in lugares
    assert "En la lista de referencias: «Pérez, L.; Romero, A.; Saiz, T.; Moya, P. (2…»" in lugares


def test_en_una_nota_al_pie():
    doc = documento.abrir(PRUEBAS["05"])

    def citar_en_la_primera_nota(nombre, raiz):
        if nombre == "/word/footnotes.xml":
            nota = next(n for n in raiz.iter(documento.w("footnote")) if n.get(documento.w("type")) is None)
            next(nota.iter(documento.w("t"))).text = "Así lo define Molina (2017, p.12)."
    documento.editar_partes_de_texto(doc, citar_en_la_primera_nota)
    (caso,) = [c for e in avisos(doc) if e.mensaje.startswith("Falta el espacio") for c in e.casos
               if c["parrafo"] is None]
    assert (caso["donde"], caso["extracto"], caso["correccion"]) == (
        "En la nota al pie 1", "Molina (2017, p.12)", "Molina (2017, p. 12)")


def test_un_titulo_de_la_portada_no_es_un_apartado():
    doc = docx.Document()
    doc.add_heading("La despoblación rural", level=1)
    doc.add_paragraph(TEXTO + " (Molina 2017).")
    doc.add_heading("Referencias", level=1)
    doc.add_paragraph(MOLINA)
    clasificacion = estructura.clasificar(doc)
    clasificacion[0].zona = "portada"   # como el título del trabajo en la portada de la universidad
    informe = Informe()
    revisar.revisar(doc, clasificacion, informe)
    assert informe.entradas[0].casos[0]["donde"] == "En el párrafo que empieza «El turismo rural es un fenómeno…»"


def test_sin_apartado():
    doc = docx.Document()
    doc.add_paragraph(TEXTO + " (Molina 2017).")
    doc.add_heading("Referencias", level=1)
    doc.add_paragraph(MOLINA)
    (aviso,) = avisos(doc)
    assert aviso.casos[0]["donde"] == "En el párrafo que empieza «El turismo rural es un fenómeno…»"


# --- Los avisos -----------------------------------------------------------------------

def test_cita_sin_referencia_y_referencia_sin_cita():
    sin_referencia, no_citada = avisos(trabajo([TEXTO + " (Castell, 2011)."], [MOLINA]))
    assert sin_referencia.mensaje.startswith("Una obra citada en el texto no está en la lista de referencias.")
    assert [c["extracto"] for c in sin_referencia.casos] == ["(Castell, 2011)"]
    assert no_citada.mensaje.startswith("Una obra de la lista de referencias no se cita en el texto.")
    assert no_citada.casos[0]["donde"].startswith("En la lista de referencias: «Molina, R. (2017).")


def test_la_pista_de_una_cita_sin_referencia():
    sin_referencia, _ = avisos(trabajo([TEXTO + " (Ramirez, 2014)."], ["Ramírez Soto, M. (2015). Fábricas. Editorial."]))
    assert sin_referencia.casos[0]["explicacion"] == ("se parece a «Ramírez Soto, M. (2015)» de la lista, pero no "
                                                      "coinciden ni los autores ni el año")


def test_varias_obras_y_todas_sus_ubicaciones():
    (aviso,) = avisos(trabajo([TEXTO + " (Castell, 2011).", "Según Castell (2011) y Pérez (2019), así es."], []))
    assert aviso.mensaje.startswith("2 obras citadas en el texto no están")
    assert [c["extracto"] for c in aviso.casos] == ["(Castell, 2011)", "Castell (2011)", "Pérez (2019)"]


def test_año_y_autores_distintos_con_su_detalle_en_cada_caso():
    año, autores = avisos(trabajo([TEXTO + " (Ramirez, 2015) y (Molina, 2016)."],
                                  ["Ramírez, M. (2015). Fábricas. Editorial.", MOLINA]))
    assert año.mensaje.startswith("El año de la cita no coincide con el de su referencia")
    assert (año.casos[0]["extracto"], año.casos[0]["explicacion"]) == ("(Molina, 2016)", "en la referencia, 2017")
    assert autores.mensaje.startswith("Los autores de la cita no están escritos como en su referencia")
    assert autores.casos[0]["explicacion"].endswith("cambian las tildes, las mayúsculas o los guiones")


def test_la_explicacion_comun_va_en_el_mensaje_y_la_propia_en_cada_caso():
    """Mismo error, distinta explicación (el año sin paréntesis, detrás del autor o no): va en cada caso."""
    (aviso,) = avisos(trabajo([TEXTO + " (Lefèvre, 2009) y (Martínez, 2019)."],
                              ["Lefèvre, A. 2009. Industrialisation. Éditions Exemple.",
                               "Martínez, J. Redes sociales. 2019. https://www.example.org/redes"]))
    assert aviso.mensaje == "Año sin paréntesis en la lista de referencias."
    assert [c["explicacion"] for c in aviso.casos] == [
        "El año va entre paréntesis, seguido de un punto.",
        "El año va entre paréntesis, seguido de un punto, justo detrás del autor."]


def test_falta_el_año_con_la_pista_del_texto():
    """El «falta el año» (7.4) y el «año distinto» (7.3) de una referencia sin año son un solo aviso."""
    entradas = avisos(trabajo([TEXTO + ", como dice Herranz (2015)."],
                              ["Herranz, M. Turismo rural. Revista de Estudios, 41, 77-96."]))
    (aviso,) = entradas
    assert aviso.mensaje == ("Falta el año en la lista de referencias. Detrás del autor va el año entre paréntesis; si "
                             "la obra no tiene fecha, «(s. f.)».")
    (caso,) = aviso.casos
    assert (caso["correccion"], caso["explicacion"]) == ("Herranz, M. (2015).", "en el texto la citas como de 2015")


def test_sin_la_pista_si_no_se_cita():
    (no_citada, falta) = avisos(trabajo([TEXTO + "."], ["Herranz, M. Turismo rural. Revista de Estudios, 41, 77-96."]))
    assert (falta.casos[0]["correccion"], falta.casos[0]["explicacion"]) == ("Herranz, M. (s. f.).", None)
    assert no_citada.mensaje.startswith("Una obra de la lista de referencias no se cita")


def test_lo_que_ya_avisa_el_6_5_no_se_repite():
    """El orden, las URL sueltas y las obras de los mismos autores y año sin letra."""
    entradas = avisos(trabajo([TEXTO + " (Molina, 2017; Collantes, 2011)."],
                              [MOLINA, "Collantes, F. (2011). Lugares. Editorial.", "www.example.org/datos"]))
    assert entradas == []
    repetidas = avisos(trabajo([TEXTO + " (Molina, 2017)."], [MOLINA, "Molina, R. (2017). Otro libro. Editorial."]))
    assert repetidas == []


def test_una_referencia_no_reconocida_que_no_es_una_url_si_se_avisa():
    (aviso,) = avisos(trabajo([TEXTO + "."], ["Redes sociales y adolescencia, varios autores, sin fecha."]))
    assert aviso.mensaje.startswith("Referencia que no se ha podido leer en la lista de referencias.")
    assert aviso.casos[0]["donde"] == "En la lista de referencias"


def test_sin_problemas_no_hay_avisos():
    assert avisos(trabajo([TEXTO + " (Molina, 2017)."], [MOLINA])) == []


# --- En el informe --------------------------------------------------------------------

def test_en_la_terminal_cada_caso_en_su_linea_y_como_mucho_diez():
    informe = Informe()
    informe.regla = "revision"
    casos = [{"donde": f"En «A», en el párrafo que empieza «Texto {i}…»", "extracto": "(Pérez et. al, 2020)",
              "correccion": "(Pérez et al., 2020)", "explicacion": None, "parrafo": i} for i in range(12)]
    informe.aviso("«et al.» mal escrito en las citas del texto.", casos)
    texto = informe.como_texto()
    assert texto.count("         – En «A»") == MAXIMO_DE_CASOS_EN_LA_TERMINAL
    assert "         – y 2 más (están todos en el informe completo)." in texto
    assert len(informe.a_dict()["avisos"][0]["casos"]) == 12
    assert caso_como_texto(casos[0]) == ("En «A», en el párrafo que empieza «Texto 0…»: «(Pérez et. al, 2020)» → "
                                         "«(Pérez et al., 2020)»")


def test_en_el_json(tmp_path):
    salida, informe = pipeline.procesar(PRUEBAS["05"], tmp_path, {"citas_tal_cual": True})   # el 7.7 lo corregiría
    datos = json.loads(salida.with_suffix(".json").read_text(encoding="utf-8"))
    assert datos["version"] == 15
    (doi,) = [a for a in datos["avisos"] if a["mensaje"].startswith("DOI con la forma antigua")]
    assert doi["nombre"] == "Citas y referencias"
    assert doi["casos"] == [{"donde": "En la lista de referencias: «Herranz, M. Turismo rural y servicios públic…»",
                             "extracto": "doi:10.0000/rerf.2015.041",
                             "correccion": "https://doi.org/10.0000/rerf.2015.041", "explicacion": None,
                             "parrafo": 71}]
    assert all("casos" not in a for a in datos["avisos"] if a["regla"] != "revision")


def test_los_avisos_van_al_final_y_el_documento_no_cambia(tmp_path):
    huella = hashlib.sha256(PRUEBAS["04"].read_bytes()).hexdigest()
    _, informe = pipeline.procesar(PRUEBAS["04"], tmp_path)
    reglas = [e.regla for e in informe.entradas if e.tipo == "aviso"]
    primera = reglas.index("revision")
    assert set(reglas[primera:]) == {"revision"}
    assert hashlib.sha256(PRUEBAS["04"].read_bytes()).hexdigest() == huella


def test_se_analiza_el_original_antes_de_las_reglas():
    """La revisión guarda lo que necesita antes de que las reglas reordenen o borren nada."""
    doc = documento.abrir(PRUEBAS["05"])
    clasificacion = estructura.clasificar(doc)
    antes = etree.tostring(doc.element)
    revision = revisar.analizar(doc, clasificacion)
    assert etree.tostring(doc.element) == antes
    for p in list(doc.element.body.iter(documento.w("p"))):   # aunque luego el documento se vacíe…
        p.getparent().remove(p)
    informe = Informe()
    revisar.avisar(revision, informe)   # …los avisos salen igual
    assert any(e.mensaje.startswith("DOI con la forma antigua") for e in informe.entradas)
