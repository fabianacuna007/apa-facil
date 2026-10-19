"""
Paso 7.6: los avisos de citas y referencias como comentarios de Word en la copia
(apa_facil/revision/comentarios.py), solo con --comentarios (decisión del
propietario). Un comentario por caso, de «APA Fácil», sobre el texto del caso;
el texto del trabajo no cambia. También documento.aislar, que deja un trozo del
texto en sus propios fragmentos para que el comentario lo marque entero.
"""

import subprocess
import tempfile
from pathlib import Path

import docx
import pytest
from lxml import etree

from apa_facil import documento, estructura, pipeline
from apa_facil.__main__ import main
from apa_facil.documento import aislar, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.revision import comentarios, revisar
from apa_facil.revision.revisar import Ancla, Revision
from herramientas import a_pdf
from herramientas.verificar import verificar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {ruta.name[:2]: ruta for ruta in sorted((RAIZ / "pruebas").glob("*.docx"))}
CON_COMENTARIOS = {"comentarios": True}


def marcados(ruta):
    """{id del comentario: texto que marca}, leyendo las marcas del documento en orden."""
    cuerpo = docx.Document(str(ruta)).element.body
    abiertos, textos = set(), {}
    for el in cuerpo.iter(w("commentRangeStart"), w("commentRangeEnd"), w("t"), w("tab")):
        nombre = etree.QName(el).localname
        if nombre == "commentRangeStart":
            abiertos.add(el.get(w("id")))
            textos.setdefault(el.get(w("id")), "")
        elif nombre == "commentRangeEnd":
            abiertos.discard(el.get(w("id")))
        else:
            for i in abiertos:
                textos[i] += el.text or "" if nombre == "t" else "\t"
    return textos


def comentarios_de(ruta):
    """[(texto marcado, texto del comentario)] en el orden del documento."""
    doc = docx.Document(str(ruta))
    por_id = {str(c.comment_id): c for c in doc.comments}
    return [(texto, por_id[i].text) for i, texto in marcados(ruta).items()]


# --- documento.aislar ------------------------------------------------------------------

def parrafo_de_prueba():
    doc = docx.Document()
    p = doc.add_paragraph()
    p.add_run("Hola mundo ")
    p.add_run("(Martínez 2019)").bold = True
    p.add_run(" fin\tcon tab")
    return p._p


def texto_entre(p, primero, ultimo):
    piezas, dentro = [], False
    for r in p.iter(w("r")):
        dentro = dentro or r is primero
        if dentro:
            piezas.append("".join(el.text or "" if el.tag == w("t") else "\t"
                                  for el in r if el.tag in (w("t"), w("tab"))))
        if r is ultimo:
            break
    return "".join(piezas)


@pytest.mark.parametrize("inicio, fin", [(5, 10), (11, 26), (0, 3), (20, 32), (26, 34), (0, 38)])
def test_aislar_deja_el_trozo_en_sus_fragmentos_sin_cambiar_el_texto(inicio, fin):
    p = parrafo_de_prueba()
    antes = texto_parrafo(p)
    primero, ultimo = aislar(p, inicio, fin)
    assert texto_entre(p, primero, ultimo) == antes[inicio:fin]
    assert texto_parrafo(p) == antes


def test_aislar_conserva_el_formato_de_cada_fragmento():
    p = parrafo_de_prueba()
    primero, ultimo = aislar(p, 15, 20)   # dentro de la parte en negrita
    assert primero is ultimo and primero.find(f"{w('rPr')}/{w('b')}") is not None


@pytest.mark.parametrize("inicio, fin", [(3, 3), (30, 99), (-1, 4)])
def test_aislar_un_trozo_vacio_o_fuera_del_parrafo(inicio, fin):
    assert aislar(parrafo_de_prueba(), inicio, fin) is None


# --- Sin la opción, nada --------------------------------------------------------------

def test_sin_la_opcion_no_hay_comentarios(tmp_path):
    copia, informe = pipeline.procesar(PRUEBAS["04"], tmp_path)
    assert list(docx.Document(str(copia)).comments) == []
    assert not any("comentario" in m for m in informe.cambios)


# --- Con la opción ----------------------------------------------------------------------

def casos_del_informe(informe):
    return [c for e in informe.entradas if e.regla == "revision" and e.casos for c in e.casos]


@pytest.mark.parametrize("clave", ["01", "03", "04", "05"])
def test_un_comentario_por_caso_sobre_su_texto(clave, tmp_path):
    copia, informe = pipeline.procesar(PRUEBAS[clave], tmp_path, CON_COMENTARIOS)
    casos = casos_del_informe(informe)
    hallados = comentarios_de(copia)
    assert len(hallados) == len(casos)
    for marcado, texto in hallados:
        assert texto.startswith("APA Fácil: ")
    # Cada caso tiene su comentario sobre su texto (o sobre la referencia entera, que empieza por él).
    libres = [" ".join(m.split()) for m, _ in hallados]
    for caso in casos:
        marca = next((m for m in libres if m == caso["extracto"]), None) or \
            next((m for m in libres if m.startswith(caso["extracto"])), None)
        assert marca is not None, caso
        libres.remove(marca)
    assert f"se han puesto {len(casos)} comentario" in " ".join(informe.cambios)


def test_los_comentarios_del_04(tmp_path):
    copia, _ = pipeline.procesar(PRUEBAS["04"], tmp_path, CON_COMENTARIOS)
    hallados = comentarios_de(copia)
    assert [m for m, _ in hallados][:9] == ["(Martínez 2019)", "(Pérez et. al, 2020)", "Gómez (s.f)", "Gómez (s.f)",
                                            "Pérez et. al, 2020", "INE (s.f)", "INE (s.f)", "Gómez, s.f.", "Gómez, s.f."]
    (coma,) = [t for m, t in hallados if m == "(Martínez 2019)"]
    assert coma == ("APA Fácil: Falta la coma antes del año en las citas del texto. Entre el autor y el año va una coma. "
                    "Corrección: «(Martínez, 2019)».")


def test_las_referencias_reordenadas_se_llevan_su_comentario(tmp_path):
    """En el 05 el 6.5 reordena la lista: los comentarios siguen a sus referencias."""
    copia, _ = pipeline.procesar(PRUEBAS["05"], tmp_path, CON_COMENTARIOS)
    de_la_lista = [m for m, _ in comentarios_de(copia)][2:]
    assert de_la_lista == ["Consejo Económico de Ejemplo (2022)", "Madrid", "Herranz, M.", "doi:10.0000/rerf.2015.041",
                           "Torres, E. (2019). Agroturismo y empleo femenino. Editorial Ejemplo."]
    (falta_año,) = [t for m, t in comentarios_de(copia) if m == "Herranz, M."]
    assert falta_año.endswith("En este caso, en el texto la citas como de 2015. Corrección: «Herranz, M. (2015).».")


@pytest.mark.parametrize("clave", ["04", "05"])
def test_el_texto_no_cambia_y_la_copia_cumple(clave, tmp_path):
    copia, _ = pipeline.procesar(PRUEBAS[clave], tmp_path, CON_COMENTARIOS)
    sin, _ = pipeline.procesar(PRUEBAS[clave], tmp_path / "sin")
    assert documento.comparar(documento.radiografia(docx.Document(str(sin))),
                              documento.radiografia(docx.Document(str(copia)))) == []
    assert verificar(copia).incumplimientos == []


def test_conserva_los_comentarios_que_ya_tenia(tmp_path):
    con_uno = tmp_path / "04_con_un_comentario.docx"
    doc = docx.Document(str(PRUEBAS["04"]))
    doc.add_comment(doc.paragraphs[6].runs, text="Revisar esta definición.", author="Profesora", initials="P")
    doc.save(str(con_uno))
    copia, informe = pipeline.procesar(con_uno, tmp_path / "salida", CON_COMENTARIOS)
    autores = [c.author for c in docx.Document(str(copia)).comments]
    assert autores.count("Profesora") == 1 and autores.count("APA Fácil") == len(casos_del_informe(informe))
    ids = [c.comment_id for c in docx.Document(str(copia)).comments]
    assert len(ids) == len(set(ids))


def test_las_notas_al_pie_no_llevan_comentario(tmp_path):
    con_cita = tmp_path / "05_cita_en_nota.docx"
    doc = docx.Document(str(PRUEBAS["05"]))

    def citar_en_la_primera_nota(nombre, raiz):
        if nombre == "/word/footnotes.xml":
            nota = next(n for n in raiz.iter(w("footnote")) if n.get(w("type")) is None)
            next(nota.iter(w("t"))).text = "Así lo define Molina (2017, p.12)."
    documento.editar_partes_de_texto(doc, citar_en_la_primera_nota)
    doc.save(str(con_cita))
    copia, informe = pipeline.procesar(con_cita, tmp_path / "salida", CON_COMENTARIOS)
    assert len(list(docx.Document(str(copia)).comments)) == len(casos_del_informe(informe)) - 1
    assert ("No llevan comentario 1 caso de las notas al pie (en ellas no se ponen comentarios): están en este "
            "informe.") in informe.avisos


# --- Cuando el texto ya no está -------------------------------------------------------

def revision_con(doc, *anclas):
    return Revision(anclas=[("Un aviso.", {"explicacion": None, "correccion": None}, a) for a in anclas])


def test_sin_sitio_no_se_adivina():
    doc = docx.Document()
    p = doc.add_paragraph("Según Molina (2017), el turismo rural ayuda.")._p
    quitado = doc.add_paragraph("Un párrafo que las reglas quitarán (Pérez 2019).")._p
    quitado.getparent().remove(quitado)
    informe = Informe()
    puestos = comentarios.comentar(doc, revision_con(doc, Ancla(p, "(Gómez, s. f.)", 0), Ancla(quitado, "(Pérez 2019)")),
                                   informe)
    assert puestos == 0 and list(doc.comments) == []
    assert informe.avisos == ["No llevan comentario 2 casos cuyo texto ya no se encuentra en la copia: están en este "
                              "informe."]


def marcas_en(p):
    """[(dónde empieza, texto marcado)] de los comentarios del párrafo."""
    marcas, abiertos, posicion = {}, set(), 0
    for el in p.iter(w("commentRangeStart"), w("commentRangeEnd"), w("t")):
        if el.tag == w("commentRangeStart"):
            abiertos.add(el.get(w("id")))
            marcas[el.get(w("id"))] = [posicion, ""]
        elif el.tag == w("commentRangeEnd"):
            abiertos.discard(el.get(w("id")))
        else:
            for i in abiertos:
                marcas[i][1] += el.text or ""
            posicion += len(el.text or "")
    return sorted(tuple(m) for m in marcas.values())


def test_una_referencia_sin_su_texto_se_marca_entera():
    doc = docx.Document()
    p = doc.add_paragraph("  Molina, R. (2017). Turismo. Editorial Ejemplo.")._p
    comentarios.comentar(doc, revision_con(doc, Ancla(p, "Madrid", referencia=True)), Informe())
    assert marcas_en(p) == [(2, "Molina, R. (2017). Turismo. Editorial Ejemplo.")]   # sin los espacios de delante


def test_si_se_repite_se_marca_la_mas_cercana_o_la_ultima():
    doc = docx.Document()
    p = doc.add_paragraph("Madrid (Pérez 2019) y otra vez (Pérez 2019) en Madrid.")._p
    comentarios.comentar(doc, revision_con(doc, Ancla(p, "(Pérez 2019)", 30), Ancla(p, "Madrid", ultima=True)),
                         Informe())
    assert marcas_en(p) == [(31, "(Pérez 2019)"), (47, "Madrid")]   # la segunda «(Pérez 2019)» y el último «Madrid»


def test_de_principio_a_fin_cada_comentario_en_su_sitio(tmp_path):
    """La misma cita mal escrita dos veces en un párrafo: un comentario en cada una. La ciudad, la del final."""
    ruta = tmp_path / "trabajo.docx"
    doc = docx.Document()
    doc.add_heading("Marco teórico", level=1)
    doc.add_paragraph("El turismo rural es un fenómeno muy estudiado (Molina 2017), y la despoblación también lo es "
                      "desde hace décadas en muchas comarcas del interior (Molina 2017).")
    doc.add_heading("Referencias", level=1)
    doc.add_paragraph("Molina, R. (2017). Madrid y el turismo rural. Editorial Ejemplo, Madrid.")
    doc.save(str(ruta))
    copia, _ = pipeline.procesar(ruta, tmp_path / "salida", CON_COMENTARIOS)
    parrafos = list(docx.Document(str(copia)).element.body.iter(w("p")))
    con_marcas = [marcas_en(p) for p in parrafos if marcas_en(p)]
    texto, referencia = con_marcas
    assert [m for _, m in texto] == ["(Molina 2017)", "(Molina 2017)"] and texto[0][0] != texto[1][0]
    assert referencia == [(len("Molina, R. (2017). Madrid y el turismo rural. Editorial Ejemplo, "), "Madrid")]


def test_texto_del_comentario():
    assert comentarios.texto_del_comentario("Falta el año en la lista de referencias.",
                                            {"explicacion": "en el texto la citas como de 2015",
                                             "correccion": "Herranz, M. (2015)."}) == (
        "APA Fácil: Falta el año en la lista de referencias. En este caso, en el texto la citas como de 2015. "
        "Corrección: «Herranz, M. (2015).».")


# --- Línea de comandos y JSON ---------------------------------------------------------

def test_opcion_en_la_linea_de_comandos(tmp_path, capsys):
    assert main([str(PRUEBAS["03"]), "--comentarios"], carpeta_salida=tmp_path) == 0
    salida = capsys.readouterr().out
    assert "se han puesto 1 comentario en la copia" in salida
    copia = tmp_path / "03_trabajo_imagenes_listas_APA.docx"
    assert [m for m, _ in comentarios_de(copia)] == ["Campos, F., & Vidal, N."]
    import json
    assert json.loads(copia.with_suffix(".json").read_text(encoding="utf-8"))["opciones"]["comentarios"] is True


def test_la_ayuda_y_el_diagnostico(capsys):
    with pytest.raises(SystemExit):
        main(["--help"])
    ayuda = " ".join(capsys.readouterr().out.split())
    assert "--comentarios" in ayuda and "como comentarios de Word en la copia" in ayuda
    assert main([str(PRUEBAS["03"]), "--diagnostico", "--comentarios"]) == 0
    assert "la opción --comentarios no se usa en el modo diagnóstico" in capsys.readouterr().out


# --- Con Word ---------------------------------------------------------------------------

SCRIPT = r"""
$ErrorActionPreference = 'Stop'
[Console]::OutputEncoding = [Text.Encoding]::UTF8   # «APA Fácil» y las citas llevan tildes
$tipo = [System.__ComObject]
$word = New-Object -ComObject Word.Application
try {
    $tipo.InvokeMember('Visible', 'SetProperty', $null, $word, @($false))
    $tipo.InvokeMember('DisplayAlerts', 'SetProperty', $null, $word, @(0))
    $documentos = $tipo.InvokeMember('Documents', 'GetProperty', $null, $word, $null)
    $doc = $tipo.InvokeMember('Open', 'InvokeMethod', $null, $documentos, @($args[0], $false, $true, $false))
    try {
        $coms = $tipo.InvokeMember('Comments', 'GetProperty', $null, $doc, $null)
        $n = $tipo.InvokeMember('Count', 'GetProperty', $null, $coms, $null)
        for ($i = 1; $i -le $n; $i++) {
            $c = $tipo.InvokeMember('Item', 'InvokeMethod', $null, $coms, @($i))
            $alcance = $tipo.InvokeMember('Scope', 'GetProperty', $null, $c, $null)
            $marcado = $tipo.InvokeMember('Text', 'GetProperty', $null, $alcance, $null)
            $autor = $tipo.InvokeMember('Author', 'GetProperty', $null, $c, $null)
            "COMENTARIO|$autor|$marcado"
        }
    } finally {
        $tipo.InvokeMember('Close', 'InvokeMethod', $null, $doc, @(0)) | Out-Null
    }
} finally {
    $tipo.InvokeMember('Quit', 'InvokeMethod', $null, $word, @(0)) | Out-Null
    [System.Runtime.InteropServices.Marshal]::ReleaseComObject($word) | Out-Null
}
"""


@pytest.mark.skipif(not a_pdf.hay_word(), reason="Word no está instalado")
def test_word_ve_los_comentarios_en_su_sitio(tmp_path):
    copia, informe = pipeline.procesar(PRUEBAS["04"], tmp_path, CON_COMENTARIOS)
    with tempfile.TemporaryDirectory() as carpeta:
        script = Path(carpeta) / "comentarios.ps1"
        script.write_text(SCRIPT, encoding="utf-8-sig")
        resultado = a_pdf._powershell("-File", str(script), str(copia), tiempo=120)
    salida = resultado.stdout.decode("utf-8", "replace").splitlines()
    filas = [linea.split("|", 2)[1:] for linea in salida if linea.startswith("COMENTARIO|")]
    assert len(filas) == len(casos_del_informe(informe)), resultado.stderr.decode("utf-8", "replace")
    assert {autor for autor, _ in filas} == {"APA Fácil"}
    assert [marcado for _, marcado in filas][:3] == ["(Martínez 2019)", "(Pérez et. al, 2020)", "Gómez (s.f)"]
