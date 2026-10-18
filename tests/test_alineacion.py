"""
Paso 1.4: texto alineado a la izquierda, sin justificar y sin guiones automáticos.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

from collections import Counter
from pathlib import Path

import docx
import pytest
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement

from apa_facil import estructura, pipeline
from apa_facil.documento import Estilos, editar_partes_de_texto, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import alineacion

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
IZQUIERDA, CENTRO, DERECHA, JUSTIFICADO = (WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                                           WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.JUSTIFY)
LARGO = "palabra " * 45   # un párrafo largo: ya no es portada


def aplicar(doc):
    informe = Informe()
    alineacion.aplicar(doc, {}, informe)
    return informe


def alineado(doc, p):
    return Estilos(doc).parrafo(p._p if hasattr(p, "_p") else p, "jc", "val") or "left"


def parrafo(doc, texto, alineacion_=None):
    p = doc.add_paragraph(texto)
    p.alignment = alineacion_
    return p


# --- Justificado ---------------------------------------------------------------------

@pytest.mark.parametrize("valor", ["both", "distribute", "mediumKashida", "thaiDistribute"])
def test_justificado_pasa_a_la_izquierda(valor):
    doc = docx.Document()
    p = parrafo(doc, LARGO)
    p._p.get_or_add_pPr().get_or_add_jc().set(w("val"), valor)
    informe = aplicar(doc)
    assert alineado(doc, p) == "left"
    assert informe.cambios == ["Alineación: 1 párrafo justificado pasa a estar alineado a la izquierda "
                               "(APA no admite texto justificado)."]


def test_estilo_justificado():
    doc = docx.Document()
    doc.styles["Normal"].paragraph_format.alignment = JUSTIFICADO
    parrafos = [parrafo(doc, LARGO) for _ in range(3)]
    informe = aplicar(doc)
    assert doc.styles["Normal"].paragraph_format.alignment == IZQUIERDA
    assert {alineado(doc, p) for p in parrafos} == {"left"}
    assert informe.cambios[0].startswith("Alineación: 3 párrafos justificados pasan")


def test_justificado_en_tablas():
    doc = docx.Document()
    celda = doc.add_table(rows=1, cols=2).rows[0].cells
    celda[0].paragraphs[0].alignment = JUSTIFICADO
    celda[1].paragraphs[0].alignment = CENTRO
    aplicar(doc)
    assert alineado(doc, celda[0].paragraphs[0]) == "left"
    assert alineado(doc, celda[1].paragraphs[0]) == "center", "lo centrado en tablas se queda"


# --- Centrado y derecha ------------------------------------------------------------------

def test_centrado_no_se_toca():
    doc = docx.Document()
    titulo = parrafo(doc, "Título del trabajo", CENTRO)
    parrafo(doc, LARGO)
    imagen = parrafo(doc, "", CENTRO)
    assert aplicar(doc).cambios == []
    assert alineado(doc, titulo) == alineado(doc, imagen) == "center"


def test_derecha_en_el_texto_pasa_a_la_izquierda():
    doc = docx.Document()
    parrafo(doc, LARGO)
    p = parrafo(doc, "Más información en la web", DERECHA)
    informe = aplicar(doc)
    assert alineado(doc, p) == "left"
    assert informe.cambios == ["Alineación: 1 párrafo alineado a la derecha pasa a la izquierda "
                               "(«Más información en la web»)."]
    assert informe.avisos == []


def test_derecha_en_la_portada_se_respeta():
    doc = docx.Document()
    parrafo(doc, "UNIVERSIDAD", CENTRO)
    autora = parrafo(doc, "Autora: Ana Gil", DERECHA)
    tutor = parrafo(doc, "Tutor: Luis Paz", DERECHA)
    tutor.add_run().add_break(WD_BREAK.PAGE)
    resto = parrafo(doc, "Fecha: junio", DERECHA)
    informe = aplicar(doc)
    assert alineado(doc, autora) == alineado(doc, tutor) == "right"
    assert alineado(doc, resto) == "left", "después del salto de página ya no es portada"
    assert len(informe.avisos) == 1
    assert informe.avisos[0].startswith("Al principio del documento hay 2 párrafos alineados a la derecha "
                                        "(«Autora: Ana Gil»). Parecen parte de la portada")


def test_un_parrafo_largo_acaba_la_portada():
    """Aunque no haya salto de página, después de un párrafo largo ya no es portada (documento 04)."""
    doc = docx.Document()
    parrafo(doc, "Título", CENTRO)
    parrafo(doc, LARGO)
    p = parrafo(doc, "Más información", DERECHA)
    aplicar(doc)
    assert alineado(doc, p) == "left"


def test_un_salto_de_seccion_acaba_la_portada():
    doc = docx.Document()
    parrafo(doc, "Autora: Ana Gil", DERECHA)
    doc.add_section(WD_SECTION.NEW_PAGE)
    p = parrafo(doc, "Fecha", DERECHA)
    aplicar(doc)
    assert alineado(doc, doc.paragraphs[0]) == "right"
    assert alineado(doc, p) == "left"


def test_salto_antes_acaba_la_portada():
    doc = docx.Document()
    parrafo(doc, "Autora: Ana Gil", DERECHA)
    p = parrafo(doc, "Fecha", DERECHA)
    p.paragraph_format.page_break_before = True
    aplicar(doc)
    assert alineado(doc, p) == "left"


# --- Lo que no se toca y los guiones ----------------------------------------------------------

def test_encabezado_y_pie_no_cambian():
    doc = docx.Document()
    doc.styles["Normal"].paragraph_format.alignment = JUSTIFICADO
    parrafo(doc, LARGO)
    encabezado = doc.sections[0].header.paragraphs[0]
    encabezado.text = "Universidad"
    pie = doc.sections[0].footer.paragraphs[0]
    pie.text = "1"
    pie.alignment = DERECHA
    aplicar(doc)
    assert alineado(doc, encabezado) == "both", "el encabezado heredaba el justificado y lo conserva"
    assert alineado(doc, pie) == "right"


def test_guiones_automaticos_se_desactivan():
    doc = docx.Document()
    parrafo(doc, LARGO)
    doc.settings.element.append(OxmlElement("w:autoHyphenation"))
    informe = aplicar(doc)
    assert doc.settings.element.find(w("autoHyphenation")) is None
    assert "Se han desactivado los guiones automáticos a final de línea (APA no los admite)." in informe.cambios


def test_guiones_opcionales_solo_se_avisan():
    doc = docx.Document()
    p = parrafo(doc, "Hidro")
    run = p.add_run()
    run._r.append(OxmlElement("w:softHyphen"))
    p.add_run("electricidad")
    informe = aplicar(doc)
    assert next(doc.element.body.iter(w("softHyphen")), None) is not None, "no se modifica el texto"
    assert informe.avisos == ["El texto tiene 1 guion opcional, que parten palabras a final de línea. APA no "
                              "los admite: búscalos en Word (Ctrl+B, «Más», «Especial», «Guion opcional») y "
                              "quítalos."]


# --- Documentos de prueba -----------------------------------------------------------------

def _alineaciones(doc):
    estilos, cuenta = Estilos(doc), Counter()
    for p in doc.element.body.iter(w("p")):
        zona = "tabla" if next(p.iterancestors(w("tc")), None) is not None else "cuerpo"
        cuenta[(zona, estilos.parrafo(p, "jc", "val") or "left")] += 1

    def ver(nombre, raiz):
        if not nombre.endswith("document.xml"):
            for p in raiz.iter(w("p")):
                cuenta[(nombre, estilos.parrafo(p, "jc", "val") or "left")] += 1
    editar_partes_de_texto(doc, ver)
    return cuenta


def _centrados(doc):
    """Los textos de los párrafos centrados del cuerpo y de las tablas, con cuántas veces sale cada uno."""
    estilos = Estilos(doc)
    return Counter(" ".join(texto_parrafo(p).split()) for p in doc.element.body.iter(w("p"))
                   if (estilos.parrafo(p, "jc", "val") or "left") == "center")


DERECHA_ESPERADA = {"05": 3}   # las tres líneas de la portada del TFG


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    antes = _alineaciones(docx.Document(str(original)))
    salida, informe = pipeline.procesar(original, tmp_path)
    despues = _alineaciones(docx.Document(str(salida)))

    assert not [k for k in despues if k[1] in alineacion.JUSTIFICADOS]
    assert despues[("cuerpo", "right")] == DERECHA_ESPERADA.get(original.name[:2], 0)
    # Lo centrado sigue centrado; lo único que se centra son los títulos de nivel 1 (paso 4.1).
    centrados_antes, centrados_despues = _centrados(docx.Document(str(original))), _centrados(docx.Document(str(salida)))
    assert not centrados_antes - centrados_despues, "lo centrado no cambia"
    nivel_1 = {" ".join(c.texto.split()) for c in estructura.clasificar(docx.Document(str(salida)))
               if c.tipo in ("titulo", "titulo_referencias", "anexo") and (c.nivel or 1) == 1}
    assert set(centrados_despues - centrados_antes) <= nivel_1
    for clave, n in antes.items():
        if clave[0].endswith(("footnotes.xml", "endnotes.xml")):   # encabezados y pies: paso 1.6
            assert despues[clave] == n, f"las notas no cambian: {clave}"
    avisos_portada = [a for a in informe.avisos if "portada" in a]
    assert len(avisos_portada) == (1 if original.name.startswith("05") else 0)
