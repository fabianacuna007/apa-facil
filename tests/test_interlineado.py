"""
Paso 1.3: interlineado doble, sin espacio antes ni después del párrafo.

Los casos concretos se prueban con documentos creados aquí mismo; después se
comprueba la regla con los 5 documentos de pruebas/ (que no se modifican).
"""

from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.shared import Pt

from apa_facil import pipeline
from apa_facil.documento import Estilos, editar_partes_de_texto, w
from apa_facil.informe import Informe
from apa_facil.reglas import interlineado
from apa_facil.reglas.interlineado import _efectivo

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
DOBLE = (480, "auto", 0, 0)
SENCILLO = (240, "auto", 0, 0)


def aplicar(doc):
    informe = Informe()
    interlineado.aplicar(doc, {}, informe)
    return informe


def espaciado(doc, p):
    return _efectivo(Estilos(doc), p._p if hasattr(p, "_p") else p)


def cuerpo(doc):
    return [p for p in doc.element.body.iter(w("p"))
            if next(p.iterancestors(w("tc"), w("txbxContent")), None) is None]


def celdas(doc):
    return [p for p in doc.element.body.iter(w("p")) if next(p.iterancestors(w("tc")), None) is not None]


def parrafo(doc, texto, linea=None, regla=None, antes=None, despues=None):
    p = doc.add_paragraph(texto)
    formato = p.paragraph_format
    if linea is not None:
        formato.line_spacing = linea
    if regla is not None:
        formato.line_spacing_rule = regla
    if antes is not None:
        formato.space_before = Pt(antes)
    if despues is not None:
        formato.space_after = Pt(despues)
    return p


def tabla(doc, linea=None, regla=None):
    t = doc.add_table(rows=2, cols=2)
    for celda in t._cells:
        formato = celda.paragraphs[0].paragraph_format
        if linea is not None:
            formato.line_spacing = linea
        if regla is not None:
            formato.line_spacing_rule = regla
        formato.space_after = Pt(6)
    return t


# --- Cuerpo -----------------------------------------------------------------------

@pytest.mark.parametrize("linea, regla", [
    (1.0, None), (1.15, None), (1.5, None), (Pt(14), WD_LINE_SPACING.EXACTLY), (Pt(12), WD_LINE_SPACING.AT_LEAST),
])
def test_todo_pasa_a_doble(linea, regla):
    doc = docx.Document()
    parrafo(doc, "Texto", linea, regla, antes=6, despues=10)
    parrafo(doc, "Otro párrafo sin formato propio")
    aplicar(doc)
    assert {espaciado(doc, p) for p in cuerpo(doc)} == {DOBLE}


def test_titulos_sin_espacio_antes():
    """El estilo de título de Word lleva espacio antes; APA no lo quiere."""
    doc = docx.Document()
    doc.add_heading("Método", level=1)
    doc.add_heading("Participantes", level=2)
    parrafo(doc, "Texto")
    aplicar(doc)
    assert {espaciado(doc, p) for p in cuerpo(doc)} == {DOBLE}


def test_quita_lo_que_anula_el_espaciado():
    doc = docx.Document()
    p = parrafo(doc, "Pegado de una web")
    espacio = p._p.get_or_add_pPr().get_or_add_spacing()
    for atributo, valor in (("beforeLines", "100"), ("afterLines", "100"),
                            ("beforeAutospacing", "1"), ("afterAutospacing", "1")):
        espacio.set(w(atributo), valor)
    aplicar(doc)
    for atributo in ("beforeLines", "afterLines", "beforeAutospacing", "afterAutospacing"):
        assert espacio.get(w(atributo)) is None
    assert espaciado(doc, p) == DOBLE


def test_informe_del_cuerpo():
    doc = docx.Document()
    parrafo(doc, "Uno", 1.5, despues=12)
    parrafo(doc, "Dos", 1.0, antes=30)
    informe = aplicar(doc)
    assert informe.cambios[0] == ("Interlineado: todo el texto pasa a doble espacio, sin espacio extra antes ni "
                                  "después de los párrafos (antes: 1,5 y sencillo, con hasta 30 pt de "
                                  "separación entre párrafos).")


def test_si_ya_cumple_no_hay_cambio():
    doc = docx.Document()
    parrafo(doc, "Uno", 2.0, antes=0, despues=0)
    estilo = doc.styles["Normal"].paragraph_format
    estilo.line_spacing, estilo.space_before, estilo.space_after = 2.0, Pt(0), Pt(0)
    assert aplicar(doc).cambios == []


# --- Tablas -------------------------------------------------------------------------

@pytest.mark.parametrize("linea, regla, esperado", [
    (1.0, None, 240), (1.5, None, 360), (2.0, None, 480),        # admitidos: se mantienen
    (1.15, None, 240), (Pt(12), WD_LINE_SPACING.EXACTLY, 240),   # otros: sencillo
])
def test_tablas(linea, regla, esperado):
    doc = docx.Document()
    parrafo(doc, "Texto")
    tabla(doc, linea, regla)
    aplicar(doc)
    assert {espaciado(doc, p) for p in celdas(doc)} == {(esperado, "auto", 0, 0)}


def test_tabla_sin_formato_no_hereda_el_doble():
    """Una tabla que heredaba el sencillo del Normal debe seguir sencilla."""
    doc = docx.Document()
    doc.styles["Normal"].paragraph_format.line_spacing = 1.0
    tabla(doc)
    aplicar(doc)
    assert {espaciado(doc, p) for p in celdas(doc)} == {SENCILLO}


def test_informe_de_tablas():
    doc = docx.Document()
    tabla(doc, 1.15)
    tabla(doc, 1.5)
    informe = aplicar(doc)
    assert informe.cambios[-1] == ("Tablas: se mantiene el interlineado 1,5, que APA admite; el interlineado 1,15 "
                                   "pasa a sencillo (APA admite sencillo, 1,5 o doble); se quita el espacio "
                                   "extra entre párrafos.")


# --- Lo que no se toca --------------------------------------------------------------------

def test_encabezado_y_pie_no_cambian():
    doc = docx.Document()
    parrafo(doc, "Texto", 1.0)
    encabezado = doc.sections[0].header.paragraphs[0]
    encabezado.text = "1"
    antes = espaciado(doc, encabezado)
    aplicar(doc)
    assert espaciado(doc, encabezado) == antes


def test_estilos_de_notas_quedan_fijados():
    doc = docx.Document()
    estilo = doc.styles.add_style("footnote text", 1)
    estilo.base_style = doc.styles["Normal"]
    aplicar(doc)
    sid = estilo.element.get(w("styleId"))
    assert interlineado._efectivo_de_estilo(Estilos(doc), sid)[0] != 480


def test_cuadros_de_texto_no_cambian():
    doc = docx.Document()
    parrafo(doc, "Texto", 1.0)
    caja = OxmlElement("w:txbxContent")
    p_caja = OxmlElement("w:p")
    caja.append(p_caja)
    doc.paragraphs[0]._p.append(caja)   # basta con que esté dentro de un w:txbxContent
    antes = espaciado(doc, p_caja)
    aplicar(doc)
    assert espaciado(doc, p_caja) == antes


# --- Documentos de prueba -----------------------------------------------------------------

TABLAS = {"02": {240}, "04": {240}, "05": {360}}


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    orig = docx.Document(str(original))
    otras_antes = _otras_partes(orig)
    salida, informe = pipeline.procesar(original, tmp_path)
    copia = docx.Document(str(salida))

    assert {espaciado(copia, p) for p in cuerpo(copia)} == {DOBLE}
    lineas = {espaciado(copia, p)[0] for p in celdas(copia)}
    assert lineas == TABLAS.get(original.name[:2], set())
    assert _otras_partes(copia) == otras_antes, "notas, encabezados y pies no cambian"
    assert any(c.startswith("Interlineado:") for c in informe.cambios)


def _otras_partes(doc):
    estilos, resultado = Estilos(doc), []

    def ver(nombre, raiz):
        if not nombre.endswith("document.xml"):
            resultado.extend((nombre, _efectivo(estilos, p)) for p in raiz.iter(w("p")))
    editar_partes_de_texto(doc, ver)
    return sorted(resultado)


def test_indice_se_define_a_doble_espacio():
    """Word rehace el índice con sus estilos "toc N" (5 pt después) si el documento no los define."""
    doc = docx.Document()
    p = parrafo(doc, "")
    run = OxmlElement("w:r")
    codigo = OxmlElement("w:instrText")
    codigo.text = r' TOC \o "1-3" \h '
    run.append(codigo)
    p._p.append(run)
    aplicar(doc)
    estilos = Estilos(doc)
    por_nombre = {estilos.nombre(i).lower(): i for i in estilos.por_id}
    for nivel in range(1, 10):
        sid = por_nombre[f"toc {nivel}"]
        assert interlineado._efectivo_de_estilo(estilos, sid) == DOBLE
    sangria = estilos.por_id[por_nombre["toc 3"]].find(f"{w('pPr')}/{w('ind')}")
    assert sangria.get(w("left")) == "440"


def test_sin_indice_no_se_crean_estilos_de_indice():
    doc = docx.Document()
    parrafo(doc, "Texto")
    aplicar(doc)
    nombres = {Estilos(doc).nombre(e).lower() for e in Estilos(doc).por_id}
    assert not nombres & {f"toc {nivel}" for nivel in range(1, 10)}
