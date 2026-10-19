"""
Paso 2.2: la cadena completa con cada documento de pruebas/.

Cada documento se procesa una sola vez (en una carpeta temporal, nunca en
salida/) y sobre esa copia se comprueba:
1. que se procesa sin errores;
2. que el contenido es idéntico al del original (texto, tablas, imágenes,
   notas y secciones), leyendo la copia desde el disco, salvo los cambios
   previstos que las reglas anotan en el informe (paso 3.5);
3. que el verificador no encuentra ningún incumplimiento (fase 1 y estilos de título), y que se
   genera su informe en JSON (paso 2.3);
4. que el original no se ha modificado;
5. que las excepciones admitidas son exactamente las esperadas. Si un detector
   empezara a "perdonar" párrafos de más, el verificador seguiría diciendo
   "cumple": esta lista lo impide.
"""

import hashlib
import json
from pathlib import Path

import docx
import pytest

from apa_facil import estructura, pipeline
from apa_facil.documento import ROTULO_DE_APENDICE, Estilos, comparar, radiografia
from herramientas.verificar import (ALINEACION, APENDICES, CITAS, INDICE, PORTADA_PROPIA, TEXTO_TITULOS, TITULO_TRABAJO, mostrar,
                                   verificar)

CARPETA = Path(__file__).resolve().parent.parent / "pruebas"
PRUEBAS = sorted(CARPETA.glob("*.docx"))

# Excepciones admitidas que debe tener cada copia: (regla, número de párrafo).
# Al añadir un documento a pruebas/, revisa a mano sus excepciones y apúntalas aquí.
# Desde el paso 4.3, los títulos numerados y en mayúsculas (que por defecto se dejan) también; desde
# el 4.4, el rótulo «Introducción» y el título del trabajo en mayúsculas. Los números son los de la
# copia, que lleva el título del trabajo añadido al principio del texto (por eso van uno por delante).
# Desde el 5.3, sin datos de portada se mantiene la de la universidad: también es una excepción admitida.
EXCEPCIONES = {
    "01_ensayo_basico": {(TEXTO_TITULOS, n) for n in (10, 13, 16)}   # «1.», «2.» y «3.»
                        | {(TITULO_TRABAJO, 7), (PORTADA_PROPIA, 1)},  # «Introducción»; la portada
    "02_informe_tablas": {(TITULO_TRABAJO, 6), (PORTADA_PROPIA, 1)},   # «Introducción»; la portada
    # «1.» a «4.» y «2.1.»; desde el 6.2, «2.1.» y lo que sigue van uno o dos más allá: el título de cada figura
    # va en su propio párrafo
    "03_trabajo_imagenes_listas": {(TEXTO_TITULOS, n) for n in (4, 10, 15, 20, 28)}
                                  | {(TITULO_TRABAJO, 4), (PORTADA_PROPIA, 1)},       # «1. Introducción»; la portada
    "04_formato_caotico": {(PORTADA_PROPIA, 1)},
    "05_tfg_completo": {(ALINEACION, 8), (ALINEACION, 9), (ALINEACION, 10)}   # portada: Autora, Tutor, Convocatoria
                       # «4.2.» y «5.» van uno más allá desde el 6.1: el título de la tabla 1 va en su propio párrafo
                       | {(TEXTO_TITULOS, n) for n in (20, 23, 24, 26, 27, 29, 31, 35, 37, 38, 64, 66)}
                       | {(TITULO_TRABAJO, 20), (PORTADA_PROPIA, 2)}          # «1. Introducción»; la portada
                       | {(INDICE, 16)}                                       # el índice (paso 5.4)
                       | {(CITAS, 33)}                                        # la cita, con sus comillas (6.3)
                       | {(APENDICES, 75)},                                   # «Anexo I», no «Apéndice» (6.7)
}


def huella(ruta):
    return hashlib.sha256(Path(ruta).read_bytes()).hexdigest()


@pytest.fixture(scope="module", params=PRUEBAS, ids=lambda p: p.stem)
def procesado(request, tmp_path_factory):
    """Procesa un documento una vez y guarda lo necesario para las comprobaciones."""
    original = request.param
    antes = huella(original)
    salida, informe = pipeline.procesar(original, tmp_path_factory.mktemp(original.stem))
    return {"original": original, "huella_antes": antes, "salida": salida, "informe": informe,
            "verificacion": verificar(salida)}


def test_se_procesa_sin_errores(procesado):
    assert procesado["salida"].exists()
    assert procesado["informe"].errores == []


def test_el_contenido_es_identico(procesado):
    antes = radiografia(docx.Document(str(procesado["original"])))
    despues = radiografia(docx.Document(str(procesado["salida"])))
    diferencias = comparar(antes, despues, procesado["informe"].cambios_previstos)
    assert diferencias == [], "\n".join(diferencias)


def test_cumple_todas_las_reglas(procesado):
    resultado = procesado["verificacion"]
    assert resultado.cumple, mostrar(resultado)


def test_la_copia_lleva_su_estructura(procesado):
    """
    Desde el paso 4.1 cada título lleva su estilo de Word, así que la copia se
    clasifica igual que el original (el verificador solo tiene la copia).
    """
    def titulos(ruta):
        doc = docx.Document(str(ruta))
        estilos = Estilos(doc)
        # Sin el título del trabajo que el paso 4.4 añade al principio del texto (lo comprueba el verificador).
        # Los de nivel 4 y 5 terminan en punto desde el paso 4.5, las mayúsculas cambian en el 4.6, el rótulo de las
        # referencias pasa a «Referencias» en el 6.4 y el del apéndice se separa de su título en el 6.7.
        def texto(c):
            if c.tipo == "titulo_referencias":
                return "rótulo"
            if c.tipo == "anexo":
                return ROTULO_DE_APENDICE.match(c.texto)["rotulo"].lower()
            return (c.texto.rstrip(".") if (c.nivel or 0) >= 4 else c.texto).lower()
        return [(c.tipo, c.nivel, texto(c)) for c in estructura.clasificar(doc)
                if c.tipo in ("titulo", "titulo_referencias", "anexo")
                and not c.motivo.startswith("título del apéndice")
                and estilos.nombre(estilos.estilo_de(c.elemento)).lower() != "title"]
    assert titulos(procesado["salida"]) == titulos(procesado["original"])


def test_el_original_no_se_toca(procesado):
    assert huella(procesado["original"]) == procesado["huella_antes"]
    assert procesado["salida"].resolve().parent != CARPETA.resolve()


def test_excepciones_esperadas(procesado):
    nombre = procesado["original"].stem
    if nombre not in EXCEPCIONES:
        pytest.fail(f"«{nombre}» no está en EXCEPCIONES: revisa a mano sus excepciones con "
                    f"herramientas/verificar.py y apúntalas en {Path(__file__).name}.")
    halladas = {(h.regla, h.parrafo) for h in procesado["verificacion"].excepciones}
    assert halladas == EXCEPCIONES[nombre], mostrar(procesado["verificacion"])


def test_genera_su_informe(procesado):
    ruta = procesado["salida"].with_suffix(".json")
    datos = json.loads(ruta.read_text(encoding="utf-8"))
    assert datos["archivo"] == procesado["original"].name
    assert datos["resumen"] == {"cambios": len(procesado["informe"].cambios),
                                "avisos": len(procesado["informe"].avisos), "errores": 0}
    assert datos["resumen"]["cambios"] > 0   # ningún original cumplía APA


def test_estan_todos_los_documentos_de_la_lista():
    # Si se borra o se renombra un documento, la lista de excepciones tiene que acompañarlo.
    assert {p.stem for p in PRUEBAS} == set(EXCEPCIONES)


# Paso 4.3: con --quitar-numeracion y --quitar-mayusculas, el texto de los títulos cambia,
# pero solo eso (la integridad lo comprueba) y desaparecen sus excepciones.
# Numeraciones quitadas y títulos con otras mayúsculas (pasos 4.3 y 4.6): cada cambio, una vez.
TITULOS_CAMBIADOS = {"01_ensayo_basico": 3 + 3, "02_informe_tablas": 0, "03_trabajo_imagenes_listas": 5 + 3,
                     "04_formato_caotico": 0 + 3 + 1, "05_tfg_completo": 12 + 8}   # 04: y el título de la tabla (6.1)


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_con_las_opciones_de_los_titulos(original, tmp_path):
    opciones = {"quitar_numeracion": True}   # las mayúsculas de APA, por defecto (paso 4.6)
    salida, informe = pipeline.procesar(original, tmp_path, opciones)   # no se entrega si la integridad falla
    resultado = verificar(salida)
    assert resultado.cumple, mostrar(resultado)
    assert not [h for h in resultado.excepciones if h.regla == TEXTO_TITULOS]
    acciones = [c for c in informe.cambios_previstos if c.accion in ("quitar_numeracion", "mayusculas_apa")]
    assert len(acciones) == TITULOS_CAMBIADOS[original.stem]
    diferencias = comparar(radiografia(docx.Document(str(original))), radiografia(docx.Document(str(salida))),
                           informe.cambios_previstos)
    assert diferencias == []
    sin_cambios_previstos = comparar(radiografia(docx.Document(str(original))),
                                     radiografia(docx.Document(str(salida))),
                                     [c for c in informe.cambios_previstos if c not in acciones])
    assert bool(sin_cambios_previstos) == bool(acciones), "sin anotarlos, la integridad los vería"


# Paso 7.5: los avisos de «Citas y referencias» (la tabla de la fase 7 del ROADMAP), por el principio de su
# mensaje y con cuántos casos. Lo que ya avisa el paso 6.5 (el orden alfabético, las referencias que son solo
# una dirección web) no se repite. El 02 es el caso de control: ninguno.
REVISION = {
    "01_ensayo_basico": [("Una obra citada en el texto no está", 1), ("Ciudad de la editorial", 2),
                         ("Año sin paréntesis", 1)],
    "02_informe_tablas": [],
    "03_trabajo_imagenes_listas": [("«&» en lugar de «y»", 1)],
    "04_formato_caotico": [("2 obras citadas en el texto no están", 3), ("Falta la coma antes del año", 1),
                           ("«et al.» mal escrito", 2), ("Fecha mal escrita", 3), ("Año sin paréntesis", 1),
                           ("Autores separados por «;»", 1), ("Falta el punto después de la fecha", 1),
                           ("Número de la revista con «nº»", 1), ("Referencia incompleta", 1)],
    "05_tfg_completo": [("Una obra citada en el texto no está", 1), ("Una obra de la lista de referencias no se cita", 1),
                        ("Falta el espacio después de «p.»", 1), ("Ciudad de la editorial", 1),
                        ("Falta el punto después del autor", 1), ("DOI con la forma antigua", 1),
                        ("Falta el año", 1)],
}


def test_avisos_de_citas_y_referencias(procesado):
    revision = [e for e in procesado["informe"].entradas if e.regla == "revision"]
    esperados = REVISION[procesado["original"].stem]
    assert len(revision) == len(esperados), [e.mensaje for e in revision]
    for entrada, (principio, casos) in zip(revision, esperados):
        assert entrada.mensaje.startswith(principio) and len(entrada.casos) == casos, entrada.mensaje
    datos = json.loads(procesado["salida"].with_suffix(".json").read_text(encoding="utf-8"))
    del_json = [a for a in datos["avisos"] if a["regla"] == "revision"]
    assert [len(a["casos"]) for a in del_json] == [casos for _, casos in esperados]
    assert all(a["nombre"] == "Citas y referencias" for a in del_json)
