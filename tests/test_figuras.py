"""
Paso 6.2: las figuras, según la web oficial de APA (Figure setup y Sample figures):
la misma disposición que las tablas (leyendas.py). «Figura N» en negrita encima y el
título en cursiva en la línea siguiente, a la izquierda y sin sangría; la nota
debajo, con «Nota.» en cursiva. «Gráfico 1» o «Imagen 2» pasan a «Figura 1» y
«Figura 2» (decisión del propietario), y también «Cuadro» a «Tabla».

Cada caso pasa por pipeline.procesar, que no entrega la copia si el contenido
cambia de una forma que no esté prevista.
"""

import base64
import io
from pathlib import Path

import docx
import pytest

from apa_facil import pipeline
from apa_facil.documento import CambioPrevisto, Estilos, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import figuras as regla
from apa_facil.reglas import tablas
from apa_facil.reglas.fuente import _runs_propios
from apa_facil.reglas.resumen import _activado, _poner
from herramientas.verificar import FIGURAS, verificar

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = {p.name[:2]: p for p in (RAIZ / "pruebas").glob("*.docx")}
# Una imagen PNG de 1 × 1 píxeles.
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==")
TEXTO = ("Los hogares rurales consumen más energía en invierno y en verano que los urbanos, según los datos de "
         "facturación recogidos durante dos años en sesenta viviendas.")


def trabajo(antes=(), despues=()):
    """Un apartado con una figura (un párrafo con solo una imagen) y, alrededor, los párrafos `antes` y `despues`."""
    doc = docx.Document()
    doc.add_heading("Resultados", level=1)
    doc.add_paragraph(TEXTO)
    for texto in antes:
        doc.add_paragraph(texto)
    doc.add_picture(io.BytesIO(PNG), width=docx.shared.Cm(8))
    for texto in despues:
        doc.add_paragraph(texto)
    doc.add_heading("Discusión", level=1)
    doc.add_paragraph(TEXTO)
    return doc


def procesar(doc, tmp_path, opciones=None):
    ruta = tmp_path / "trabajo.docx"
    doc.save(str(ruta))
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", opciones)
    return docx.Document(str(salida)), informe


def imagenes(doc):
    return [p for p in doc.element.body.iter(w("p")) if next(p.iter(w("drawing")), None) is not None
            and not texto_parrafo(p).strip()]


def alrededor(doc, k=0):
    """Los textos de los párrafos justo encima (dos) y justo debajo (uno) de la imagen k, sin los vacíos."""
    ps = list(doc.element.body.iter(w("p")))
    i = ps.index(imagenes(doc)[k])
    encima = [texto_parrafo(p) for p in ps[:i] if texto_parrafo(p).strip()]
    debajo = [texto_parrafo(p) for p in ps[i + 1:] if texto_parrafo(p).strip()]
    return encima[-2:], debajo[:1]


def parrafo(doc, inicio):
    return next(p for p in doc.element.body.iter(w("p")) if texto_parrafo(p).startswith(inicio))


def formato(doc, p):
    """[(texto, cursiva, negrita)] de cada fragmento con texto."""
    estilos = Estilos(doc)
    return [("".join(t.text or "" for t in r.iter(w("t"))), _activado(estilos, r, p, "i"),
             _activado(estilos, r, p, "b")) for r in _runs_propios(p)
            if "".join(t.text or "" for t in r.iter(w("t")))]


def incumplimientos(doc):
    return [h.mensaje for h in verificar(doc).incumplimientos if h.regla == FIGURAS]


# --- Los documentos de prueba ------------------------------------------------------------------

@pytest.fixture(scope="module")
def copia_03(tmp_path_factory):
    return pipeline.procesar(PRUEBAS["03"], tmp_path_factory.mktemp("figuras"))


def test_las_figuras_del_03(copia_03):
    salida, informe = copia_03
    doc = docx.Document(str(salida))
    assert alrededor(doc, 0)[0] == ["Figura 1", "Consumo Eléctrico Medio Mensual (kWh)"]
    assert alrededor(doc, 1)[0] == ["Figura 2", "Fuentes de Energía de los Hogares (Elaboración Propia)"]
    assert formato(doc, parrafo(doc, "Figura 1")) == [("Figura 1", False, True)]
    assert formato(doc, parrafo(doc, "Consumo Eléctrico")) == [("Consumo Eléctrico Medio Mensual (kWh)", True, False)]
    assert "Figuras: APA llama «Figura» a todas las figuras: «Gráfico 1» → «Figura 1», «Imagen 2» → «Figura 2». " \
           "Si el texto las nombra de otra forma («gráfico 1»), cámbialo tú." in informe.cambios
    assert any(a.startswith("El título de 1 figura lleva «(Elaboración Propia)»") for a in informe.avisos)


def test_la_copia_del_03_cumple_y_el_original_no(copia_03):
    salida, _ = copia_03
    assert incumplimientos(docx.Document(str(salida))) == []
    assert len(incumplimientos(docx.Document(str(PRUEBAS["03"])))) == 2


def test_el_logo_de_la_portada_no_es_una_figura(tmp_path):
    _, informe = pipeline.procesar(PRUEBAS["05"], tmp_path)
    assert not [e for e in informe.entradas if e.regla == "figuras"]


# --- Casos hechos a mano --------------------------------------------------------------------------

def test_leyenda_debajo_sube_cambia_de_palabra_y_se_separa(tmp_path):
    doc, _ = procesar(trabajo(despues=["Gráfico 1: consumo por meses", TEXTO]), tmp_path)
    assert alrededor(doc) == (["Figura 1", "Consumo por Meses"], [TEXTO])
    assert incumplimientos(doc) == []


def test_sin_numero_en_ninguna_figura_se_numeran(tmp_path):
    doc, informe = procesar(trabajo(antes=["Consumo por meses"]), tmp_path)
    assert alrededor(doc)[0] == ["Figura 1", "Consumo por Meses"]
    assert any(c.startswith("Figuras: se ha añadido el número a 1 figura") for c in informe.cambios)


def test_nota_debajo_de_la_figura(tmp_path):
    doc, _ = procesar(trabajo(antes=["Figura 1", "Consumo"], despues=["Fuente: elaboración propia."]), tmp_path)
    assert alrededor(doc)[1] == ["Nota. Fuente: elaboración propia."]
    assert formato(doc, parrafo(doc, "Nota."))[0] == ("Nota.", True, False)
    assert imagenes(doc)[0].find(f"{w('pPr')}/{w('keepNext')}") is not None, "la imagen, pegada a su nota"
    assert incumplimientos(doc) == []


def test_con_titulos_tal_cual(tmp_path):
    doc, _ = procesar(trabajo(antes=["Figura 1. Consumo por meses"]), tmp_path, {"titulos_tal_cual": True})
    assert alrededor(doc)[0] == ["Figura 1", "Consumo por meses"]


def test_una_imagen_en_medio_de_una_frase_no_es_una_figura():
    doc = docx.Document()
    p = doc.add_paragraph("El símbolo ")
    p.add_run().add_picture(io.BytesIO(PNG), width=docx.shared.Cm(0.5))
    p.add_run(" indica que la vivienda tiene placas solares, como se ve en el resto de los ejemplos del trabajo.")
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    assert informe.entradas == []


def test_un_solo_recordatorio_para_tablas_y_figuras():
    doc = trabajo(antes=["Figura 1", "Consumo"])
    doc.add_paragraph("Tabla 1")
    doc.add_paragraph("Notas")
    doc.add_table(rows=2, cols=2)
    informe = Informe()
    tablas.aplicar(doc, {}, informe)
    regla.aplicar(doc, {}, informe)
    recordatorios = [a for a in informe.avisos if a.startswith("Comprueba en Word")]
    assert recordatorios == ["Comprueba en Word que cada tabla y cada figura queda arriba o abajo de la página, y no "
                             "en medio del texto, como recomienda APA (también puede ir en una página para ella sola). "
                             "La herramienta no puede saberlo: depende de cómo reparte Word las páginas."]


def test_figura_sin_leyenda():
    informe = Informe()
    regla.aplicar(trabajo(), {}, informe)
    assert any(a.startswith("La figura 1 no lleva número ni título") for a in informe.avisos)


def test_el_verificador_ve_un_numero_sin_negrita(tmp_path):
    doc, _ = procesar(trabajo(antes=["Figura 1", "Consumo"]), tmp_path)
    for r in _runs_propios(parrafo(doc, "Figura 1")):
        _poner(r, "b", False)
    assert incumplimientos(doc) == ["El número de la figura: no va en negrita."]


# --- Integridad -------------------------------------------------------------------------------------

def _radiografias(antes, despues):
    """Radiografías de dos listas de párrafos: un texto, o None para un párrafo con solo una imagen."""
    def radio(lista):
        doc = docx.Document()
        for texto in lista:
            if texto is None:
                doc.add_picture(io.BytesIO(PNG), width=docx.shared.Cm(4))
            else:
                doc.add_paragraph(texto)
        return radiografia(doc)
    return radio(antes), radio(despues)


def test_integridad_admite_subir_a_una_imagen_y_cambiar_la_palabra():
    antes, despues = _radiografias(["Texto.", None, "Gráfico 1 – Consumo", "Fuente: INE."],
                                   ["Texto.", "Figura 1", "Consumo", None, "Nota. Fuente: INE."])
    previstos = [CambioPrevisto(3, "mover_leyenda", desde=(2,)),
                 CambioPrevisto(3, "cambiar_rotulo", quitado="Gráfico", nuevo="Figura"),
                 CambioPrevisto(3, "dividir_leyenda", quitado=" – ", nuevo="Consumo"),
                 CambioPrevisto(4, "añadir_nota", nuevo="Nota. ")]
    assert comparar(antes, despues, previstos) == []


@pytest.mark.parametrize("cambio", [
    CambioPrevisto(1, "cambiar_rotulo", quitado="Gráfico", nuevo="Tabla"),     # un gráfico no es una tabla
    CambioPrevisto(1, "cambiar_rotulo", quitado="Gráfico", nuevo="Esquema"),   # solo la palabra de APA
    CambioPrevisto(1, "cambiar_rotulo", quitado="Gráf", nuevo="Figura"),       # la palabra entera
])
def test_integridad_rechaza_otro_cambio_de_palabra(cambio):
    antes, despues = _radiografias(["Gráfico 1"], [cambio.nuevo + " 1"])
    assert comparar(antes, despues, [cambio])


def test_integridad_no_sube_a_un_parrafo_vacio_sin_imagen():
    antes, despues = _radiografias(["Texto.", "", "Figura 1"], ["Texto.", "Figura 1", ""])
    assert comparar(antes, despues, [CambioPrevisto(3, "mover_leyenda", desde=(2,))])


def test_integridad_no_pone_nota_lejos_de_una_tabla_o_imagen():
    antes, despues = _radiografias(["Texto.", "Fuente: INE."], ["Texto.", "Nota. Fuente: INE."])
    assert comparar(antes, despues, [CambioPrevisto(2, "añadir_nota", nuevo="Nota. ")])
