"""
Paso 9.7: el texto pegado se convierte en un .docx (un párrafo por línea) y se
procesa con el mismo núcleo, sin cambiar el texto.
"""

import re

import docx
import pytest

from apa_facil import datos_portada, desde_texto, estructura, pipeline
from apa_facil.documento import ErrorDocumento, w

TRABAJO = """El trabajo infantil en la Revolución Industrial

Introducción
La Revolución Industrial transformó la vida de las familias trabajadoras. Muchos niños empezaron a \
trabajar en fábricas y minas desde edades muy tempranas (Thompson, 1963).

1. Las condiciones de trabajo
Las jornadas superaban a menudo las doce horas diarias. Según Hobsbawm (1962), los salarios \
infantiles eran una parte importante de los ingresos del hogar.

Conclusiones
El trabajo infantil fue una pieza clave del primer capitalismo industrial.

Referencias
Thompson, E. P. (1963). The making of the English working class. Victor Gollancz.
Hobsbawm, E. J. (1962). The age of revolution: Europe 1789-1848. Weidenfeld & Nicolson.
"""


def crear(texto, tmp_path):
    return docx.Document(desde_texto.crear(texto, tmp_path / desde_texto.NOMBRE))


def test_una_linea_un_parrafo_y_el_texto_tal_cual(tmp_path):
    texto = "Título\r\n\r\n  Con espacios\ty un tabulador  \nOtra línea\rUna más\x0bTras Mayús+Intro\x0cTras un salto de página"
    doc = crear(texto, tmp_path)
    assert [p.text for p in doc.paragraphs] == [
        "Título", "  Con espacios\ty un tabulador  ", "Otra línea", "Una más", "Tras Mayús+Intro",
        "Tras un salto de página"]


def test_las_lineas_vacias_o_de_espacios_no_se_copian(tmp_path):
    doc = crear("Uno.\n\n   \n\t\nDos.", tmp_path)
    assert [p.text for p in doc.paragraphs] == ["Uno.", "Dos."]


def test_es_un_a4(tmp_path):
    seccion = crear("Un párrafo.", tmp_path).sections[0]
    assert (round(seccion.page_width.cm, 2), round(seccion.page_height.cm, 2)) == (21.0, 29.7)


def test_la_primera_linea_es_el_titulo(tmp_path):
    # Con un salto de página detrás, el núcleo la toma por la portada que solo tiene el título,
    # y «Introducción» es un título del texto, no parte de la portada.
    doc = crear(TRABAJO, tmp_path)
    assert doc.paragraphs[0]._p.find(f".//{w('br')}").get(w("type")) == "page"
    tipos = {c.texto: (c.zona, c.tipo) for c in estructura.clasificar(doc) if c.texto}
    assert tipos["El trabajo infantil en la Revolución Industrial"][0] == "portada"
    assert tipos["Introducción"] == ("cuerpo", "titulo")
    assert tipos["Conclusiones"] == ("cuerpo", "titulo")


@pytest.mark.parametrize("texto", [
    "Este trabajo empieza directamente con un párrafo.\nY sigue con otro.",   # acaba en punto
    "Una sola línea",                                                         # nada detrás
    " ".join(["palabra"] * 20) + "\nOtra línea",                              # demasiado larga
])
def test_sin_salto_si_no_empieza_por_un_titulo(texto, tmp_path):
    doc = crear(texto, tmp_path)
    assert doc.element.body.find(f".//{w('br')}") is None


@pytest.mark.parametrize("texto, mensaje", [
    ("", "No has pegado ningún texto"),
    (" \n\t\n ", "No has pegado ningún texto"),
    ("a" * (desde_texto.MAXIMO + 1), "el máximo es 500.000"),
    ("Hola\x07 mundo", "U+0007), después de «Hola»"),
], ids=["vacio", "solo-espacios", "demasiado-largo", "caracter-de-control"])
def test_textos_que_no_se_pueden_convertir(texto, mensaje, tmp_path):
    with pytest.raises(ErrorDocumento, match=re.escape(mensaje)):
        desde_texto.crear(texto, tmp_path / "x.docx")
    assert not (tmp_path / "x.docx").exists()


def test_se_procesa_con_el_nucleo_sin_perder_texto(tmp_path):
    entrada = desde_texto.crear(TRABAJO, tmp_path / desde_texto.NOMBRE)
    datos = {"titulo": "El trabajo infantil en la Revolución Industrial", "autores": ["Ana Pérez"]}
    salida, informe = pipeline.procesar(entrada, tmp_path / "salida", {"portada": datos_portada.combinar(None, datos)})
    # El pipeline ya comprueba la integridad del texto; aquí, que el cuerpo está entero y en orden.
    textos = [p.text for p in docx.Document(salida).paragraphs]
    cuerpo = textos.index("Introducción")
    assert textos[cuerpo - 1] == "El Trabajo Infantil en la Revolución Industrial"   # título del trabajo (4.4 y 4.6)
    assert textos[cuerpo + 1].startswith("La Revolución Industrial transformó")
    assert textos[-2:] == [   # las referencias, ordenadas (6.5)
        "Hobsbawm, E. J. (1962). The age of revolution: Europe 1789-1848. Weidenfeld & Nicolson.",
        "Thompson, E. P. (1963). The making of the English working class. Victor Gollancz."]
    assert not informe.errores


# --- Títulos del texto pegado (decisión del propietario, 2026-10-20: interpretar si es un título) ---

SIN_FORMATO = """El impacto de las redes sociales en la autoestima adolescente
Introducción
Las redes sociales forman parte de la vida diaria de los adolescentes. Según Pérez (2020), la mayoría las usa a diario.
Marco teórico
Antecedentes
Los primeros estudios se centraron en el tiempo de uso (Gómez y Ruiz, 2018). Más tarde se estudió el tipo de uso.
Los factores que se han estudiado son:
La edad
El sexo
La frecuencia de uso
Estos factores explican una parte de las diferencias entre adolescentes, pero no todas.
¿Qué dicen los estudios recientes?
Los estudios recientes apuntan a que el uso pasivo se asocia con una autoestima más baja.
Según Gómez y Ruiz (2018)
Esto coincide con lo observado en otros países europeos.
Conclusiones
El uso de las redes sociales no es bueno ni malo en sí mismo: depende de cómo se usen.
Referencias
Gómez, A. y Ruiz, B. (2018). Adolescentes en red. Editorial Ejemplo.
Pérez, L. (2020). Uso de redes sociales en España. Revista de Psicología, 12(3), 45-60.
"""


def test_en_el_texto_pegado_se_interpretan_los_titulos(tmp_path):
    clasificados = estructura.clasificar(crear(SIN_FORMATO, tmp_path))
    titulos = {c.texto: (c.nivel, c.confianza) for c in clasificados if c.tipo == "titulo"}
    assert titulos == {
        "Introducción": (1, "media"),
        "Marco teórico": (1, "media"),
        "Antecedentes": (2, "media"),   # justo debajo de otro título
        "¿Qué dicen los estudios recientes?": (1, "media"),
        "Conclusiones": (1, "media"),
        "La frecuencia de uso": (1, "baja"),   # final de una lista sin viñetas: solo se avisa
        "Según Gómez y Ruiz (2018)": (1, "baja"),   # lleva una cita: no es un título
    }


def test_los_titulos_del_texto_pegado_reciben_su_estilo(tmp_path):
    salida, _ = pipeline.procesar(desde_texto.crear(SIN_FORMATO, tmp_path / desde_texto.NOMBRE), tmp_path / "salida")
    estilos = {p.text: p.style.name for p in docx.Document(str(salida)).paragraphs}
    assert estilos["Introducción"] == estilos["Marco Teórico"] == estilos["Conclusiones"] == "Heading 1"
    assert estilos["Antecedentes"] == "Heading 2"
    assert estilos["La edad"] != "Heading 1" and estilos["La frecuencia de uso"] != "Heading 1"


def test_con_formato_no_se_interpreta_por_la_forma(tmp_path):
    # Si el documento tiene algún formato (una negrita), la forma sola no basta: como antes, confianza baja.
    doc = crear(SIN_FORMATO, tmp_path)
    doc.paragraphs[2].runs[0].bold = True
    doc.save(str(tmp_path / "con_negrita.docx"))
    clasificados = estructura.clasificar(docx.Document(str(tmp_path / "con_negrita.docx")))
    assert {c.texto: c.confianza for c in clasificados if c.tipo == "titulo"}["Conclusiones"] == "baja"
