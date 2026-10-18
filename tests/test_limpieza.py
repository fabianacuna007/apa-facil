"""
Paso 3.5: limpieza de formato.

- Párrafos vacíos de relleno: se borran; los que hacen algo (saltos, imágenes,
  marcadores, junto a una tabla o imagen, en la portada) se conservan.
- Tabuladores y espacios al principio de los párrafos de texto: se quitan.
- Color, resaltado, fondo y subrayado del texto: se quitan, salvo en enlaces,
  tablas, portada y código.
- Integridad: los cambios de contenido se anotan como previstos y la
  comparación acepta exactamente esos y ninguno más.

Los casos se prueban con documentos creados aquí; después, los 5 de pruebas/.
"""

import base64
import io
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_BREAK, WD_COLOR_INDEX
from docx.oxml import parse_xml
from docx.shared import Cm, RGBColor

from apa_facil import documento, estructura, pipeline
from apa_facil.documento import CambioPrevisto, ErrorDocumento, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import limpieza, sangria

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
TEXTO = ("Este es un párrafo normal de texto con bastantes palabras para que no parezca un título "
         "ni una línea de portada, y que termina en punto como cualquier frase.")
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def documento_con_portada():
    doc = docx.Document()
    doc.add_paragraph("Portada").runs[0].add_break(WD_BREAK.PAGE)
    return doc


def limpiar(doc):
    informe = Informe()
    antes = radiografia(doc)
    limpieza.aplicar(doc, {}, informe)
    # Lo que haga la limpieza tiene que pasar la comprobación de integridad.
    assert comparar(antes, radiografia(doc), informe.cambios_previstos) == []
    return informe


def sigue(doc, p):
    return p._p.getparent() is not None


# --- 1. Párrafos vacíos -----------------------------------------------------------------

def test_borra_los_vacios_de_relleno():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    vacio = doc.add_paragraph()
    con_espacios = doc.add_paragraph("   ")
    doc.add_paragraph(TEXTO)
    informe = limpiar(doc)
    assert not sigue(doc, vacio) and not sigue(doc, con_espacios)
    assert informe.cambios == ["Limpieza: se han eliminado 2 párrafos vacíos que se usaban para dejar espacio. "
                               "En APA no se dejan líneas en blanco de más: el espacio lo da el interlineado doble."]
    assert [(c.parrafo, c.accion) for c in informe.cambios_previstos] == [(3, "borrar"), (4, "borrar")]


def test_conserva_el_vacio_con_salto_de_pagina():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    salto = doc.add_paragraph()
    salto.add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph(TEXTO)
    assert limpiar(doc).cambios == [] and sigue(doc, salto)


def test_conserva_el_vacio_con_salto_de_seccion():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    seccion = doc.add_paragraph()
    seccion._p.get_or_add_pPr().append(parse_xml(f'<w:sectPr xmlns:w="{W}"/>'))
    doc.add_paragraph(TEXTO)
    limpiar(doc)
    assert sigue(doc, seccion)


def test_conserva_el_vacio_con_salto_antes():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    vacio = doc.add_paragraph()
    vacio.paragraph_format.page_break_before = True
    doc.add_paragraph(TEXTO)
    limpiar(doc)
    assert sigue(doc, vacio)


def test_conserva_el_vacio_con_un_marcador():
    """Un marcador puede ser el destino de un enlace o de una referencia cruzada."""
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    marcado = doc.add_paragraph()
    marcado._p.append(parse_xml(f'<w:bookmarkStart xmlns:w="{W}" w:id="1" w:name="destino"/>'))
    marcado._p.append(parse_xml(f'<w:bookmarkEnd xmlns:w="{W}" w:id="1"/>'))
    doc.add_paragraph(TEXTO)
    limpiar(doc)
    assert sigue(doc, marcado)


def test_conserva_el_vacio_junto_a_una_tabla_y_a_una_imagen():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    lejos = doc.add_paragraph()
    antes_de_la_tabla = doc.add_paragraph()
    doc.add_table(rows=1, cols=1)
    despues_de_la_tabla = doc.add_paragraph()
    doc.add_paragraph(TEXTO)
    doc.add_picture(io.BytesIO(PNG), width=Cm(3))
    despues_de_la_imagen = doc.add_paragraph()
    doc.add_paragraph(TEXTO)
    limpiar(doc)
    assert not sigue(doc, lejos), "de dos vacíos seguidos, solo se conserva el que toca la tabla"
    assert sigue(doc, antes_de_la_tabla) and sigue(doc, despues_de_la_tabla) and sigue(doc, despues_de_la_imagen)


def test_no_toca_la_portada_ni_las_tablas():
    doc = docx.Document()
    doc.add_paragraph("Mi trabajo")
    hueco = doc.add_paragraph()
    doc.add_paragraph("Autora: Ana López").runs[0].add_break(WD_BREAK.PAGE)
    celda = doc.add_table(rows=1, cols=1).rows[0].cells[0]
    celda.add_paragraph()
    doc.add_paragraph(TEXTO)
    limpiar(doc)
    assert sigue(doc, hueco)
    assert len(celda.paragraphs) == 2


def test_la_estructura_queda_sin_los_parrafos_borrados():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    doc.add_paragraph()
    doc.add_paragraph(TEXTO)
    clasificados = estructura.clasificar(doc)
    limpieza.aplicar(doc, {}, Informe(), clasificados)
    assert [c.tipo for c in clasificados] == ["portada", "cuerpo", "cuerpo"]
    assert all(c.elemento.getparent() is not None for c in clasificados)


# --- 2. Tabuladores y espacios al principio ---------------------------------------------

@pytest.mark.parametrize("inicio", ["\t", "\t\t", "    ", "\t  "])
def test_quita_la_sangria_hecha_a_mano(inicio):
    doc = documento_con_portada()
    p = doc.add_paragraph(inicio + TEXTO)
    informe = limpiar(doc)
    assert texto_parrafo(p._p) == TEXTO
    assert [(c.accion, c.quitado) for c in informe.cambios_previstos] == [("recortar", inicio)]
    assert informe.cambios[0].startswith("Limpieza: se han quitado los tabuladores o espacios del principio de "
                                         "1 párrafo («Este es un párrafo normal de texto…»)")


def test_quita_el_tabulador_y_los_espacios_repartidos_en_varios_fragmentos():
    doc = documento_con_portada()
    p = doc.add_paragraph()
    p.add_run().add_tab()
    p.add_run("  ")
    p.add_run(" " + TEXTO)
    informe = limpiar(doc)
    assert texto_parrafo(p._p) == TEXTO
    assert informe.cambios_previstos[0].quitado == "\t   "


def test_no_quita_un_solo_espacio_ni_toca_titulos_ni_listas():
    doc = documento_con_portada()
    un_espacio = doc.add_paragraph(" " + TEXTO)
    lista = doc.add_paragraph("\t• Comparación social")
    titulo = doc.add_heading("\tMétodo", level=1)
    doc.add_paragraph(TEXTO)
    informe = limpiar(doc)
    assert informe.cambios_previstos == []
    assert texto_parrafo(un_espacio._p).startswith(" ")
    assert texto_parrafo(lista._p).startswith("\t") and texto_parrafo(titulo._p).startswith("\t")


def test_el_parrafo_limpio_recibe_la_sangria():
    doc = documento_con_portada()
    p = doc.add_paragraph("\t" + TEXTO)
    clasificados = estructura.clasificar(doc)
    informe = Informe()
    limpieza.aplicar(doc, {}, informe, clasificados)
    sangria.aplicar(doc, {}, informe, clasificados)
    assert p.paragraph_format.first_line_indent == Cm(1.27)
    assert not any("tabulador" in a for a in informe.avisos)


# --- 3. Color, resaltado y subrayado ----------------------------------------------------

def parrafo_con(doc, texto, color=None, resaltado=None, subrayado=False, estilo=None):
    p = doc.add_paragraph()
    r = p.add_run(texto, style=estilo)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    if resaltado:
        r.font.highlight_color = resaltado
    r.font.underline = subrayado or None
    return r


def marcas(r):
    rpr = r._r.find(w("rPr"))
    return {e for e in ("color", "highlight", "shd", "u") if rpr is not None and rpr.find(w(e)) is not None}


def test_quita_color_resaltado_y_subrayado():
    doc = documento_con_portada()
    rojo = parrafo_con(doc, TEXTO, color="C00000")
    amarillo = parrafo_con(doc, TEXTO, resaltado=WD_COLOR_INDEX.YELLOW)
    subrayado = parrafo_con(doc, "Introducción", subrayado=True)
    doc.add_paragraph(TEXTO)
    informe = limpiar(doc)
    assert marcas(rojo) == marcas(amarillo) == marcas(subrayado) == set()
    assert informe.cambios == [
        "Limpieza: se ha quitado el color de 1 fragmento de texto («Este es un párrafo normal de texto…»). "
        "APA usa texto negro; los enlaces conservan el suyo.",
        "Limpieza: se ha quitado el resaltado o fondo de color de 1 fragmento de texto («Este es un párrafo normal "
        "de texto…»). APA no usa resaltados ni fondos de color; los enlaces conservan el suyo.",
        "Limpieza: se ha quitado el subrayado de 1 fragmento de texto («Introducción»). APA no usa subrayado; los "
        "enlaces conservan el suyo."]


def test_quita_el_fondo_de_color_del_texto():
    """Caso del documento 04: un fondo amarillo puesto como sombreado del texto."""
    doc = documento_con_portada()
    r = parrafo_con(doc, TEXTO)
    r._r.get_or_add_rPr().append(parse_xml(f'<w:shd xmlns:w="{W}" w:val="clear" w:color="auto" w:fill="FFFF00"/>'))
    assert "resaltado o fondo de color" in limpiar(doc).cambios[0]
    assert marcas(r) == set()


def test_conserva_los_enlaces():
    doc = documento_con_portada()
    de_estilo = parrafo_con(doc, "pulsa aquí", color="0563C1", subrayado=True)
    doc.styles.add_style("Hyperlink", docx.enum.style.WD_STYLE_TYPE.CHARACTER)
    de_estilo.style = doc.styles["Hyperlink"]
    a_mano = parrafo_con(doc, "www.example.org/datos-uso-redes-2024", color="0563C1", subrayado=True)
    hipervinculo = doc.add_paragraph()
    hipervinculo._p.append(parse_xml(
        f'<w:hyperlink xmlns:w="{W}"><w:r><w:rPr><w:color w:val="0563C1"/><w:u w:val="single"/></w:rPr>'
        '<w:t>Ministerio</w:t></w:r></w:hyperlink>'))
    doc.add_paragraph(TEXTO)
    assert limpiar(doc).cambios == []
    assert marcas(de_estilo) == marcas(a_mano) == {"color", "u"}
    assert len(list(hipervinculo._p.iter(w("color")))) == 1


def test_no_toca_portada_tablas_ni_codigo():
    doc = docx.Document()
    portada = parrafo_con(doc, "MI TRABAJO", color="C00000", subrayado=True)
    portada._r.append(parse_xml(f'<w:br xmlns:w="{W}" w:type="page"/>'))
    celda = doc.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0].add_run("Variable")
    celda.font.color.rgb = RGBColor.from_string("FFFFFF")
    codigo = parrafo_con(doc, "for i in range(10): print(i)", color="0000FF")
    codigo.font.name = "Courier New"
    doc.add_paragraph(TEXTO)
    assert limpiar(doc).cambios == []
    assert marcas(portada) == {"color", "u"} and marcas(celda) == {"color"} and marcas(codigo) == {"color"}


def test_no_cambia_los_estilos():
    doc = documento_con_portada()
    doc.styles["Heading 1"].font.color.rgb = RGBColor.from_string("2E74B5")
    doc.add_heading("Método", level=1)
    doc.add_paragraph(TEXTO)
    limpiar(doc)
    assert str(doc.styles["Heading 1"].font.color.rgb) == "2E74B5"


# --- 4. Integridad: cambios previstos ---------------------------------------------------

def doc_y_radiografia():
    doc = documento_con_portada()
    doc.add_paragraph("\t" + TEXTO)
    doc.add_paragraph()
    doc.add_paragraph(TEXTO)
    return doc, radiografia(doc)


def test_comparar_acepta_exactamente_lo_previsto():
    doc, antes = doc_y_radiografia()
    informe = Informe()
    limpieza.aplicar(doc, {}, informe)
    despues = radiografia(doc)
    assert comparar(antes, despues) != [], "sin los previstos, se ve la diferencia"
    assert comparar(antes, despues, informe.cambios_previstos) == []


def test_comparar_no_deja_borrar_un_parrafo_con_texto():
    doc, antes = doc_y_radiografia()
    p = doc.paragraphs[3]._p
    p.getparent().remove(p)
    diferencias = comparar(antes, radiografia(doc), [CambioPrevisto(4, "borrar")])
    assert any("Párrafo 4: una regla anotó «borrar»" in d for d in diferencias)


def test_comparar_no_deja_recortar_algo_que_no_es_espacio():
    doc, antes = doc_y_radiografia()
    diferencias = comparar(antes, radiografia(doc), [CambioPrevisto(2, "recortar", "\tEste")])
    assert any("Párrafo 2: una regla anotó «recortar»" in d for d in diferencias)


def test_comparar_ve_un_cambio_no_anotado():
    doc, antes = doc_y_radiografia()
    informe = Informe()
    limpieza.aplicar(doc, {}, informe)
    doc.paragraphs[-1].runs[0].text = "Otro texto."
    assert any("el texto ha cambiado" in d for d in comparar(antes, radiografia(doc), informe.cambios_previstos))


def test_si_la_limpieza_no_anota_un_cambio_no_se_entrega_la_copia(tmp_path, monkeypatch):
    """A la contra: sin anotar los cambios previstos, el pipeline no entrega la copia del 04."""
    monkeypatch.setattr(Informe, "cambio_previsto", lambda self, *args, **kwargs: None)
    original = next(p for p in PRUEBAS if p.name.startswith("04"))
    with pytest.raises(ErrorDocumento, match="no coincide con el original"):
        pipeline.procesar(original, tmp_path)
    assert not (tmp_path / "04_formato_caotico_APA.docx").exists()


# --- Los 5 documentos de prueba ---------------------------------------------------------

ESPERADO = {   # (vacíos borrados, párrafos con tabulador limpiados, mensajes de limpieza)
    "01": (0, 0, ["subrayado de 1 fragmento"]),
    "02": (0, 0, ["color de 2 fragmentos"]),
    "03": (0, 0, []),
    "04": (6, 3, ["eliminado 6 párrafos vacíos", "principio de 3 párrafos", "color de 2 fragmentos",
                  "fondo de color de 1 fragmento"]),
    "05": (0, 0, []),
}


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_documentos_de_prueba(original, tmp_path):
    salida, informe = pipeline.procesar(original, tmp_path)
    borrados, recortados, mensajes = ESPERADO[original.name[:2]]
    acciones = [c.accion for c in informe.cambios_previstos]
    assert (acciones.count("borrar"), acciones.count("recortar")) == (borrados, recortados)
    de_limpieza = [m for m in informe.cambios if m.startswith("Limpieza")]
    assert len(de_limpieza) == len(mensajes)
    for mensaje, esperado in zip(de_limpieza, mensajes):
        assert esperado in mensaje
    # La copia ya no tiene vacíos de relleno ni tabuladores al principio de los párrafos de texto.
    copia = docx.Document(str(salida))
    for c in estructura.clasificar(copia):
        if c.tipo == "cuerpo":
            assert not texto_parrafo(c.elemento).startswith(("\t", "  ")), c.texto
