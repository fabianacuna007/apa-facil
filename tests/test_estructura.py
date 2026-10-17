"""
Paso 3.1: clasificador de párrafos.

Primero se prueba cada tipo con documentos creados aquí mismo; después se
compara la clasificación de los 5 documentos de pruebas/ con la escrita a
mano en clasificacion_esperada.py.
"""

import base64
import io
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.shared import Cm, Pt
from lxml import etree

from apa_facil import documento, estructura
from apa_facil.estructura import ALTA, BAJA, MEDIA, TIPOS, ZONAS, clasificar

from clasificacion_esperada import ESPERADO

PRUEBAS = sorted((Path(__file__).resolve().parent.parent / "pruebas").glob("*.docx"))
PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")
TEXTO = ("Este es un párrafo normal de texto con bastantes palabras para que no parezca un título "
         "ni una línea de portada, y que termina en punto como cualquier frase.")
CITA = " ".join(["Una cita larga copiada de otra obra con muchas palabras seguidas."] * 5)
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

# Diferencias admitidas porque las resuelve un paso posterior: (documento, párrafo) → qué falta.
# Ahora no hay ninguna (el nivel de «Factores de riesgo:» del 04 se resolvió en el paso 3.2).
PENDIENTES = {}


def documento_con_portada():
    """Un documento cuya primera página es una portada (para que no estorbe en las pruebas)."""
    doc = docx.Document()
    doc.add_paragraph("Portada").runs[0].add_break(WD_BREAK.PAGE)
    return doc


def de(doc, p):
    """La clasificación de un párrafo creado con python-docx."""
    return next(c for c in clasificar(doc) if c.elemento is p._p)


def negrita(doc, texto):
    p = doc.add_paragraph()
    p.add_run(texto).bold = True
    return p


def imagen(doc):
    doc.add_picture(io.BytesIO(PNG), width=Cm(5))
    return doc.paragraphs[-1]


# --- Todos los párrafos tienen una clasificación completa ------------------------------

@pytest.mark.parametrize("ruta", PRUEBAS, ids=lambda r: r.name)
def test_cada_parrafo_tiene_tipo_zona_confianza_y_motivo(ruta):
    for c in clasificar(documento.abrir(ruta)):
        assert c.tipo in TIPOS and c.zona in ZONAS and c.confianza in (ALTA, MEDIA, BAJA)
        assert c.motivo
        assert (c.nivel is not None) == (c.tipo == "titulo")
        assert c.nivel is None or 1 <= c.nivel <= 5


def test_clasificar_no_cambia_el_documento():
    doc = documento.abrir(PRUEBAS[-1])
    antes = etree.tostring(doc.element)
    clasificar(doc)
    assert etree.tostring(doc.element) == antes


def test_no_clasifica_los_parrafos_de_las_tablas():
    doc = documento_con_portada()
    doc.add_table(rows=1, cols=1).rows[0].cells[0].paragraphs[0].add_run("Celda")
    assert "Celda" not in [c.texto for c in clasificar(doc)]


# --- Vacíos e imágenes ------------------------------------------------------------------

def test_parrafo_vacio():
    doc = documento_con_portada()
    p = doc.add_paragraph("   ")
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("vacio", ALTA)


def test_vacio_con_salto_lo_dice_en_el_motivo():
    doc = documento_con_portada()
    doc.add_paragraph(TEXTO)
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)
    assert de(doc, p).tipo == "vacio"
    assert "salto" in de(doc, p).motivo


def test_parrafo_con_solo_una_imagen_no_es_vacio():
    doc = documento_con_portada()
    p = imagen(doc)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("imagen", ALTA)


# --- Portada ----------------------------------------------------------------------------

def test_portada_hasta_el_salto_de_pagina():
    doc = docx.Document()
    titulo = doc.add_paragraph("Mi trabajo")
    autor = doc.add_paragraph("Autora: Ana López")
    autor.runs[0].add_break(WD_BREAK.PAGE)
    texto = doc.add_paragraph(TEXTO)
    assert [(de(doc, p).tipo, de(doc, p).zona) for p in (titulo, autor)] == [("portada", "portada")] * 2
    assert de(doc, titulo).confianza == ALTA
    assert (de(doc, texto).tipo, de(doc, texto).zona) == ("cuerpo", "cuerpo")


def test_portada_sin_salto_acaba_antes_del_primer_titulo():
    """Caso del documento 01: «Introducción» no es parte de la portada."""
    doc = docx.Document()
    portada = doc.add_paragraph("Curso 2025-2026")
    titulo = negrita(doc, "Introducción")
    doc.add_paragraph(TEXTO + " " + TEXTO)
    assert (de(doc, portada).tipo, de(doc, portada).confianza) == ("portada", MEDIA)
    assert (de(doc, titulo).tipo, de(doc, titulo).zona) == ("titulo", "cuerpo")


def test_titulo_mas_grande_que_el_texto_no_es_portada():
    """Caso del documento 04: un título sin negrita, pero con letra más grande que el texto."""
    doc = docx.Document()
    doc.add_paragraph("Nombre: Irene")
    titulo = doc.add_paragraph()
    titulo.add_run("¿Qué son las redes sociales?").font.size = Pt(14)
    doc.add_paragraph().add_run(TEXTO + " " + TEXTO).font.size = Pt(11)
    assert de(doc, titulo).tipo == "titulo"


def test_portada_acaba_en_un_titulo_de_word():
    """Caso del documento 02: el «Resumen» con estilo Título 1 no es parte de la portada."""
    doc = docx.Document()
    portada = doc.add_paragraph("Hábitos de sueño", style="Title")
    resumen = doc.add_heading("Resumen", level=1)
    assert de(doc, portada).tipo == "portada"
    assert (de(doc, resumen).tipo, de(doc, resumen).zona) == ("titulo", "resumen")


def test_imagen_y_vacio_en_la_portada():
    doc = docx.Document()
    logo = imagen(doc)
    hueco = doc.add_paragraph()
    doc.add_paragraph("Universidad").runs[0].add_break(WD_BREAK.PAGE)
    assert (de(doc, logo).tipo, de(doc, logo).zona) == ("imagen", "portada")
    assert (de(doc, hueco).tipo, de(doc, hueco).zona) == ("vacio", "portada")


def test_sin_portada():
    doc = docx.Document()
    p = doc.add_paragraph(TEXTO + " " + TEXTO)
    assert (de(doc, p).tipo, de(doc, p).zona) == ("cuerpo", "cuerpo")


# --- Títulos ----------------------------------------------------------------------------

@pytest.mark.parametrize("nivel", [1, 2, 3, 4, 5])
def test_titulo_de_word_con_su_nivel(nivel):
    doc = documento_con_portada()
    p = doc.add_heading("Método", level=nivel)
    c = de(doc, p)
    assert (c.tipo, c.nivel, c.confianza) == ("titulo", nivel, ALTA)
    assert "heading" in c.motivo.lower()


def test_titulo_de_nivel_6_de_word_se_queda_en_5():
    doc = documento_con_portada()
    p = doc.add_heading("Detalle", level=6)
    assert de(doc, p).nivel == 5


def test_titulo_falso_en_negrita_con_numeracion():
    doc = documento_con_portada()
    p = negrita(doc, "1.2. Cambios en la organización del trabajo")
    c = de(doc, p)
    assert (c.tipo, c.nivel, c.confianza) == ("titulo", 2, MEDIA)


def test_titulo_falso_en_mayusculas():
    doc = documento_con_portada()
    p = doc.add_paragraph("CONCLUSIONES.")
    assert de(doc, p).tipo == "titulo"


def test_frase_corta_con_enlace_no_es_titulo():
    """Caso del documento 04: «Más información en: www…»."""
    doc = documento_con_portada()
    p = doc.add_paragraph("Más información en: www.example.org/redes")
    assert de(doc, p).tipo == "cuerpo"


# --- Títulos falsos: detección por indicios (paso 3.2) ----------------------------------

def parrafo(doc, texto, tamano=11, negrita=False, cursiva=False, centrado=False):
    p = doc.add_paragraph()
    fragmento = p.add_run(texto)
    fragmento.font.size, fragmento.bold, fragmento.italic = Pt(tamano), negrita, cursiva
    if centrado:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def titulo_y_texto(doc, texto, **formato):
    """Un título falso seguido de un párrafo de texto normal (11 pt)."""
    p = parrafo(doc, texto, **formato)
    parrafo(doc, TEXTO)
    return p


def test_titulo_solo_con_letra_mas_grande():
    """Caso del documento 04: sin negrita, pero con letra más grande que el texto."""
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "Efectos sobre la autoestima", tamano=13)
    c = de(doc, p)
    assert (c.tipo, c.confianza) == ("titulo", ALTA)
    assert "más grande que el texto (13 pt frente a 11 pt)" in c.motivo


def test_titulo_numerado_sin_negrita():
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "1. Introducción")
    c = de(doc, p)
    assert (c.tipo, c.nivel, c.confianza) == ("titulo", 1, ALTA)
    assert "numerado como un apartado" in c.motivo


def test_apartado_seguido_de_su_subapartado():
    """«1. Marco teórico» seguido directamente de «1.1. Antecedentes»: los dos son títulos."""
    doc = documento_con_portada()
    apartado = parrafo(doc, "1. Marco teórico")
    subapartado = titulo_y_texto(doc, "1.1. Antecedentes")
    assert [(de(doc, p).tipo, de(doc, p).nivel) for p in (apartado, subapartado)] == [("titulo", 1), ("titulo", 2)]


@pytest.mark.parametrize("lineas", [["1. Primera idea", "2. Segunda idea", "3. Tercera idea"],
                                    ["1. ¿Desde cuándo funciona su negocio?", "2. ¿Cómo ha evolucionado?"],
                                    ["a) Primera opción", "b) Segunda opción"]])
def test_lineas_numeradas_seguidas_son_una_lista(lineas):
    """Caso del anexo del documento 05: aunque la última vaya seguida de texto, es una lista."""
    doc = documento_con_portada()
    parrafos = [parrafo(doc, linea) for linea in lineas]
    parrafo(doc, TEXTO)
    assert [de(doc, p).tipo for p in parrafos] == ["lista"] * len(lineas)


def test_vineta_en_negrita_no_es_titulo():
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "• Idea clave", negrita=True)
    assert de(doc, p).tipo == "lista"


def test_linea_sin_formato_seguida_de_texto_es_titulo_dudoso():
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "Introducción")
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("titulo", BAJA)


def test_linea_sin_formato_entre_lineas_cortas_no_es_titulo():
    doc = documento_con_portada()
    p = parrafo(doc, "Gracias por su atención")
    parrafo(doc, "Madrid, junio de 2026")
    assert de(doc, p).tipo == "cuerpo"


def test_frase_corta_en_negrita_con_punto_es_titulo_dudoso():
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "Esto es importante.", negrita=True)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("titulo", BAJA)
    assert "aunque acaba en punto" in de(doc, p).motivo


def test_frase_corta_sin_formato_con_punto_no_es_titulo():
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "Esto es importante.")
    assert de(doc, p).tipo == "cuerpo"


def test_el_tamano_del_texto_no_cuenta_la_portada():
    """La portada suele tener letra grande: no debe cambiar el tamaño del texto normal."""
    doc = docx.Document()
    portada = parrafo(doc, "Título del trabajo en grande", tamano=20)
    portada.runs[0].add_break(WD_BREAK.PAGE)
    p = titulo_y_texto(doc, "Método", tamano=14)
    assert "14 pt frente a 11 pt" in de(doc, p).motivo


@pytest.mark.parametrize("texto, esperado", [
    ("1. Introducción", ("arábiga", "1", 1)),
    ("2.3. Resultados", ("arábiga", "2.3", 2)),
    ("1.1.1 Objetivo general", ("arábiga", "1.1.1", 3)),
    ("3) Propuestas", ("arábiga", "3", 1)),
    ("II. Marco teórico", ("romana", "II", 1)),
    ("a) Evolución demográfica", ("letra", "a", None)),
    ("B. Segundo grupo", ("letra", "B", None)),
    ("2026 fue un año difícil", None),
    ("Introducción", None),
])
def test_numeracion_de_apartado(texto, esperado):
    assert estructura._numeracion(texto) == esperado


# --- Títulos falsos: niveles (paso 3.2) -------------------------------------------------

def niveles(doc, parrafos):
    return [de(doc, p).nivel for p in parrafos]


def test_nivel_por_tamano():
    """Caso del documento 04: los de 13 pt son nivel 1; el de 11 pt en negrita, nivel 2."""
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "Efectos sobre la autoestima", tamano=13),
               titulo_y_texto(doc, "Factores de riesgo", negrita=True),
               titulo_y_texto(doc, "Conclusión", tamano=13)]
    assert niveles(doc, titulos) == [1, 2, 1]
    assert "menos destacado que el de los títulos de nivel 1" in de(doc, titulos[1]).motivo


def test_tres_niveles_por_formato():
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "Marco teórico", tamano=16, negrita=True),
               titulo_y_texto(doc, "Antecedentes", tamano=13, negrita=True),
               titulo_y_texto(doc, "Estudios previos", negrita=True, cursiva=True)]
    assert niveles(doc, titulos) == [1, 2, 3]


def test_mismo_tamano_centrado_negrita_y_cursiva():
    """Como en APA: centrado y negrita (1), a la izquierda en negrita (2), negrita y cursiva (3)."""
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "Método", negrita=True, centrado=True),
               titulo_y_texto(doc, "Participantes", negrita=True),
               titulo_y_texto(doc, "Criterios de inclusión", negrita=True, cursiva=True)]
    assert niveles(doc, titulos) == [1, 2, 3]


def test_un_solo_formato_es_nivel_1():
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "Introducción", negrita=True), titulo_y_texto(doc, "Conclusión", negrita=True)]
    assert niveles(doc, titulos) == [1, 1]
    assert "único formato de título" in de(doc, titulos[0]).motivo


def test_mismo_formato_que_un_titulo_numerado():
    """Caso del documento 01: «Introducción» se ve igual que «1. Cambios…», que es nivel 1."""
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "Introducción", tamano=12, negrita=True),
               titulo_y_texto(doc, "1. Cambios en el trabajo", tamano=12, negrita=True),
               titulo_y_texto(doc, "1.1. Los talleres", negrita=True),
               titulo_y_texto(doc, "Otro detalle", negrita=True)]
    assert niveles(doc, titulos) == [1, 1, 2, 2]
    assert "mismo formato que los demás títulos de nivel 1" in de(doc, titulos[0]).motivo
    assert "mismo formato que los demás títulos de nivel 2" in de(doc, titulos[3]).motivo


def test_mismo_formato_que_un_titulo_de_word():
    """En la plantilla, «Título 2» de Word es de 13 pt en negrita: un título falso igual es nivel 2."""
    doc = documento_con_portada()
    doc.add_heading("Método", level=1)
    doc.add_heading("Participantes", level=2)
    parrafo(doc, TEXTO)
    falso = titulo_y_texto(doc, "Instrumentos", tamano=13, negrita=True)
    assert de(doc, falso).nivel == 2


def test_formato_menos_destacado_que_los_de_word():
    doc = documento_con_portada()
    doc.add_heading("Método", level=1)          # 14 pt en negrita
    parrafo(doc, TEXTO)
    falso = titulo_y_texto(doc, "Participantes", tamano=12, negrita=True)
    assert de(doc, falso).nivel == 2


def test_letra_va_por_debajo_del_titulo_anterior():
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "1. Resultados", negrita=True),
               titulo_y_texto(doc, "a) Primer grupo", negrita=True),
               titulo_y_texto(doc, "b) Segundo grupo", negrita=True)]
    assert niveles(doc, titulos) == [1, 2, 2]


def test_numeros_romanos_por_encima_de_los_arabigos():
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "I. Marco teórico", negrita=True),
               titulo_y_texto(doc, "1. Antecedentes", negrita=True)]
    assert niveles(doc, titulos) == [1, 2]


def test_no_se_salta_ningun_nivel():
    doc = documento_con_portada()
    titulos = [titulo_y_texto(doc, "1. Introducción", negrita=True),
               titulo_y_texto(doc, "1.1.1. Detalle", negrita=True)]
    assert niveles(doc, titulos) == [1, 2]
    c = de(doc, titulos[1])
    assert c.confianza == BAJA
    assert "no se puede saltar un nivel" in c.motivo


def test_el_primer_titulo_puede_no_ser_de_nivel_1():
    """Sin título anterior no hay salto que corregir."""
    doc = documento_con_portada()
    p = titulo_y_texto(doc, "1.2. Cambios en la organización", negrita=True)
    assert de(doc, p).nivel == 2


def test_los_titulos_de_word_no_cambian_de_nivel():
    doc = documento_con_portada()
    doc.add_heading("Método", level=1)
    parrafo(doc, TEXTO)
    p = doc.add_heading("Detalle", level=3)
    assert de(doc, p).nivel == 3


# --- Listas -----------------------------------------------------------------------------

def test_lista_de_word():
    doc = documento_con_portada()
    p = doc.add_paragraph("Primer elemento", style="List Bullet")
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("lista", ALTA)


@pytest.mark.parametrize("texto", ["• Comparación social", "- Ciberacoso", "1. ¿Desde cuándo funciona?",
                                   "a) Primera opción"])
def test_lista_escrita_a_mano(texto):
    doc = documento_con_portada()
    p = doc.add_paragraph(texto)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("lista", MEDIA)


# --- Leyendas ---------------------------------------------------------------------------

def test_leyenda_de_tabla_debajo_de_la_tabla():
    """Caso del documento 02: «Tabla 1: …» va debajo de la tabla."""
    doc = documento_con_portada()
    doc.add_table(rows=1, cols=1)
    p = doc.add_paragraph("Tabla 1: Estadísticos descriptivos")
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("leyenda_tabla", ALTA)


def test_leyenda_de_tabla_sin_tabla_al_lado():
    doc = documento_con_portada()
    p = doc.add_paragraph("Tabla 3. Datos de uso")
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("leyenda_tabla", MEDIA)


def test_titulo_sin_numero_encima_de_una_tabla():
    """Caso del documento 04: «Datos de uso por red social» encima de la tabla."""
    doc = documento_con_portada()
    p = doc.add_paragraph("Datos de uso por red social")
    doc.add_paragraph()
    doc.add_table(rows=1, cols=1)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("leyenda_tabla", MEDIA)


@pytest.mark.parametrize("texto", ["Figura 1. Consumo", "Gráfico 1 – Consumo eléctrico", "Imagen 2: fuentes"])
def test_leyenda_de_figura(texto):
    doc = documento_con_portada()
    imagen(doc)
    p = doc.add_paragraph(texto)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("leyenda_figura", ALTA)


@pytest.mark.parametrize("texto", ["Nota. Elaboración propia.", "Fuente: padrón municipal."])
def test_nota_de_tabla_y_de_figura(texto):
    doc = documento_con_portada()
    doc.add_table(rows=1, cols=1)
    de_tabla = doc.add_paragraph(texto)
    doc.add_paragraph(TEXTO)
    imagen(doc)
    de_figura = doc.add_paragraph(texto)
    assert de(doc, de_tabla).tipo == "leyenda_tabla"
    assert de(doc, de_figura).tipo == "leyenda_figura"


def test_estilo_de_leyenda_de_word():
    doc = documento_con_portada()
    imagen(doc)
    p = doc.add_paragraph("Consumo eléctrico medio", style="Caption")
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("leyenda_figura", ALTA)


def test_frase_que_habla_de_una_tabla_es_cuerpo():
    doc = documento_con_portada()
    p = doc.add_paragraph("La tabla 2 muestra el rendimiento según las horas de sueño.")
    assert de(doc, p).tipo == "cuerpo"


# --- Citas en bloque --------------------------------------------------------------------

def test_cita_en_bloque():
    doc = documento_con_portada()
    p = doc.add_paragraph(CITA)
    p.paragraph_format.left_indent = Cm(1.27)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("cita_bloque", ALTA)


def test_parrafo_corto_con_sangria_izquierda_es_cita_dudosa():
    doc = documento_con_portada()
    p = doc.add_paragraph("Una cita corta con sangría, que no llega a cuarenta palabras.")
    p.paragraph_format.left_indent = Cm(1.27)
    assert (de(doc, p).tipo, de(doc, p).confianza) == ("cita_bloque", MEDIA)


# --- Zonas: resumen, índice, referencias y anexos ---------------------------------------

def test_zona_de_resumen_hasta_el_siguiente_titulo():
    doc = documento_con_portada()
    rotulo = negrita(doc, "RESUMEN")
    resumen = doc.add_paragraph(TEXTO)
    clave = doc.add_paragraph("Palabras clave: sueño, rendimiento.")
    titulo = doc.add_heading("Introducción", level=1)
    texto = doc.add_paragraph(TEXTO)
    assert [(de(doc, p).tipo, de(doc, p).zona) for p in (rotulo, resumen, clave, titulo, texto)] == [
        ("titulo", "resumen"), ("cuerpo", "resumen"), ("cuerpo", "resumen"), ("titulo", "cuerpo"),
        ("cuerpo", "cuerpo")]
    assert de(doc, rotulo).nivel == 1


def test_zona_de_resumen_acaba_en_el_salto_de_pagina():
    doc = documento_con_portada()
    negrita(doc, "Resumen")
    doc.add_paragraph(TEXTO).runs[0].add_break(WD_BREAK.PAGE)
    texto = doc.add_paragraph(TEXTO)
    assert de(doc, texto).zona == "cuerpo"


def test_indice_automatico_de_word():
    """Caso del documento 05: un campo TOC que ocupa dos párrafos, dentro de un control de contenido."""
    doc = documento_con_portada()
    doc.element.body.insert(len(doc.element.body) - 1, parse_xml(
        f'<w:sdt xmlns:w="{W}"><w:sdtPr><w:alias w:val="Índice"/></w:sdtPr><w:sdtContent>'
        '<w:p><w:r><w:fldChar w:fldCharType="begin"/><w:instrText xml:space="preserve">TOC \\o "1-3"'
        '</w:instrText><w:fldChar w:fldCharType="separate"/></w:r></w:p>'
        '<w:p><w:r><w:fldChar w:fldCharType="end"/></w:r></w:p></w:sdtContent></w:sdt>'))
    fin = doc.add_paragraph()
    fin.add_run().add_break(WD_BREAK.PAGE)
    texto = doc.add_paragraph(TEXTO)
    clasificados = clasificar(doc)
    assert [(c.tipo, c.zona) for c in clasificados[1:]] == [
        ("indice", "indice"), ("indice", "indice"), ("vacio", "indice"), ("cuerpo", "cuerpo")]
    assert clasificados[-1].elemento is texto._p


def test_campo_toc_sin_control_de_contenido():
    doc = documento_con_portada()
    p = doc.add_paragraph()
    p._p.append(parse_xml(f'<w:fldSimple xmlns:w="{W}" w:instr="TOC \\o &quot;1-3&quot;"/>'))
    assert (de(doc, p).tipo, de(doc, p).zona) == ("indice", "indice")


def test_indice_escrito_a_mano_hasta_el_salto():
    doc = documento_con_portada()
    rotulo = negrita(doc, "Índice")
    linea = doc.add_paragraph("1. Introducción\t3")
    ultima = doc.add_paragraph("2. Método\t5")
    ultima.runs[0].add_break(WD_BREAK.PAGE)
    titulo = doc.add_heading("Introducción", level=1)
    assert [(de(doc, p).tipo, de(doc, p).zona) for p in (rotulo, linea, ultima, titulo)] == [
        ("indice", "indice"), ("indice", "indice"), ("indice", "indice"), ("titulo", "cuerpo")]


@pytest.mark.parametrize("rotulo", ["Referencias", "BIBLIOGRAFÍA", "Webgrafía:", "Fuentes consultadas"])
def test_rotulo_de_referencias(rotulo):
    doc = documento_con_portada()
    p = doc.add_paragraph(rotulo)
    assert (de(doc, p).tipo, de(doc, p).zona, de(doc, p).confianza) == ("titulo_referencias", "referencias", ALTA)


def test_referencias_hasta_el_siguiente_titulo():
    doc = documento_con_portada()
    doc.add_heading("Referencias", level=1)
    con_francesa = doc.add_paragraph("Molina, R. (2017). Turismo y territorio. Editorial Ejemplo.")
    con_francesa.paragraph_format.left_indent = Cm(1.27)
    con_francesa.paragraph_format.first_line_indent = Cm(-1.27)
    sin_francesa = doc.add_paragraph("Ortega, L. y Blanco, R. (2018). El tiempo de la fábrica.")
    url = doc.add_paragraph("https://www.historiaejemplo.es/revolucion")
    rara = doc.add_paragraph("Apuntes de clase del profesor")
    titulo = doc.add_heading("Agradecimientos", level=1)
    assert [(de(doc, p).tipo, de(doc, p).confianza) for p in (con_francesa, sin_francesa, url, rara)] == [
        ("referencia", ALTA), ("referencia", ALTA), ("referencia", MEDIA), ("referencia", BAJA)]
    assert (de(doc, titulo).tipo, de(doc, titulo).zona) == ("titulo", "cuerpo")


def test_titulo_falso_acaba_la_lista_de_referencias():
    doc = documento_con_portada()
    doc.add_paragraph("Referencias")
    doc.add_paragraph("Ortega, L. (2018). El tiempo de la fábrica.")
    titulo = negrita(doc, "Agradecimientos")
    assert (de(doc, titulo).tipo, de(doc, titulo).zona) == ("titulo", "cuerpo")


def test_anexos():
    doc = documento_con_portada()
    doc.add_paragraph("Referencias")
    doc.add_paragraph("Ortega, L. (2018). El tiempo de la fábrica.")
    rotulo = doc.add_heading("Anexo I. Guion de las entrevistas", level=1)
    pregunta = doc.add_paragraph("1. ¿Desde cuándo funciona su negocio?")
    tabla = doc.add_paragraph("Tabla A1. Datos de las entrevistas")
    assert [(de(doc, p).tipo, de(doc, p).zona) for p in (rotulo, pregunta, tabla)] == [
        ("anexo", "anexos"), ("lista", "anexos"), ("leyenda_tabla", "anexos")]


def test_frase_que_empieza_por_anexo_no_es_rotulo():
    doc = documento_con_portada()
    p = doc.add_paragraph("Anexo a este trabajo se incluyen las entrevistas completas.")
    assert (de(doc, p).tipo, de(doc, p).zona) == ("cuerpo", "cuerpo")


# --- Los 5 documentos de prueba ---------------------------------------------------------

def test_todos_los_documentos_de_prueba_tienen_clasificacion_esperada():
    assert sorted(ESPERADO) == [r.name for r in PRUEBAS]


@pytest.mark.parametrize("ruta", PRUEBAS, ids=lambda r: r.name)
def test_clasificacion_de_los_documentos_de_prueba(ruta):
    obtenido = [(c.tipo, c.nivel, c.zona) for c in clasificar(documento.abrir(ruta))]
    esperado = ESPERADO[ruta.name]
    assert len(obtenido) == len(esperado)
    fallos = [f"párrafo {i}: se esperaba {e}, sale {o}"
              for i, (e, o) in enumerate(zip(esperado, obtenido))
              if e != o and (ruta.name, i) not in PENDIENTES]
    assert not fallos, "\n".join(fallos)


@pytest.mark.parametrize("documento_y_parrafo", sorted(PENDIENTES))
def test_pendientes_siguen_pendientes(documento_y_parrafo):
    """Si un pendiente se arregla, hay que quitarlo de PENDIENTES (y la prueba de arriba lo vigilará)."""
    nombre, i = documento_y_parrafo
    c = clasificar(documento.abrir(PRUEBAS[0].parent / nombre))[i]
    assert (c.tipo, c.nivel, c.zona) != ESPERADO[nombre][i]
