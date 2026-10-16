"""
Paso 0.5: la prueba de integridad.

Comprueba tres cosas:
1. la ida y vuelta conserva todo el contenido de los 5 documentos;
2. la radiografía ve lo que tiene que ver (si no viera nada, la comparación
   "pasaría" siempre sin servir para nada);
3. la comparación detecta un cambio de verdad, y entonces no se entrega la copia.
"""

from pathlib import Path

import docx
import pytest

from apa_facil import documento, pipeline
from apa_facil.documento import ErrorDocumento, comparar, radiografia, w

CARPETA = Path(__file__).resolve().parent.parent / "pruebas"
PRUEBAS = sorted(CARPETA.glob("*.docx"))


def abrir(nombre):
    return docx.Document(str(CARPETA / nombre))


def texto_completo(r):
    return "\n".join(t for _, t in r.textos)


@pytest.mark.parametrize("original", PRUEBAS, ids=lambda p: p.stem)
def test_la_ida_y_vuelta_conserva_todo(original, tmp_path):
    salida, _ = pipeline.procesar(original, tmp_path)
    antes = radiografia(docx.Document(str(original)))
    despues = radiografia(docx.Document(str(salida)))
    assert comparar(antes, despues) == []


@pytest.mark.parametrize("nombre, atributo, valor", [
    ("02_informe_tablas.docx", "tablas", 2),
    ("03_trabajo_imagenes_listas.docx", "imagenes", 2),
    ("04_formato_caotico.docx", "secciones", 3),
    ("05_tfg_completo.docx", "notas_al_pie", 3),
    ("05_tfg_completo.docx", "secciones", 3),
    ("05_tfg_completo.docx", "tablas", 1),
    ("05_tfg_completo.docx", "imagenes", 1),
])
def test_la_radiografia_cuenta_bien(nombre, atributo, valor):
    assert getattr(radiografia(abrir(nombre)), atributo) == valor


@pytest.mark.parametrize("nombre, fragmento", [
    ("02_informe_tablas.docx", "Horas de sueño (entre semana)"),              # celda de tabla
    ("05_tfg_completo.docx", "Se entiende por municipio rural"),              # nota al pie
    ("05_tfg_completo.docx", "Valdeserranos"),                                # celda de tabla
    ("04_formato_caotico.docx", "pertenencia,\nespecialmente"),               # salto de línea manual
    ("04_formato_caotico.docx", "\tLas redes sociales son plataformas"),      # tabulador inicial
])
def test_la_radiografia_ve_el_texto(nombre, fragmento):
    assert fragmento in texto_completo(radiografia(abrir(nombre)))


def test_detecta_una_palabra_cambiada_y_una_tabla_borrada():
    doc = abrir("02_informe_tablas.docx")
    antes = radiografia(doc)

    primer_texto = next(t for t in doc.element.body.iter(w("t")) if (t.text or "").strip())
    primer_texto.text = primer_texto.text.replace("sueño", "SUEÑO")
    tabla = next(doc.element.body.iter(w("tbl")))
    tabla.getparent().remove(tabla)

    diferencias = comparar(antes, radiografia(doc))
    assert any("Había 2 tablas y ahora hay 1" in d for d in diferencias)
    assert any("el texto ha cambiado" in d and "SUEÑO" in d for d in diferencias)


def test_si_el_contenido_cambia_no_se_entrega_la_copia(tmp_path, monkeypatch):
    guardar_de_verdad = documento.guardar

    def guardar_estropeando(doc, ruta):
        texto = next(t for t in doc.element.body.iter(w("t")) if (t.text or "").strip())
        texto.text = "texto alterado"
        guardar_de_verdad(doc, ruta)

    monkeypatch.setattr(documento, "guardar", guardar_estropeando)
    with pytest.raises(ErrorDocumento, match="no coincide con el original"):
        pipeline.procesar(PRUEBAS[0], tmp_path)
    assert not (tmp_path / f"{PRUEBAS[0].stem}_APA.docx").exists()
