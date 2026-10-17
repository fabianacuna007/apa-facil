"""
Paso 2.1: el verificador comprueba las reglas de la fase 1 y dice dónde falla.

- Con los originales de pruebas/: debe encontrar los incumplimientos que ya
  conocemos por el inspector.
- Con una copia procesada del documento 01 (que cumple todo) a la que se le
  introduce un solo fallo: debe encontrar ese fallo, y solo ese, en su sitio.
"""

import shutil
from pathlib import Path

import docx
import pytest
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt

from apa_facil import pipeline
from apa_facil.documento import poner_tamano, propiedades_por_defecto, w
from herramientas.verificar import (ALINEACION, FUENTE, INTERLINEADO, MARGENES, PAGINACION, SANGRIA,
                                    main, mostrar, verificar)

CARPETA = Path(__file__).resolve().parent.parent / "pruebas"
PRUEBAS = sorted(CARPETA.glob("*.docx"))

# Párrafos del documento 01 (numerados como en el verificador).
PORTADA = 2        # "Alumno: Javier Moreno Díaz"
TITULO = 6         # "Introducción"
TEXTO = 7          # "La Revolución Industrial supuso una transformación…"


@pytest.fixture(scope="module")
def base(tmp_path_factory):
    """La copia procesada del documento 01, que cumple todas las reglas."""
    carpeta = tmp_path_factory.mktemp("salida")
    salida, _ = pipeline.procesar(CARPETA / "01_ensayo_basico.docx", carpeta)
    return salida


@pytest.fixture
def doc(base, tmp_path):
    """Un documento nuevo cada vez (abierto desde una copia), para estropearlo."""
    copia = tmp_path / "copia.docx"
    shutil.copy(base, copia)
    return docx.Document(str(copia))


def parrafo(doc, numero):
    return list(doc.element.body.iter(w("p")))[numero - 1]


def incumplimientos(doc):
    return verificar(doc).incumplimientos


def unico(doc, regla):
    """El único incumplimiento del documento, que debe ser de `regla`."""
    hallados = incumplimientos(doc)
    assert [h.regla for h in hallados] == [regla], mostrar(verificar(doc))
    return hallados[0]


def espaciado(p, **atributos):
    s = p.get_or_add_pPr().get_or_add_spacing()
    for nombre, valor in atributos.items():
        s.set(w(nombre), str(valor))


# --- La base cumple -------------------------------------------------------------

def test_la_copia_procesada_cumple(doc):
    resultado = verificar(doc)
    assert resultado.cumple, mostrar(resultado)


# --- Originales: incumplimientos conocidos --------------------------------------

@pytest.mark.parametrize("ruta", PRUEBAS, ids=lambda p: p.stem)
def test_ningun_original_cumple(ruta):
    resultado = verificar(ruta)
    assert not resultado.cumple
    assert any(h.regla == MARGENES for h in resultado.incumplimientos)   # ninguno tenía 2,54 cm


def test_01_margenes_de_1_50_y_8_justificados():
    resultado = verificar(CARPETA / "01_ensayo_basico.docx")
    margenes = [h for h in resultado.incumplimientos if h.regla == MARGENES]
    assert len(margenes) == 1 and "superior 1,50 cm" in margenes[0].mensaje
    justificados = [h for h in resultado.incumplimientos if h.regla == ALINEACION]
    assert len(justificados) == 8 and all("justificado" in h.mensaje for h in justificados)


def test_02_interlineado_1_15_y_numero_en_el_pie():
    resultado = verificar(CARPETA / "02_informe_tablas.docx")
    assert any("Interlineado 1,15" in h.mensaje for h in resultado.incumplimientos if h.regla == INTERLINEADO)
    assert any("pie" in h.donde and "número de página" in h.mensaje
               for h in resultado.incumplimientos if h.regla == PAGINACION)


def test_03_encabezado_con_texto():
    resultado = verificar(CARPETA / "03_trabajo_imagenes_listas.docx")
    encabezado = [h for h in resultado.incumplimientos if h.regla == PAGINACION and h.donde.endswith("encabezado")]
    assert len(encabezado) == 1 and "Facultad de Ciencias" in encabezado[0].mensaje


def test_04_margenes_de_cada_seccion():
    resultado = verificar(CARPETA / "04_formato_caotico.docx")
    margenes = {h.seccion: h.mensaje for h in resultado.incumplimientos if h.regla == MARGENES}
    assert set(margenes) == {1, 2, 3}
    assert "superior 1,27 cm" in margenes[2]


def test_05_portada_a_la_derecha_es_una_excepcion():
    resultado = verificar(CARPETA / "05_tfg_completo.docx")
    derecha = [h for h in resultado.excepciones if h.regla == ALINEACION]
    assert [h.parrafo for h in derecha] == [8, 9, 10]
    assert not any("derecha" in h.mensaje for h in resultado.incumplimientos if h.regla == ALINEACION)


# --- 1.1 Márgenes -----------------------------------------------------------------

def test_margen_distinto(doc):
    doc.sections[0].left_margin = Cm(3)
    h = unico(doc, MARGENES)
    assert h.seccion == 1 and "izquierdo 3,00 cm" in h.mensaje


def test_margen_casi_igual_se_tolera(doc):
    doc.sections[0].top_margin = Cm(2.55)   # 0,01 cm de diferencia: no se nota
    assert incumplimientos(doc) == []


def test_encuadernacion(doc):
    doc.sections[0].gutter = Cm(1)
    assert "encuadernación de 1,00 cm" in unico(doc, MARGENES).mensaje


# --- 1.2 Fuente -------------------------------------------------------------------

def test_fragmento_en_otra_fuente(doc):
    run = docx.text.paragraph.Paragraph(parrafo(doc, TEXTO), None).runs[0]
    run.font.name = "Calibri"
    h = unico(doc, FUENTE)
    assert h.parrafo == TEXTO and "Calibri" in h.mensaje and "debe ser Arial 11 pt" in h.mensaje


def test_fragmento_en_otro_tamano(doc):
    run = docx.text.paragraph.Paragraph(parrafo(doc, TEXTO + 1), None).runs[0]
    run.font.size = Pt(14)
    h = unico(doc, FUENTE)
    assert h.parrafo == TEXTO + 1 and "Arial 14 pt" in h.mensaje


def test_tamano_heredado_de_los_valores_por_defecto(doc):
    # No basta con mirar el formato directo: en el 01 el tamaño viene de la capa
    # más profunda (los valores por defecto del documento), y también cuenta.
    poner_tamano(propiedades_por_defecto(doc), 9)
    for p in doc.element.body.iter(w("p")):
        for sz in p.iter(w("sz"), w("szCs")):
            sz.getparent().remove(sz)
    for estilo in doc.styles.element.iter(w("style")):
        for sz in estilo.iter(w("sz"), w("szCs")):
            sz.getparent().remove(sz)
    assert all(h.regla == FUENTE for h in incumplimientos(doc))
    assert any(h.parrafo == TEXTO and "Arial 9 pt" in h.mensaje for h in incumplimientos(doc))


def test_codigo_en_monoespaciada_es_una_excepcion(doc):
    nuevo = doc.add_paragraph()
    run = nuevo.add_run("total = suma(lista);")
    run.font.name = "Courier New"
    run.font.size = Pt(10)
    nuevo.paragraph_format.first_line_indent = 0   # tampoco es un párrafo de texto para la sangría
    resultado = verificar(doc)
    assert not any(h.regla == FUENTE for h in resultado.incumplimientos), mostrar(resultado)
    assert any(h.regla == FUENTE and "Código" in h.mensaje for h in resultado.excepciones)


def test_fuente_del_numero_de_pagina(doc):
    for r in doc.sections[0].header._element.iter(w("r")):
        rpr = r.get_or_add_rPr()
        fuentes = rpr.get_or_add_rFonts()
        fuentes.set(w("ascii"), "Georgia")
        fuentes.set(w("hAnsi"), "Georgia")
    h = unico(doc, FUENTE)
    assert h.donde == "Sección 1 · encabezado" and "Georgia" in h.mensaje


# --- 1.3 Interlineado -------------------------------------------------------------

def test_interlineado_1_5(doc):
    espaciado(parrafo(doc, TEXTO), line=360, lineRule="auto")
    h = unico(doc, INTERLINEADO)
    assert h.parrafo == TEXTO and "Interlineado 1,5 (debe ser doble)" in h.mensaje


def test_espacio_despues(doc):
    espaciado(parrafo(doc, TEXTO), after=120)
    assert "6 pt de espacio después" in unico(doc, INTERLINEADO).mensaje


def test_espacio_automatico_de_la_web(doc):
    espaciado(parrafo(doc, TEXTO), beforeAutospacing=1)
    assert "Espacio automático" in unico(doc, INTERLINEADO).mensaje


def test_tablas_admiten_sencillo_1_5_y_doble_pero_no_1_15(doc):
    tabla = doc.add_table(rows=1, cols=4)
    for celda, linea in zip(tabla.rows[0].cells, (240, 360, 480, 276)):
        celda.paragraphs[0].add_run("dato")
        espaciado(celda.paragraphs[0]._p, line=linea, lineRule="auto", before=0, after=0)
    h = unico(doc, INTERLINEADO)
    assert h.donde.endswith("(en una tabla)") and "Interlineado 1,15" in h.mensaje


# --- 1.4 Alineación ----------------------------------------------------------------

def test_justificado(doc):
    parrafo(doc, TEXTO).get_or_add_pPr().get_or_add_jc().set(w("val"), "both")
    h = unico(doc, ALINEACION)
    assert h.parrafo == TEXTO and "justificado" in h.mensaje


def test_a_la_derecha_fuera_de_la_portada(doc):
    parrafo(doc, TEXTO).get_or_add_pPr().get_or_add_jc().set(w("val"), "right")
    assert "derecha" in unico(doc, ALINEACION).mensaje


def test_a_la_derecha_en_la_portada_es_una_excepcion(doc):
    parrafo(doc, PORTADA).get_or_add_pPr().get_or_add_jc().set(w("val"), "right")
    resultado = verificar(doc)
    assert resultado.cumple, mostrar(resultado)
    assert [h.parrafo for h in resultado.excepciones] == [PORTADA]


def test_guiones_automaticos(doc):
    guiones = OxmlElement("w:autoHyphenation")
    doc.settings.element.append(guiones)
    assert unico(doc, ALINEACION).donde == "Ajustes del documento"


# --- 1.5 Sangría -------------------------------------------------------------------

def test_parrafo_de_texto_sin_sangria(doc):
    parrafo(doc, TEXTO).get_or_add_pPr().get_or_add_ind().set(w("firstLine"), "0")
    h = unico(doc, SANGRIA)
    assert h.parrafo == TEXTO and "tiene ninguna" in h.mensaje


def test_sangria_de_1_cm(doc):
    parrafo(doc, TEXTO).get_or_add_pPr().get_or_add_ind().set(w("firstLine"), "567")
    assert "tiene de 1,00 cm" in unico(doc, SANGRIA).mensaje


def test_titulo_con_sangria(doc):
    parrafo(doc, TITULO).get_or_add_pPr().get_or_add_ind().set(w("firstLine"), "720")
    h = unico(doc, SANGRIA)
    assert h.parrafo == TITULO and "no debe llevarla" in h.mensaje


# --- 1.6 Número de página ----------------------------------------------------------

def test_primera_pagina_diferente(doc):
    doc.sections[0].different_first_page_header_footer = True
    assert "Primera página diferente" in unico(doc, PAGINACION).mensaje


def test_pares_e_impares(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    assert "Pares e impares" in unico(doc, PAGINACION).mensaje


def test_numero_a_la_izquierda(doc):
    doc.sections[0].header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    assert "no está a la derecha" in unico(doc, PAGINACION).mensaje


def test_texto_en_el_encabezado(doc):
    doc.sections[0].header.add_paragraph("Universidad de Prueba")
    assert "Universidad de Prueba" in unico(doc, PAGINACION).mensaje


def test_numero_en_el_pie(doc):
    campo = OxmlElement("w:fldSimple")
    campo.set(w("instr"), " PAGE ")
    doc.sections[0].footer.paragraphs[0]._p.append(campo)
    h = unico(doc, PAGINACION)
    assert h.donde == "Sección 1 · pie" and "número de página" in h.mensaje


def test_numeracion_reiniciada(doc):
    tipo = doc.sections[0]._sectPr.find(w("pgNumType"))
    if tipo is None:
        tipo = OxmlElement("w:pgNumType")
        doc.sections[0]._sectPr.append(tipo)
    tipo.set(w("start"), "3")
    assert "se reinicia en 3" in unico(doc, PAGINACION).mensaje


def test_seccion_nueva_hereda_el_encabezado(doc):
    doc.add_section()
    assert incumplimientos(doc) == []


# --- Línea de comandos --------------------------------------------------------------

def test_main_codigos_de_salida(base, capsys):
    assert main([str(base)]) == 0
    assert main([str(CARPETA / "01_ensayo_basico.docx")]) == 1
    assert main([]) == 2
    assert main(["no_existe.docx"]) == 1
    salida = capsys.readouterr().out
    assert "cumple todas las reglas de la fase 1" in salida
    assert "✘ 1.1 Márgenes" in salida and "No se encuentra el archivo" in salida


def test_mostrar_resume_cuando_hay_muchos():
    texto = mostrar(verificar(CARPETA / "02_informe_tablas.docx"))
    assert "… y " in texto and "más." in texto
