"""
Paso 8.1: el servidor (POST /procesar) da lo mismo que la línea de comandos.
Se llama sin arrancarlo, con el cliente de pruebas de FastAPI.
"""

import base64
import io
import threading
import zipfile
from pathlib import Path

import docx
import pytest
from fastapi.testclient import TestClient

from apa_facil import documento, pipeline
from servidor import app as servidor_app, seguridad
from servidor.app import app

RAIZ = Path(__file__).resolve().parent.parent
PRUEBAS = sorted((RAIZ / "pruebas").glob("*.docx"))
cliente = TestClient(app)


def enviar(ruta, nombre=None, **campos):
    with open(ruta, "rb") as f:
        return cliente.post("/procesar", data=campos, files={"archivo": (nombre or ruta.name, f)})


@pytest.mark.parametrize("ruta", PRUEBAS, ids=lambda r: r.name)
def test_procesa_igual_que_la_linea_de_comandos(ruta, tmp_path):
    respuesta = enviar(ruta)
    assert respuesta.status_code == 200, respuesta.text
    datos = respuesta.json()
    assert datos["documento"]["nombre"] == f"{ruta.stem}_APA.docx"
    assert datos["informe"]["archivo"] == ruta.name
    assert datos["informe"]["copia"] == f"{ruta.stem}_APA.docx"

    # La copia que devuelve se abre y es la misma que guarda la línea de comandos
    # (el pipeline ya ha comprobado que conserva el texto del original).
    copia = docx.Document(io.BytesIO(base64.b64decode(datos["documento"]["contenido"])))
    original = documento.abrir(ruta)
    assert len(copia.tables) == len(original.tables)
    assert len(copia.inline_shapes) == len(original.inline_shapes)

    # El informe es el mismo que da la línea de comandos (salvo la fecha).
    _, informe = pipeline.procesar(ruta, tmp_path)
    esperado, recibido = informe.a_dict(), datos["informe"]
    for clave in ("cambios", "avisos", "errores", "resumen"):
        assert recibido[clave] == esperado[clave]
    assert datos["resumen"] == informe.como_texto()


def test_las_opciones_llegan_al_nucleo():
    datos = enviar(PRUEBAS[1], fuente="arial", comentarios="true", quitar_indice="true",
                   titulo="Mi trabajo", autor=["Ana Pérez", "Luis Gómez"], fecha="16/10/2026").json()
    opciones = datos["informe"]["opciones"]
    assert opciones["fuente"] == "Arial"
    assert opciones["comentarios"] and opciones["quitar_indice"] and not opciones["mantener_portada"]
    assert "titulo" in opciones["portada"]["dados"] and "autores" in opciones["portada"]["dados"]


def test_campos_vacios_cuentan_como_no_enviados():
    datos = enviar(PRUEBAS[1], fuente="", titulo="", fecha="").json()
    assert datos["informe"]["opciones"]["fuente"] is None
    assert datos["informe"]["opciones"]["portada"] is None


@pytest.mark.parametrize("campos, palabra", [({"fuente": "Comic Sans"}, "fuentes recomendadas"),
                                             ({"fecha": "mañana"}, "fecha")])
def test_opciones_no_validas(campos, palabra):
    respuesta = enviar(PRUEBAS[1], **campos)
    assert respuesta.status_code == 422
    assert palabra in respuesta.json()["error"]


def test_archivo_que_no_es_word(tmp_path):
    pdf = tmp_path / "trabajo.pdf"
    pdf.write_bytes(b"%PDF-1.4 no es un Word")
    respuesta = enviar(pdf)
    assert respuesta.status_code == 422
    assert respuesta.json()["error"] == "«trabajo.pdf» no es un documento de Word. Solo se admiten archivos .docx."

    renombrado = tmp_path / "trabajo.docx"
    renombrado.write_bytes(pdf.read_bytes())
    error = enviar(renombrado).json()["error"]
    assert "«trabajo.docx» está dañado" in error and "documento.docx" not in error


def test_nombre_reservado_de_windows():
    respuesta = enviar(PRUEBAS[1], nombre="con.docx")
    assert respuesta.status_code == 200
    assert respuesta.json()["documento"]["nombre"] == "con_APA.docx"


def test_falta_el_archivo():
    respuesta = cliente.post("/procesar", data={"fuente": "Arial"})
    assert respuesta.status_code == 422
    assert respuesta.json() == {"error": "Falta el documento: sube un archivo de Word (.docx)."}


# --- Paso 8.2: seguridad de entrada --------------------------------------------------


def un_zip(ruta, piezas):
    with zipfile.ZipFile(ruta, "w", zipfile.ZIP_DEFLATED) as z:
        for nombre, contenido in piezas.items():
            z.writestr(nombre, contenido)
    return ruta


def test_archivo_vacio(tmp_path):
    vacio = tmp_path / "trabajo.docx"
    vacio.write_bytes(b"")
    respuesta = enviar(vacio)
    assert respuesta.status_code == 422
    assert respuesta.json()["error"].startswith("«trabajo.docx» está vacío")


def test_archivo_demasiado_grande(tmp_path):
    grande = tmp_path / "trabajo.docx"
    grande.write_bytes(b"PK" + b"\0" * (seguridad.TAMANO_MAXIMO + 2 * seguridad.MB))
    respuesta = enviar(grande)
    assert respuesta.status_code == 413
    assert "más de 10 MB" in respuesta.json()["error"]


def test_archivo_algo_mayor_que_el_maximo(tmp_path):
    # Cabe en la petición (menos de 10 MB + 1 MB) pero el archivo pasa de 10 MB.
    grande = tmp_path / "trabajo.docx"
    grande.write_bytes(b"PK" + b"\0" * (seguridad.TAMANO_MAXIMO + 1000))
    respuesta = enviar(grande)
    assert respuesta.status_code == 413
    assert respuesta.json()["error"].startswith("«trabajo.docx» ocupa más de 10 MB")


def test_bomba_zip(tmp_path):
    bomba = un_zip(tmp_path / "trabajo.docx", {"[Content_Types].xml": "<Types/>",
                                                "word/document.xml": b"\0" * (50 * seguridad.MB)})
    assert bomba.stat().st_size < seguridad.MB
    respuesta = enviar(bomba)
    assert respuesta.status_code == 422
    assert "contenido comprimido anómalo" in respuesta.json()["error"]


def test_zip_que_no_es_word(tmp_path):
    respuesta = enviar(un_zip(tmp_path / "fotos.docx", {"foto.txt": "hola"}))
    assert respuesta.status_code == 422
    assert "«fotos.docx» está dañado" in respuesta.json()["error"]


def test_word_con_contrasena(tmp_path):
    ole = tmp_path / "trabajo.docx"
    ole.write_bytes(seguridad.OLE + b"\0" * 1000)
    respuesta = enviar(ole)
    assert respuesta.status_code == 422
    assert "contraseña" in respuesta.json()["error"]


def test_tiempo_maximo(monkeypatch):
    monkeypatch.setattr(servidor_app, "TIEMPO_MAXIMO", 0.01)
    respuesta = enviar(PRUEBAS[1])
    assert respuesta.status_code == 504
    assert "ha tardado demasiado" in respuesta.json()["error"]


def test_servidor_ocupado(monkeypatch):
    monkeypatch.setattr(servidor_app, "TURNOS", threading.BoundedSemaphore(1))
    monkeypatch.setattr(servidor_app, "ESPERA_MAXIMA", 0.01)
    servidor_app.TURNOS.acquire()
    respuesta = enviar(PRUEBAS[1])
    assert respuesta.status_code == 503
    assert "en cola" in respuesta.json()["error"]


# --- Paso 8.3: privacidad ------------------------------------------------------------

NOMBRE_PRIVADO = "TFG de Ana Pérez.docx"


@pytest.fixture
def temporal(tmp_path, monkeypatch):
    carpeta = tmp_path / "temporal"
    carpeta.mkdir()
    monkeypatch.setattr(servidor_app, "CARPETA_TEMPORAL", str(carpeta))
    return carpeta


@pytest.mark.parametrize("caso", ["bien", "dañado", "bomba", "tiempo"])
def test_no_queda_nada_en_la_carpeta_temporal(caso, temporal, tmp_path, monkeypatch):
    ruta = PRUEBAS[4]
    if caso == "dañado":
        ruta = tmp_path / "x.docx"
        ruta.write_bytes(b"PK no es un zip")
    elif caso == "bomba":
        ruta = un_zip(tmp_path / "x.docx", {"[Content_Types].xml": "<Types/>",
                                             "word/document.xml": b"\0" * (50 * seguridad.MB)})
    elif caso == "tiempo":
        monkeypatch.setattr(servidor_app, "TIEMPO_MAXIMO", 0.5)
    respuesta = enviar(ruta)
    assert respuesta.status_code == {"bien": 200, "dañado": 422, "bomba": 422, "tiempo": 504}[caso]
    assert list(temporal.iterdir()) == []


def test_el_registro_no_lleva_nombre_ni_texto(temporal, caplog, monkeypatch):
    texto = next(p.text for p in documento.abrir(PRUEBAS[4]).paragraphs if len(p.text) > 20)
    assert texto
    caplog.set_level("DEBUG")
    assert enviar(PRUEBAS[4], nombre=NOMBRE_PRIVADO).status_code == 200

    def falla(ruta, nombre):
        raise RuntimeError(f"fallo con el texto {texto} de {nombre}")
    monkeypatch.setattr(seguridad, "comprobar", falla)
    respuesta = enviar(PRUEBAS[4], nombre=NOMBRE_PRIVADO)
    assert respuesta.status_code == 500
    assert texto not in respuesta.text and "RuntimeError" not in respuesta.text
    assert "Error inesperado en el servidor (RuntimeError)" in caplog.text
    assert "Ana Pérez" not in caplog.text and texto not in caplog.text
    assert list(temporal.iterdir()) == []


# --- Paso 8.4: más archivos inválidos y casos de la web --------------------------------

@pytest.mark.parametrize("nombre, contenido, mensaje", [
    ("trabajo.doc", b"\xd0\xcf\x11\xe0" + b"\0" * 100, "«trabajo.doc» está en el formato antiguo de Word (.doc)."),
    ("~$trabajo.docx", b"PK", "«~$trabajo.docx» es un archivo temporal que Word crea"),
    ("notas.txt", b"hola", "«notas.txt» no es un documento de Word."),
    ("trabajo", b"PK", "«trabajo» no es un documento de Word."),
])
def test_nombres_no_validos(nombre, contenido, mensaje, tmp_path):
    ruta = tmp_path / "archivo"
    ruta.write_bytes(contenido)
    respuesta = enviar(ruta, nombre=nombre)
    assert respuesta.status_code == 422
    assert respuesta.json()["error"].startswith(mensaje)


def test_pdf_de_verdad_renombrado(tmp_path):
    # Un PDF real (cabecera, objetos y final), con la extensión cambiada.
    pdf = tmp_path / "Trabajo final.docx"
    pdf.write_bytes(b"%PDF-1.7\n1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
                    b"2 0 obj << /Type /Pages /Kids [] /Count 0 >> endobj\ntrailer << /Root 1 0 R >>\n%%EOF\n")
    respuesta = enviar(pdf)
    assert respuesta.status_code == 422
    assert respuesta.json() == {"error": "«Trabajo final.docx» está dañado o no es un documento de Word de verdad "
                                         "(aunque termine en .docx)."}


def test_nombre_con_ruta_del_navegador():
    respuesta = enviar(PRUEBAS[0], nombre="C:\fakepath\mi trabajo.docx")
    assert respuesta.status_code == 200
    assert respuesta.json()["documento"]["nombre"] == "mi trabajo_APA.docx"
    assert respuesta.json()["informe"]["archivo"] == "mi trabajo.docx"


def test_comentarios_en_la_copia():
    datos = enviar(PRUEBAS[0], comentarios="true").json()
    with zipfile.ZipFile(io.BytesIO(base64.b64decode(datos["documento"]["contenido"]))) as z:
        assert "word/comments.xml" in z.namelist()


def test_los_errores_son_json_con_mensaje_en_espanol(tmp_path):
    vacio = tmp_path / "a.docx"
    vacio.write_bytes(b"")
    for respuesta in (enviar(vacio), cliente.post("/procesar"), enviar(PRUEBAS[1], fuente="Papyrus")):
        assert respuesta.headers["content-type"] == "application/json"
        assert list(respuesta.json()) == ["error"]
