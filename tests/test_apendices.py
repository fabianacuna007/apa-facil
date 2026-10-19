"""
Paso 6.7: los apéndices, según la web oficial de APA (Appendices setup): cada uno
en una página nueva, con el rótulo («Apéndice», o «Apéndice A», «B»… si hay
varios) y, en la línea siguiente, su título, los dos centrados y en negrita.
«Anexo» se deja por defecto y se avisa (ROADMAP y decisión del propietario);
--rotulo-apendice lo cambia.
"""

from pathlib import Path

import docx
import pytest

from apa_facil import estructura, pipeline
from apa_facil.documento import CambioPrevisto, Estilos, comparar, radiografia, texto_parrafo, w
from apa_facil.informe import Informe
from apa_facil.reglas import apendices as regla
from apa_facil.reglas.alineacion import _salto_antes
from herramientas.verificar import APENDICES, verificar

RAIZ = Path(__file__).resolve().parent.parent
TFG = RAIZ / "pruebas" / "05_tfg_completo.docx"
TEXTO = ("El sueño cumple funciones esenciales en la consolidación de la memoria y en la regulación emocional, según "
         "los estudios revisados (Navarro y Gil, 2020).")


def trabajo(*apendices):
    """Un trabajo con referencias y, al final, los apéndices: (rótulo, [párrafos]) cada uno."""
    doc = docx.Document()
    doc.add_heading("Introducción", level=1)
    doc.add_paragraph(TEXTO)
    doc.add_heading("Referencias", level=1)
    doc.add_paragraph("Navarro, A. y Gil, S. (2020). Sueño, memoria y aprendizaje. Editorial Ejemplo.")
    for rotulo, parrafos in apendices:
        doc.add_heading(rotulo, level=1)
        for texto in parrafos:
            doc.add_paragraph(texto)
    return doc


def procesar(doc, tmp_path, opciones=None):
    ruta = tmp_path / "trabajo.docx"
    doc.save(str(ruta))
    salida, informe = pipeline.procesar(ruta, tmp_path / "salida", opciones)
    return docx.Document(str(salida)), informe


def parrafo(doc, inicio):
    return next(p for p in doc.element.body.iter(w("p")) if texto_parrafo(p).startswith(inicio))


def rotulos_y_titulos(doc):
    return [(texto_parrafo(r.elemento), texto_parrafo(t.elemento) if t is not None else None)
            for r, t in regla.apendices(estructura.clasificar(doc))]


def incumplimientos(doc):
    return [h.mensaje for h in verificar(doc).incumplimientos if h.regla == APENDICES]


# --- El 05 ---------------------------------------------------------------------------------------

@pytest.fixture(scope="module")
def tfg(tmp_path_factory):
    return pipeline.procesar(TFG, tmp_path_factory.mktemp("apendices"))


def test_el_anexo_del_05(tfg):
    salida, informe = tfg
    doc = docx.Document(str(salida))
    estilos = Estilos(doc)
    assert rotulos_y_titulos(doc) == [("Anexo I", "Guion de las Entrevistas")]
    rotulo, titulo = parrafo(doc, "Anexo I"), parrafo(doc, "Guion de las")
    assert estilos.nombre(estilos.estilo_de(rotulo)).lower() == "heading 1"
    assert _salto_antes(estilos, rotulo) and not _salto_antes(estilos, titulo)
    assert estilos.parrafo(titulo, "jc", "val") == "center"
    assert texto_parrafo(rotulo.getprevious()).strip(), "ya no está el párrafo vacío del salto"
    assert any(a.startswith("APA llama «Apéndice» a los anexos") for a in informe.avisos)
    assert incumplimientos(doc) == []
    assert incumplimientos(docx.Document(str(TFG))) == ["El rótulo y el título del apéndice van en la misma línea."]


def test_el_05_con_rotulo_apendice(tmp_path):
    salida, informe = pipeline.procesar(TFG, tmp_path, {"rotulo_apendice": True})
    doc = docx.Document(str(salida))
    assert rotulos_y_titulos(doc) == [("Apéndice", "Guion de las Entrevistas")]
    assert any(c.startswith("Apéndices: «Anexo I» → «Apéndice»") for c in informe.cambios)
    assert not [h for h in verificar(doc).excepciones if h.regla == APENDICES]


# --- Casos hechos a mano --------------------------------------------------------------------------

def test_varios_apendices_con_letra(tmp_path):
    original = trabajo(("Anexo 1: cuestionario", [TEXTO]), ("Anexo 2: entrevistas", [TEXTO]))
    doc, _ = procesar(original, tmp_path, {"rotulo_apendice": True})
    assert rotulos_y_titulos(doc) == [("Apéndice A", "Cuestionario"), ("Apéndice B", "Entrevistas")]
    assert incumplimientos(doc) == []


def test_rotulo_y_titulo_ya_separados(tmp_path):
    doc, informe = procesar(trabajo(("Apéndice", ["Guion de las entrevistas", TEXTO])), tmp_path)
    assert rotulos_y_titulos(doc) == [("Apéndice", "Guion de las Entrevistas")]
    assert not any(a.startswith("APA llama «Apéndice»") for a in informe.avisos)
    assert incumplimientos(doc) == []


def test_con_titulos_tal_cual(tmp_path):
    doc, _ = procesar(trabajo(("Anexo I. Guion de las entrevistas", [TEXTO])), tmp_path, {"titulos_tal_cual": True})
    assert rotulos_y_titulos(doc) == [("Anexo I", "Guion de las entrevistas")]


def test_apendice_sin_titulo(tmp_path):
    doc, informe = procesar(trabajo(("Anexo", [TEXTO])), tmp_path)
    assert any(a.startswith("«Anexo» no tiene título") for a in informe.avisos)
    assert incumplimientos(doc) == ["El apéndice no tiene título en la línea siguiente al rótulo."]


def test_tabla_de_un_apendice_sin_letra():
    doc = trabajo(("Anexo I. Datos", [TEXTO, "Tabla 3", "Población por municipio"]))
    doc.add_table(rows=2, cols=2)
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    assert any(a.startswith("En los apéndices, APA numera las tablas y figuras con la letra") for a in informe.avisos)


def test_apendice_antes_de_las_referencias():
    doc = docx.Document()
    doc.add_paragraph(TEXTO)
    doc.add_heading("Anexo I. Datos", level=1)
    doc.add_paragraph(TEXTO)
    doc.add_heading("Referencias", level=1)
    doc.add_paragraph("Navarro, A. y Gil, S. (2020). Sueño, memoria y aprendizaje. Editorial Ejemplo.")
    informe = Informe()
    regla.aplicar(doc, {}, informe)
    assert any(a.startswith("Los apéndices van al final del trabajo, después de las referencias") for a in informe.avisos)


# --- Integridad ----------------------------------------------------------------------------------------

def _radiografias(antes, despues):
    def radio(textos):
        doc = docx.Document()
        for texto in textos:
            doc.add_paragraph(texto)
        return radiografia(doc)
    return radio(antes), radio(despues)


@pytest.mark.parametrize("antes, despues, cambios, admitido", [
    (["Anexo I. Guion"], ["Anexo I", "Guion"], [CambioPrevisto(1, "dividir_apendice", quitado=". ", nuevo="Guion")],
     True),
    (["Anexo I. Guion"], ["Apéndice", "Guion"], [CambioPrevisto(1, "dividir_apendice", quitado=". ", nuevo="Guion"),
                                                 CambioPrevisto(1, "rotulo_de_apendice", nuevo="Apéndice")], True),
    (["Anexo 2"], ["Apéndice B"], [CambioPrevisto(1, "rotulo_de_apendice", nuevo="Apéndice B")], True),
    (["Anexo I. Guion"], ["Anexo I", "Otro"], [CambioPrevisto(1, "dividir_apendice", quitado=". ", nuevo="Otro")],
     False),                                                       # otro título
    (["Anexo 2"], ["Anexo B"], [CambioPrevisto(1, "rotulo_de_apendice", nuevo="Anexo B")], False),   # no es de APA
    (["Anexo I. Guion"], ["Apéndice"], [CambioPrevisto(1, "rotulo_de_apendice", nuevo="Apéndice")], False),  # se perdería el título
    (["Conclusiones"], ["Apéndice"], [CambioPrevisto(1, "rotulo_de_apendice", nuevo="Apéndice")], False),
])
def test_integridad(antes, despues, cambios, admitido):
    a, d = _radiografias(antes, despues)
    assert (comparar(a, d, cambios) == []) == admitido


def test_el_verificador_ve_que_no_empieza_pagina(tmp_path):
    doc, _ = procesar(trabajo(("Anexo I. Datos", [TEXTO])), tmp_path)
    salto = parrafo(doc, "Anexo I").find(f"{w('pPr')}/{w('pageBreakBefore')}")
    salto.getparent().remove(salto)
    assert incumplimientos(doc) == ["El apéndice no empieza en una página nueva."]
